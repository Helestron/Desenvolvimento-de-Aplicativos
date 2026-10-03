"""Apoio aos testes da transcrição: sinais sintéticos, modelo dublê e uma
configuração apontando para uma pasta temporária.

Nada aqui usa rede, microfone ou tela. O "modelo dublê" imita a interface
do WhisperModel (`transcribe(audio, **kw) -> (segmentos, info)`): acha as
rajadas de som no áudio recebido e devolve, para cada uma, o texto
associado à frequência fundamental dela - assim os testes conferem que cada
fala chegou com o texto e o falante certos, sem modelo de verdade.
"""

from __future__ import annotations

import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np

TAXA = 16000
NUMERO = "0700123-83.2024.8.02.0001"        # dígito verificador correto
NUMERO_CI = "0700123-45.2024.8.02.0001"     # o do CI (dígito não confere)

TEXTOS = {
    200: "Bom dia a todos.",
    300: "Sem perguntas, Excelência.",
    400: "Eu vi o acidente.",
    500: "Esta fala aconteceu durante a pausa.",
    250: "Nada mais.",
    700: "Legendas pela comunidade Amara.org",
}


# ------------------------------------------------------------- sinais
def ruido(duracao: float, amp: float = 0.002, semente: int = 0) -> np.ndarray:
    rng = np.random.default_rng(semente)
    return (rng.normal(0.0, amp, int(round(duracao * TAXA)))).astype(np.float32)


def fala(duracao: float, f0: float, amp: float = 0.1) -> np.ndarray:
    """Som "de fala": harmônicos de f0 com sílabas (~5 por segundo)."""
    n = int(round(duracao * TAXA))
    t = np.arange(n) / TAXA
    tom = (np.sin(2 * np.pi * f0 * t) + 0.5 * np.sin(2 * np.pi * 2 * f0 * t)
           + 0.25 * np.sin(2 * np.pi * 3 * f0 * t)) / 1.75
    silabas = 0.35 + 0.65 * np.abs(np.sin(2 * np.pi * 2.5 * t))
    rampa = np.minimum(1.0, np.minimum(t, t[-1] - t) / 0.02) if n > 1 else np.ones(n)
    return (amp * tom * silabas * rampa).astype(np.float32)


def montar(partes: list[tuple], semente: int = 0) -> np.ndarray:
    """partes: ("silencio", dur) | ("fala", dur, f0) | ("pausa_curta", dur)."""
    blocos = []
    for parte in partes:
        if parte[0] == "fala":
            blocos.append(fala(parte[1], parte[2]))
        else:
            blocos.append(np.zeros(int(round(parte[1] * TAXA)), dtype=np.float32))
    audio = np.concatenate(blocos)
    return (audio + ruido(audio.size / TAXA, semente=semente)).astype(np.float32)


def gravar_wav(caminho: Path, audio: np.ndarray, taxa: int = TAXA) -> Path:
    import soundfile as sf

    sf.write(str(caminho), audio, taxa, subtype="PCM_16")
    return caminho


def roteiro_audiencia() -> np.ndarray:
    """O WAV da audiência do teste de ponta a ponta (18 s)."""
    return montar([
        ("silencio", 0.5), ("fala", 2.0, 200),     # 0.5-2.5   Juiz(a)
        ("silencio", 1.8), ("fala", 2.0, 300),     # 4.3-6.3   Promotor(a)
        ("silencio", 1.7), ("fala", 2.5, 400),     # 8.0-10.5  Testemunha
        ("silencio", 1.5), ("fala", 1.5, 500),     # 12.0-13.5 (gravação pausada)
        ("silencio", 1.5), ("fala", 1.5, 250),     # 15.0-16.5 Testemunha
        ("silencio", 1.5),
    ])


# ------------------------------------------------------------ dublês
@dataclass
class Palavra:
    start: float
    end: float
    word: str


@dataclass
class SegmentoDuble:
    start: float
    end: float
    text: str
    no_speech_prob: float = 0.02
    avg_logprob: float = -0.25
    compression_ratio: float = 1.3
    words: list | None = None


@dataclass
class InfoDuble:
    language: str = "pt"
    language_probability: float = 1.0
    duration: float = 0.0


def rajadas(audio: np.ndarray, taxa: int = TAXA) -> list[tuple[float, float]]:
    """Trechos com som bem acima do ruído (pausas < 0,5 s não separam)."""
    q = int(0.03 * taxa)
    n = audio.size // q
    if n == 0:
        return []
    quadros = audio[:n * q].reshape(n, q)
    energia = 10 * np.log10(np.mean(np.square(quadros, dtype=np.float64), axis=1) + 1e-12)
    ativo = energia > max(np.max(energia) - 25.0, -45.0)
    trechos: list[list[int]] = []
    for i, a in enumerate(ativo):
        if not a:
            continue
        if trechos and i - trechos[-1][1] <= int(0.5 / 0.03):
            trechos[-1][1] = i
        else:
            trechos.append([i, i])
    return [(a * q / taxa, (b + 1) * q / taxa) for a, b in trechos if (b - a + 1) * 0.03 >= 0.2]


def frequencia(audio: np.ndarray, taxa: int = TAXA) -> float:
    janela = audio * np.hanning(audio.size)
    espectro = np.abs(np.fft.rfft(janela))
    freqs = np.fft.rfftfreq(audio.size, 1 / taxa)
    espectro[freqs < 120] = 0
    return float(freqs[int(np.argmax(espectro))])


def texto_para(f: float, textos: dict[int, str] = TEXTOS) -> str | None:
    chave = min(textos, key=lambda k: abs(k - f))
    return textos[chave] if abs(chave - f) < 30 else None


@dataclass
class ModeloDuble:
    """Imita o WhisperModel. Guarda os argumentos de cada chamada."""

    textos: dict[int, str] = field(default_factory=lambda: dict(TEXTOS))
    chamadas: list[dict] = field(default_factory=list)
    falhar_em: int = -1         # índice da chamada que levanta erro (testes de falha)

    def transcribe(self, audio, **kw):
        self.chamadas.append(dict(kw))
        if len(self.chamadas) - 1 == self.falhar_em:
            raise RuntimeError("falha simulada do modelo")
        audio = np.asarray(audio, dtype=np.float32)
        segmentos = []
        for ini, fim in rajadas(audio):
            pedaco = audio[int(ini * TAXA):int(fim * TAXA)]
            texto = texto_para(frequencia(pedaco))
            if texto is None:   # ruído: o VAD do Whisper descartaria
                continue
            palavras = None
            if kw.get("word_timestamps"):
                termos = texto.split()
                passo = (fim - ini) / max(1, len(termos))
                palavras = [Palavra(ini + i * passo, ini + (i + 1) * passo, " " + w)
                            for i, w in enumerate(termos)]
            segmentos.append(SegmentoDuble(round(ini, 2), round(fim, 2), " " + texto,
                                           words=palavras))
        return iter(segmentos), InfoDuble(duration=audio.size / TAXA)


# ------------------------------------------------------- configuração
class PastaTemporaria:
    """Uma pasta temporária com config.ini cujo acervo fica dentro dela."""

    def __init__(self, **ajustes: str):
        self.raiz = Path(tempfile.mkdtemp(prefix="assessor-transcricao-"))
        self.acervo = self.raiz / "Acervo"
        self.ini = self.raiz / "config.ini"
        linhas = ["[geral]", f"pasta_acervo = {self.acervo}",
                  f"pasta_sigilosos = {self.raiz / 'Sigilosos'}",
                  "[unidade]", "magistrado = Fulana de Tal", "cargo = Juíza de Direito",
                  "vara = 1ª Vara Cível", "comarca = Maceió", "tribunal = TJAL",
                  "[transcricao]"]
        padrao = {"modelo_ao_vivo": "small", "modelo_revisao": "medium",
                  "separar_falantes": "false", "salvar_audio": "true", "marcar_tempo": "true"}
        padrao.update(ajustes)
        linhas += [f"{k} = {v}" for k, v in padrao.items()]
        self.ini.write_text("\n".join(linhas) + "\n", encoding="utf-8")

    def config(self):
        from app.nucleo.config import Config

        return Config(self.ini)

    def apagar(self) -> None:
        shutil.rmtree(self.raiz, ignore_errors=True)


def ler_docx(caminho: Path) -> tuple[list[str], list[list[list[str]]], str]:
    """(parágrafos, tabelas [[células]], texto do rodapé) de um DOCX."""
    from docx import Document

    doc = Document(str(caminho))
    paragrafos = [p.text for p in doc.paragraphs]
    tabelas = [[[c.text for c in linha.cells] for linha in t.rows] for t in doc.tables]
    rodape = " ".join(p.text for p in doc.sections[0].footer.paragraphs)
    return paragrafos, tabelas, rodape


def ficha(tabelas) -> dict[str, str]:
    return {linha[0]: linha[1] for linha in tabelas[0]} if tabelas else {}
