"""Testes do núcleo: caminhos, número CNJ, tribunais, configuração, senhas e listas."""

from __future__ import annotations

import configparser
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

from helestron.nucleo import (caminhos, cnj, cofre_senhas, config, listas, sigilo, sistema,
                             tribunais)

REPOSITORIO = Path(caminhos.__file__).resolve().parents[2]

# Lê as constantes de caminhos.py num processo novo: elas são calculadas na
# importação, a partir do ambiente e do sys.prefix.
_SONDA_CAMINHOS = r"""
import json, sys
prefixo = sys.argv[1]
if prefixo:
    sys.prefix = prefixo
from helestron.nucleo import caminhos
nomes = ["PACOTE", "INSTALADO", "INSTALACAO", "DADOS", "RECURSOS", "WEB", "MODELOS_EMBUTIDOS",
         "LOCAL", "ARQUIVO_CONFIG", "LOGS", "PERFIS", "ARQUIVO_SENHAS", "TEMP", "MODELOS",
         "ARQUIVO_PAUTA", "ARQUIVO_INSTANCIA", "BASE_USUARIO"]
saida = {n: (getattr(caminhos, n) if isinstance(getattr(caminhos, n), bool)
             else str(getattr(caminhos, n))) for n in nomes}
saida["python"] = str(caminhos.python_exe())
saida["pythonw"] = str(caminhos.python_exe(janela=True))
saida["relativo"] = str(caminhos.resolver("Outra", "Acervo"))
saida["vazio"] = str(caminhos.resolver("", "Acervo"))
print(json.dumps(saida))
"""


def ler_caminhos(ambiente: dict, prefixo: str = "") -> dict:
    env = {k: v for k, v in os.environ.items()
           if k not in ("HELESTRON_LOCAL", "HELESTRON_DADOS", "LOCALAPPDATA", "OneDrive",
                        "OneDriveCommercial", "OneDriveConsumer", "USERPROFILE")}
    env.update(ambiente)
    env["PYTHONPATH"] = str(REPOSITORIO)
    r = subprocess.run([sys.executable, "-c", _SONDA_CAMINHOS, prefixo], capture_output=True,
                       text=True, env=env, cwd=str(REPOSITORIO), timeout=120)
    if r.returncode != 0:
        raise AssertionError(r.stderr)
    return json.loads(r.stdout)


class TestCaminhos(unittest.TestCase):
    """O contrato da especificação (seção 4): nada do usuário na pasta do programa."""

    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.base = Path(self.dir.name)

    def tearDown(self):
        self.dir.cleanup()

    def test_variaveis_dos_testes_isolam_tudo(self):
        local, dados = self.base / "local", self.base / "docs"
        c = ler_caminhos({"HELESTRON_LOCAL": str(local), "HELESTRON_DADOS": str(dados)})
        self.assertFalse(c["INSTALADO"])
        self.assertEqual(c["INSTALACAO"], str(REPOSITORIO))
        self.assertEqual(c["PACOTE"], str(REPOSITORIO / "helestron"))
        self.assertEqual(c["WEB"], str(REPOSITORIO / "helestron" / "web"))
        self.assertEqual(c["DADOS"], str(REPOSITORIO / "helestron" / "dados"))
        self.assertEqual(c["RECURSOS"], str(REPOSITORIO / "helestron" / "recursos"))
        self.assertEqual(c["MODELOS_EMBUTIDOS"], str(REPOSITORIO / "modelos"))
        self.assertEqual(c["LOCAL"], str(local))
        for nome, rel in (("ARQUIVO_CONFIG", "config.ini"), ("LOGS", "Logs"), ("PERFIS", "perfis"),
                          ("ARQUIVO_SENHAS", "credenciais.json"), ("TEMP", "temp"),
                          ("MODELOS", "modelos"), ("ARQUIVO_PAUTA", "pauta.sqlite3"),
                          ("ARQUIVO_INSTANCIA", "instancia.json")):
            self.assertEqual(c[nome], str(local / rel), nome)
        self.assertEqual(c["BASE_USUARIO"], str(dados))
        # caminho relativo do config.ini: dentro de Documentos\Helestron
        self.assertEqual(c["relativo"], str(dados / "Outra"))
        self.assertEqual(c["vazio"], str(dados / "Acervo"))
        self.assertFalse(list(self.base.iterdir()), "importar caminhos.py não cria pasta nenhuma")

    def test_padroes_localappdata_e_documentos(self):
        casa = self.base / "casa"
        c = ler_caminhos({"LOCALAPPDATA": str(self.base / "AppData" / "Local"), "HOME": str(casa)})
        self.assertEqual(c["LOCAL"], str(self.base / "AppData" / "Local" / "Helestron"))
        if sys.platform != "win32":
            self.assertEqual(c["BASE_USUARIO"], str(casa / "Documents" / "Helestron"))
            # sem LOCALAPPDATA (fora do Windows): ~/.helestron
            c = ler_caminhos({"HOME": str(casa)})
            self.assertEqual(c["LOCAL"], str(casa / ".helestron"))

    @unittest.skipIf(sys.platform == "win32", "no Windows, a pasta Documentos vem da API")
    def test_documentos_no_onedrive_muda_a_base(self):
        # Documentos redirecionado para o OneDrive: a base vai para o perfil,
        # fora da sincronização (que trava arquivo em uso).
        casa = self.base / "casa"
        c = ler_caminhos({"HOME": str(casa), "OneDrive": str(casa / "Documents"),
                          "USERPROFILE": str(casa)})
        self.assertEqual(c["BASE_USUARIO"], str(casa / "Helestron"))

    def test_documentos_no_google_drive_muda_a_base(self):
        # Documentos dentro do Google Drive (modo espelho, %USERPROFILE%\Meu
        # Drive): a pasta padrão dos sigilosos e da pauta não pode ficar na
        # nuvem - a base vai para o perfil, como com o OneDrive.
        casa = self.base / "casa"
        sem_onedrive = {"OneDrive": "", "OneDriveCommercial": "", "OneDriveConsumer": ""}
        with mock.patch.dict(os.environ, dict(sem_onedrive, HOME=str(casa), USERPROFILE=str(casa),
                                              HELESTRON_DADOS="")):
            for documentos in (casa / "Meu Drive" / "Documentos", casa / "My Drive (2)" / "Docs",
                               casa / "Google Drive" / "Documentos"):
                with self.subTest(documentos=documentos), \
                        mock.patch.object(caminhos, "pasta_documentos", return_value=documentos):
                    self.assertTrue(caminhos.no_google_drive(documentos))
                    self.assertEqual(caminhos._base_usuario(), casa / "Helestron")
            with mock.patch.object(caminhos, "pasta_documentos", return_value=casa / "Documents"):
                self.assertFalse(caminhos.no_google_drive(casa / "Documents"))
                self.assertEqual(caminhos._base_usuario(), casa / "Documents" / "Helestron")

    def test_instalado_pelo_manifesto(self):
        programa = self.base / "Programs" / "Helestron"
        programa.mkdir(parents=True)
        for nome in ("manifesto.json", "python.exe", "pythonw.exe"):
            (programa / nome).write_text("{}", encoding="utf-8")
        c = ler_caminhos({"HELESTRON_LOCAL": str(self.base / "l"),
                          "HELESTRON_DADOS": str(self.base / "d")}, prefixo=str(programa))
        self.assertTrue(c["INSTALADO"])
        self.assertEqual(c["INSTALACAO"], str(programa))
        self.assertEqual(c["MODELOS_EMBUTIDOS"], str(programa / "modelos"))
        self.assertEqual(c["python"], str(programa / "python.exe"))
        self.assertEqual(c["pythonw"], str(programa / "pythonw.exe"))

    def test_resolver(self):
        with mock.patch.object(caminhos, "BASE_USUARIO", self.base):
            self.assertEqual(caminhos.resolver("Acervo 2", "Acervo"), self.base / "Acervo 2")
            self.assertEqual(caminhos.resolver(None, "Acervo"), self.base / "Acervo")
            self.assertEqual(caminhos.resolver(f'"{self.base / "x"}"', "Acervo"), self.base / "x")
        with mock.patch.dict(os.environ, {"HELESTRON_TESTE": str(self.base / "var")}):
            self.assertEqual(caminhos.resolver("$HELESTRON_TESTE" if os.name != "nt"
                                               else "%HELESTRON_TESTE%", "Acervo"),
                             self.base / "var")

    def test_onedrive(self):
        with mock.patch.dict(os.environ, {"OneDrive": str(self.base / "OneDrive - TJ")}):
            self.assertTrue(caminhos.dentro_do_onedrive(self.base / "OneDrive - TJ" / "Acervo"))
            self.assertFalse(caminhos.dentro_do_onedrive(self.base / "Outra"))

    def test_enderecos_corrigidos_ficam_na_pasta_de_dados(self):
        # Fora da pasta do programa: a atualização não os apaga.
        self.assertEqual(tribunais.ARQUIVO_LOCAL.parent, caminhos.LOCAL)
        alvo = self.base / "Local" / "ainda-nao-existe" / "enderecos-locais.json"
        with mock.patch.object(tribunais, "ARQUIVO_LOCAL", alvo):
            tribunais.definir_endereco("eproc:TJAL", "1g", "https://eproc1g.tjal.jus.br/eproc/")
            self.assertEqual(json.loads(alvo.read_text(encoding="utf-8")),
                             {"eproc:TJAL": {"1g": "https://eproc1g.tjal.jus.br/eproc/"}})
            tribunais.definir_endereco("eproc:TJAL", "1g", "")
            self.assertEqual(json.loads(alvo.read_text(encoding="utf-8")), {})


class _nada:
    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False


class TestCNJ(unittest.TestCase):
    def test_formatos_aceitos(self):
        for texto in ("0800072-12.2024.8.02.0056", "08000721220248020056",
                      "Proc. 0800072–12.2024.8.02.0056 (Word)", "0800072 12 2024 8 02 0056"):
            self.assertEqual(cnj.ler(texto).formatado, "0800072-12.2024.8.02.0056", texto)

    def test_digito_verificador(self):
        self.assertTrue(cnj.ler("0800072-12.2024.8.02.0056").digito_confere)
        self.assertFalse(cnj.ler("0800072-13.2024.8.02.0056").digito_confere)

    def test_dependente_so_com_barra(self):
        # Regressão da base: "- 2ª Vara" virava o incidente 02.
        self.assertEqual(cnj.ler("0700123-45.2024.8.02.0001 - 2ª Vara").dependente, "")
        self.assertEqual(cnj.ler("0700123-45.2024.8.02.0001-2024").dependente, "")
        n = cnj.ler("0700123-45.2024.8.02.0001/0003")
        self.assertEqual(n.dependente, "03")
        self.assertEqual(n.nome_arquivo, "0700123-45.2024.8.02.0001-03")

    def test_sem_casar_dentro_de_numero_maior(self):
        with self.assertRaises(cnj.NumeroInvalido):
            cnj.ler("9070012345202480200019")

    def test_varios_na_mesma_linha_e_sem_repetir(self):
        achados = cnj.extrair_todos("a 0700123-45.2024.8.02.0001; b 07001234520248020001 "
                                    "c 5001234-56.2023.4.04.7100")
        self.assertEqual([n.formatado for n in achados],
                         ["0700123-45.2024.8.02.0001", "5001234-56.2023.4.04.7100"])

    def test_ler_none(self):
        with self.assertRaises(cnj.NumeroInvalido):
            cnj.ler(None)  # type: ignore[arg-type]

    def test_nome_de_arquivo_com_dependente(self):
        # Regressão: o "-NN" do nome do arquivo (Numero.nome_arquivo) era
        # ignorado, e o PDF do incidente ficava com a chave do principal.
        principal = "0700123-83.2024.8.02.0001"
        for nome, esperado in ((principal, principal),
                               (f"{principal}-01", f"{principal}-01"),
                               (f"{principal}/0001", f"{principal}-01"),
                               (f"{principal} (2)", principal),
                               (f"{principal}-01 (2)", f"{principal}-01"),
                               (f"{principal} 2025-03-10 14h00 - revisão", principal),
                               (f"{principal} - 2ª Vara", principal)):
            self.assertEqual(cnj.ler_nome_arquivo(nome).nome_arquivo, esperado, nome)
        with self.assertRaises(cnj.NumeroInvalido):
            cnj.ler_nome_arquivo("INDICE")


class TestTribunais(unittest.TestCase):
    def test_detecta_pelo_numero(self):
        t = tribunais.por_numero(cnj.ler("0700123-45.2024.8.02.0001"))
        self.assertEqual((t.sigla, t.sistema), ("TJAL", "esaj"))
        self.assertIsNotNone(t.alternativo)
        self.assertEqual(t.alternativo.sistema, "eproc")
        self.assertEqual(t.portal, "esaj:TJAL")

    def test_secao_judiciaria_federal(self):
        for origem, host in (("7000", "jfpr"), ("7100", "jfrs"), ("7200", "jfsc")):
            n = cnj.ler(f"5001234-56.2023.4.04.{origem}")
            self.assertIn(host, tribunais.por_numero(n).url_para(n))

    def test_nao_suportado(self):
        t = tribunais.por_numero(cnj.ler("0001234-56.2023.8.05.0001"))
        self.assertFalse(t.suportado)

    def test_catalogo_valido(self):
        dados = json.loads(tribunais.ARQUIVO.read_text(encoding="utf-8"))
        chaves = [t["chave"] for t in dados["tribunais"]]
        self.assertEqual(len(chaves), len(set(chaves)))
        for t in tribunais.carregar():
            if t.suportado:
                self.assertTrue(t.url_para(), t.sigla)


class TestConfig(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.arq = Path(self.dir.name) / "config.ini"

    def tearDown(self):
        self.dir.cleanup()

    def test_cria_com_comentarios_e_padroes(self):
        c = config.Config(self.arq)
        self.assertTrue(self.arq.exists())
        self.assertIn("; Pasta com tudo", self.arq.read_text(encoding="utf-8"))
        self.assertTrue(c.flag("download", "pular_baixados"))
        self.assertEqual(c.texto("transcricao", "modelo_ao_vivo"), "small")

    def test_definir_preserva_comentarios_e_aceita_porcento(self):
        c = config.Config(self.arq)
        antes = self.arq.read_text(encoding="utf-8").count(";")
        c.definir("geral", "nome_usuario", "Dr. Fulano 100% ; teste")
        c.definir("download", "pausa_entre_processos", 7)
        self.assertEqual(c.texto("geral", "nome_usuario"), "Dr. Fulano 100% ; teste")
        self.assertEqual(c.inteiro("download", "pausa_entre_processos"), 7)
        self.assertEqual(self.arq.read_text(encoding="utf-8").count(";"), antes + 1)

    def test_bom_e_valor_multilinha(self):
        self.arq.write_text("\ufeff[geral]\nnome_usuario = A\n  continua\n", encoding="utf-8")
        c = config.Config(self.arq)
        c.definir("geral", "nome_usuario", "B")
        self.assertEqual(c.texto("geral", "nome_usuario"), "B")
        self.assertEqual(self.arq.read_text(encoding="utf-8").count("[geral]"), 1)

    def test_recarregar_nao_expoe_configuracao_vazia(self):
        # Regressão: recarregar() trocava o parser por um VAZIO antes de ler;
        # outra thread lia, nesse intervalo, os padrões (pasta do acervo
        # errada) ou valores ainda em montagem (AttributeError).
        c = config.Config(self.arq)
        c.definir("geral", "pasta_acervo", r"D:\Gabinete\Acervo")
        vistos = []
        original = configparser.ConfigParser.read_string

        def lendo(cp, *args, **kwargs):
            vistos.append(c.texto("geral", "pasta_acervo"))
            return original(cp, *args, **kwargs)

        with mock.patch.object(configparser.ConfigParser, "read_string", lendo):
            c.recarregar()
        self.assertEqual(vistos, [r"D:\Gabinete\Acervo"])
        self.assertEqual(c.texto("geral", "pasta_acervo"), r"D:\Gabinete\Acervo")

    def test_arquivo_em_ansi_ou_utf16_nao_volta_ao_padrao(self):
        acervo = Path(self.dir.name) / "Gabinete" / "Acervo"     # absoluto também no Windows
        texto = (config.modelo_ini()
                 .replace("nome_usuario = \n", "nome_usuario = Dra. Conceição\n")
                 .replace("pasta_acervo = Acervo\n", f"pasta_acervo = {acervo}\n"))
        for codificacao in ("cp1252", "utf-16"):
            with self.subTest(codificacao=codificacao):
                self.arq.write_bytes(texto.replace("\n", "\r\n").encode(codificacao))
                with self.assertLogs("nucleo.config", "WARNING") if codificacao == "cp1252" \
                        else _nada():
                    c = config.Config(self.arq)
                self.assertEqual(c.texto("geral", "nome_usuario"), "Dra. Conceição")
                self.assertEqual(c.pasta_acervo, acervo)
                c.definir("unidade", "comarca", "Maceió")      # antes: UnicodeDecodeError
                self.assertIn("Dra. Conceição", self.arq.read_text(encoding="utf-8"))
                self.assertEqual(config.Config(self.arq).texto("unidade", "comarca"), "Maceió")

    def test_so_leitura_nao_cria_o_arquivo(self):
        c = config.Config(self.arq, criar=False)
        self.assertFalse(self.arq.exists())
        self.assertEqual(c.texto("transcricao", "modelo_ao_vivo"), "small")

    def test_pastas_uma_dentro_da_outra(self):
        base = Path(self.dir.name)
        self.assertEqual(config.conflito_de_pastas(base / "Acervo", base / "Sigilosos"), "")
        for acervo, sigilosos in ((base, base / "Sigilosos"), (base / "Acervo", base),
                                  (base / "Acervo", base / "Acervo")):
            self.assertIn("compartilhado com a IA", config.conflito_de_pastas(acervo, sigilosos))

    def test_definir_insiste_quando_o_arquivo_esta_preso(self):
        # Regressão: o antivírus ou o OneDrive seguram o config.ini por um
        # instante, e o os.replace falhava de primeira com "acesso negado".
        c = config.Config(self.arq)
        original = config.os.replace
        falhas = [PermissionError(13, "Acesso negado")] * 2

        def replace(origem, destino):
            if falhas:
                raise falhas.pop()
            return original(origem, destino)

        with mock.patch.object(config.os, "replace", replace), \
                mock.patch.object(config.time, "sleep") as dormir:
            c.definir("unidade", "comarca", "Maceió")
        self.assertEqual(dormir.call_count, 2)
        self.assertEqual(config.Config(self.arq).texto("unidade", "comarca"), "Maceió")

    def test_definir_com_arquivo_preso_de_vez_explica_e_nao_deixa_resto(self):
        c = config.Config(self.arq)
        with mock.patch.object(config.os, "replace", side_effect=PermissionError(13, "Acesso negado")), \
                mock.patch.object(config.time, "sleep") as dormir:
            with self.assertRaises(PermissionError) as ctx:
                c.definir("unidade", "comarca", "Maceió")
        self.assertEqual(dormir.call_count, config.TENTATIVAS_TROCA - 1)
        self.assertGreaterEqual(config.TENTATIVAS_TROCA * config.ESPERA_TROCA_S, 2)
        mensagem = str(ctx.exception)
        self.assertIn("Não consegui salvar a configuração", mensagem)
        self.assertIn("antivírus, OneDrive", mensagem)
        self.assertNotIn("Errno", mensagem)
        self.assertFalse(self.arq.with_name(self.arq.name + ".tmp").exists())
        self.assertEqual(c.texto("unidade", "comarca"), "")

    def test_valor_invalido_vira_padrao(self):
        self.arq.write_text("[download]\npausa_entre_processos = muito\npular_baixados = talvez\n",
                            encoding="utf-8")
        c = config.Config(self.arq)
        self.assertEqual(c.inteiro("download", "pausa_entre_processos"), 3)
        self.assertTrue(c.flag("download", "pular_baixados"))

    def test_padroes_do_helestron(self):
        base = Path(self.dir.name) / "Documentos" / "Helestron"
        with mock.patch.object(caminhos, "BASE_USUARIO", base):
            c = config.Config(self.arq)
            self.assertEqual(c.pasta_acervo, base / "Acervo")
            self.assertEqual(c.pasta_processos, base / "Acervo" / "Processos")
            self.assertEqual(c.pasta_sigilosos, base / "Sigilosos")
            self.assertEqual(c.pasta_pauta, base / "Pauta")
            self.assertEqual(c.conflito_de_pastas(), "")
        self.assertTrue(c.flag("pauta", "monitorar"))
        self.assertEqual(c.inteiro("pauta", "intervalo_horas"), 6)
        self.assertEqual(c.inteiro("pauta", "dias_atras"), 7)
        self.assertEqual(c.inteiro("pauta", "dias_a_frente"), 60)
        self.assertFalse(c.flag("pauta", "incluir_partes_sigilosos"))
        texto = self.arq.read_text(encoding="utf-8")
        self.assertIn("Helestron - configuração", texto)
        self.assertIn("tela de Ajustes", texto)
        self.assertIn("[pauta]", texto)
        self.assertNotIn("Assessor", texto)
        self.assertNotIn("Configurações", texto)

    def test_rotulos_da_pauta_para_a_tela_de_ajustes(self):
        chaves = {(s, c) for s, c, _, _ in config.ESQUEMA}
        for chave in ("monitorar", "intervalo_horas", "dias_atras", "dias_a_frente",
                      "incluir_partes_sigilosos", "pasta"):
            self.assertIn(("pauta", chave), config.ROTULOS)
        for (secao, chave), (rotulo, tipo, ajuda) in config.ROTULOS.items():
            with self.subTest(chave=chave):
                self.assertIn((secao, chave), chaves)
                self.assertIn(tipo, ("texto", "flag", "inteiro", "pasta", "escolha"))
                self.assertTrue(rotulo and ajuda and ajuda.endswith("."))

    def test_pasta_da_pauta_fora_do_acervo(self):
        base = Path(self.dir.name)
        acervo, sigilosos = base / "Acervo", base / "Sigilosos"
        self.assertEqual(config.conflito_de_pastas(acervo, sigilosos, base / "Pauta"), "")
        for pauta in (acervo, acervo / "Pauta"):
            with self.subTest(pauta=pauta):
                frase = config.conflito_de_pastas(acervo, sigilosos, pauta)
                self.assertIn("pauta exportada", frase)
                self.assertIn("compartilhado com a IA", frase)
        # a pauta pode ficar junto dos sigilosos (nenhum dos dois vai para a IA)
        self.assertEqual(config.conflito_de_pastas(acervo, sigilosos, sigilosos / "Pauta"), "")
        c = config.Config(self.arq)
        c.definir("geral", "pasta_acervo", str(acervo))
        c.definir("geral", "pasta_sigilosos", str(sigilosos))
        c.definir("pauta", "pasta", str(acervo / "Pauta"))
        self.assertIn("pauta exportada", c.conflito_de_pastas())

    def test_criar_pastas_inclui_a_pauta(self):
        base = Path(self.dir.name)
        with mock.patch.object(caminhos, "BASE_USUARIO", base / "Docs"), \
                mock.patch.object(caminhos, "LOGS", base / "Local" / "Logs"):
            config.Config(self.arq).criar_pastas()
        for pasta in ("Docs/Acervo/Processos", "Docs/Acervo/Transcricoes", "Docs/Sigilosos",
                      "Docs/Pauta", "Local/Logs"):
            self.assertTrue((base / pasta).is_dir(), pasta)


class TestCofreSenhas(unittest.TestCase):
    def test_guarda_obtem_apaga(self):
        with tempfile.TemporaryDirectory() as d:
            cofre = cofre_senhas.CofreSenhas(Path(d) / "c.json")
            cofre.guardar("esaj:TJAL", "12345678900", "s3nh@;%")
            self.assertEqual(cofre.obter("esaj:TJAL"), ("12345678900", "s3nh@;%"))
            self.assertNotIn("s3nh@", (Path(d) / "c.json").read_text(encoding="utf-8"))
            cofre.apagar("esaj:TJAL")
            self.assertEqual(cofre.obter("esaj:TJAL"), ("", ""))

    def test_arquivo_estragado_nao_derruba(self):
        with tempfile.TemporaryDirectory() as d:
            arq = Path(d) / "c.json"
            arq.write_text("{isto não é json", encoding="utf-8")
            self.assertEqual(cofre_senhas.CofreSenhas(arq).obter("x"), ("", ""))


class TestListas(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.d = Path(self.dir.name)

    def tearDown(self):
        self.dir.cleanup()

    def test_texto_colado_com_senha(self):
        lei = listas.ler_texto("0800072-12.2024.8.02.0056 ; SENHA1\n"
                               "0700123-45.2024.8.02.0001; 0700001-11.2023.8.02.0001\n"
                               "# comentário 0700999-99.2024.8.02.0001\n")
        self.assertEqual(len(lei.processos), 3)
        self.assertEqual(lei.senhas, {"0800072-12.2024.8.02.0056": "SENHA1"})

    def test_segunda_coluna_nao_vira_senha(self):
        lei = listas.ler_texto("Por fim 0800072-12.2024.8.02.0056 ; Audiência\n")
        self.assertEqual(lei.senhas, {})
        # tabela colada do Excel, sem o cabeçalho "Senha"
        lei = listas.ler_texto("0800072-12.2024.8.02.0056\tCível\n"
                               "0700123-45.2024.8.02.0001\tFamília\n")
        self.assertEqual(len(lei.processos), 2)
        self.assertEqual(lei.senhas, {})

    def test_numero_e_senha_em_todas_as_linhas_sem_cabecalho(self):
        # Só linhas "número ; senha": o leitor de CSV pegava o atalho e perdia
        # as senhas.
        lei = listas.ler_texto("0800072-12.2024.8.02.0056 ; SENHA1\n"
                               "0700123-45.2024.8.02.0001 ; SENHA2\n")
        self.assertEqual(lei.senhas, {"0800072-12.2024.8.02.0056": "SENHA1",
                                      "0700123-45.2024.8.02.0001": "SENHA2"})
        (self.d / "s.csv").write_text("0800072-12.2024.8.02.0056;XY99\n"
                                      "0700123-45.2024.8.02.0001;ZW77\n", encoding="utf-8")
        lei = listas.ler_arquivo(self.d / "s.csv")
        self.assertEqual(lei.senhas, {"0800072-12.2024.8.02.0056": "XY99",
                                      "0700123-45.2024.8.02.0001": "ZW77"})

    def test_docx_na_ordem_do_documento(self):
        from docx import Document

        doc = Document()
        doc.add_paragraph("Primeiro: 0700123-45.2024.8.02.0001")
        t = doc.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Processo", "Senha"
        t.cell(1, 0).text, t.cell(1, 1).text = "0800072-12.2024.8.02.0056", "ABC123"
        doc.add_paragraph("Último: 5001234-56.2023.4.04.7100")
        doc.save(self.d / "l.docx")
        lei = listas.ler_arquivo(self.d / "l.docx")
        self.assertEqual([n.formatado for n in lei.processos],
                         ["0700123-45.2024.8.02.0001", "0800072-12.2024.8.02.0056",
                          "5001234-56.2023.4.04.7100"])
        self.assertEqual(lei.senhas, {"0800072-12.2024.8.02.0056": "ABC123"})

    def test_xlsx_coluna_preferida_corrompido_e_aba_oculta(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.append(["Nº do processo", "Processo de origem"])
        ws.append(["0800072-12.2024.8.02.0056", "0700001-11.2023.8.02.0001"])
        ws.append([8000721220248020056, None])     # número: o Excel corrompe
        oculta = wb.create_sheet("oculta")
        oculta.sheet_state = "hidden"
        oculta.append(["0700999-99.2024.8.02.0001"])
        wb.save(self.d / "l.xlsx")
        lei = listas.ler_arquivo(self.d / "l.xlsx")
        self.assertEqual([n.formatado for n in lei.processos], ["0800072-12.2024.8.02.0056"])
        self.assertTrue(lei.corrompidos)
        self.assertTrue(any("oculta" in a for a in lei.avisos))

    def test_xls_que_e_html_em_utf16(self):
        html = ("<html><table><tr><td>0800072-12.2024.8.02.0056</td></tr>"
                "<tr><td>07001234520248020001</td></tr></table></html>")
        (self.d / "x.xls").write_bytes(html.encode("utf-16"))
        lei = listas.ler_arquivo(self.d / "x.xls")
        self.assertEqual(len(lei.processos), 2)

    def test_csv_ponto_e_virgula(self):
        (self.d / "l.csv").write_text("processo;senha\n0800072-12.2024.8.02.0056;XY99\n",
                                      encoding="cp1252")
        lei = listas.ler_arquivo(self.d / "l.csv")
        self.assertEqual(lei.senhas, {"0800072-12.2024.8.02.0056": "XY99"})

    def test_ods(self):
        conteudo = (
            '<?xml version="1.0"?><office:document-content '
            'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
            'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
            'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body>'
            '<office:spreadsheet><table:table table:name="A"><table:table-row>'
            '<table:table-cell office:value-type="string"><text:p>0800072-12.2024.8.02.0056'
            '</text:p></table:table-cell></table:table-row></table:table>'
            '</office:spreadsheet></office:body></office:document-content>')
        with zipfile.ZipFile(self.d / "l.ods", "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            z.writestr("content.xml", conteudo)
        lei = listas.ler_arquivo(self.d / "l.ods")
        self.assertEqual([n.formatado for n in lei.processos], ["0800072-12.2024.8.02.0056"])

    def test_pdf_com_quebra_no_meio_do_numero(self):
        try:
            import pymupdf
        except ImportError:
            import fitz as pymupdf
        doc = pymupdf.open()
        doc.new_page().insert_text((72, 72), "Processo 0800072-12.2024.8.02.\n0056 pauta")
        doc.save(str(self.d / "l.pdf"))
        lei = listas.ler_arquivo(self.d / "l.pdf")
        self.assertEqual([n.formatado for n in lei.processos], ["0800072-12.2024.8.02.0056"])

    def test_arquivo_sem_numero(self):
        (self.d / "v.txt").write_text("nada aqui", encoding="utf-8")
        with self.assertRaises(listas.ListaInvalida):
            listas.ler_arquivo(self.d / "v.txt")

    def test_links_compartilhados(self):
        self.assertEqual(
            listas.url_de_download("https://docs.google.com/spreadsheets/d/AbC-1_x/edit#gid=0"),
            "https://docs.google.com/spreadsheets/d/AbC-1_x/export?format=xlsx")
        self.assertIn("download=1", listas.url_de_download(
            "https://tjal-my.sharepoint.com/:x:/g/personal/x/EaBc?e=xyz"))


class TestSistema(unittest.TestCase):
    def test_nome_seguro_e_destino_livre(self):
        self.assertEqual(sistema.nome_seguro('a/b:c*?"<>|'), "a_b_c_")
        with tempfile.TemporaryDirectory() as d:
            p = Path(d)
            (p / "x.docx").write_text("1")
            self.assertEqual(sistema.destino_livre(p, "x", ".docx").name, "x (2).docx")

    def test_frase_de_reinstalar(self):
        self.assertIn("Helestron-Setup", sistema.REINSTALAR)
        self.assertIn("sem apagar os seus dados", sistema.REINSTALAR)

    def test_ambiente_sem_chaves(self):
        import os

        os.environ["ANTHROPIC_API_KEY"] = "sk-teste"
        try:
            self.assertNotIn("ANTHROPIC_API_KEY", sistema.ambiente_sem_chaves())
        finally:
            del os.environ["ANTHROPIC_API_KEY"]



def _numero(seq: str, tr: str = "02", dependente: str = "") -> cnj.Numero:
    corpo = f"{seq}20248{tr}0001"
    dv = 98 - int(corpo + "00") % 97
    texto = f"{seq}-{dv:02d}.2024.8.{tr}.0001"
    return cnj.ler(f"{texto}/{dependente}" if dependente else texto)


def pauta_com_sigiloso(arquivo: Path, *numeros: cnj.Numero, sigiloso: bool = True) -> None:
    """Grava no banco da pauta uma audiência de cada processo (sigilosa ou não)."""
    from datetime import date

    from helestron.pauta import modelos as pm
    from helestron.pauta.armazem import Armazem

    armazem = Armazem(arquivo)
    try:
        armazem.gravar([pm.nova(sistema="esaj", tribunal="TJAL", data_=date(2026, 10, 8),
                                processo=n.formatado, hora=f"{9 + i:02d}:30",
                                tipo_original="Conciliação", situacao_original="Designada",
                                partes="M. A. S. x J. R. S.", sigiloso=sigiloso)
                        for i, n in enumerate(numeros)], "esaj-tjal", None)
    finally:
        armazem.fechar()


class TestSigilo(unittest.TestCase):
    """A regra única do processo sigiloso: autos, transcrição, gravação ou
    diário na pasta dos sigilosos, ou a pauta marcando o segredo de justiça.
    Antes, o compartilhamento e o download só olhavam os PDFs da pasta: o
    processo que a pauta (ou a transcrição sigilosa) dava como sigiloso ia
    para a IA, para a nuvem e para o acervo."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        self.tmp = Path(self._tmp.name)
        self.sigilosos = self.tmp / "Sigilosos"
        self.pauta = self.tmp / "local" / "pauta.sqlite3"
        p = mock.patch.object(caminhos, "ARQUIVO_PAUTA", self.pauta)
        p.start()
        self.addCleanup(p.stop)
        self.addCleanup(sigilo.esquecer_pauta)
        self.cfg = mock.Mock(pasta_sigilosos=self.sigilosos)

    def _arquivo(self, rel: str) -> Path:
        caminho = self.sigilosos / rel
        caminho.parent.mkdir(parents=True, exist_ok=True)
        caminho.write_bytes(b"x")
        return caminho

    def test_pasta_conta_autos_transcricao_gravacao_e_diario(self):
        autos, lote, docx, diario, gravacao = (_numero(s) for s in (
            "0700101", "0700102", "0700103", "0700104", "0700105"))
        incidente = _numero("0700106", dependente="01")
        fundo = _numero("0700107")
        self._arquivo(f"{autos.nome_arquivo}.pdf")
        self._arquivo(f"Lote 1/{lote.nome_arquivo}.pdf")
        self._arquivo(f"Transcricoes/{docx.nome_arquivo} (2).docx")
        self._arquivo(f"Transcricoes/_audio/{diario.nome_arquivo} 2026-09-16 14h00.jsonl")
        self._arquivo(f"Transcricoes/_audio/{gravacao.nome_arquivo} 2026-09-16 15h00.flac")
        self._arquivo(f"Lote 1/{incidente.nome_arquivo}.pdf")
        # fundo demais não conta: a pasta pode ter sido apontada para os Documentos (o
        # preparo leva até subpastas do lote: test_sigilo.TestPastaFunda)
        self._arquivo(f"a/b/c/d/e/{fundo.nome_arquivo}.pdf")
        esperado = {n.nome_arquivo for n in (autos, lote, docx, diario, gravacao, incidente)}
        self.assertEqual(sigilo.chaves_na_pasta(self.sigilosos), esperado)
        for n in (autos, lote, docx, diario, gravacao, incidente):
            self.assertTrue(sigilo.na_pasta(self.sigilosos, n), n)
            self.assertEqual(sigilo.motivo(self.cfg, n), sigilo.MOTIVO_PASTA)
        # o incidente não torna sigiloso o principal (o contrário, sim: o
        # incidente herda o sigilo do principal)
        self.assertFalse(sigilo.na_pasta(self.sigilosos, _numero("0700106")))
        self.assertTrue(sigilo.na_pasta(self.sigilosos, _numero("0700101", dependente="02")))
        self.assertEqual(sigilo.motivo(self.cfg, _numero("0700101", dependente="02")),
                         sigilo.MOTIVO_PASTA_PRINCIPAL)
        self.assertFalse(sigilo.na_pasta(self.sigilosos, fundo))
        self.assertFalse(sigilo.processo_sigiloso(self.cfg, fundo))
        # acervo (por engano) dentro da pasta de sigilosos: o que é dele não vira sigiloso
        self.assertEqual(sigilo.chaves_na_pasta(self.sigilosos, self.sigilosos / "Lote 1"),
                         esperado - {lote.nome_arquivo, incidente.nome_arquivo})

    def test_pauta_marca_o_processo_e_a_consulta_nao_cria_o_banco(self):
        x, y, publico = _numero("0700777"), _numero("0700778"), _numero("0700779")
        incidente = _numero("0700777", dependente="01")
        self.assertEqual(sigilo.chaves_da_pauta(), set())
        self.assertFalse(self.pauta.exists(), "a consulta criou o banco da pauta")
        self.pauta.parent.mkdir(parents=True)
        pauta_com_sigiloso(self.pauta, x, incidente)
        pauta_com_sigiloso(self.pauta, publico, sigiloso=False)
        self.assertEqual(sigilo.chaves_da_pauta(), {x.nome_arquivo, incidente.nome_arquivo})
        self.assertEqual(sigilo.motivo(self.cfg, x), sigilo.MOTIVO_PAUTA)
        self.assertTrue(sigilo.na_pauta(x.formatado))
        self.assertFalse(sigilo.processo_sigiloso(self.cfg, publico))
        self.assertFalse(sigilo.na_pauta(y, pauta=None))
        self.assertEqual(sigilo.chaves_sigilosas(self.sigilosos, pauta=None), set())
        # O segredo decretado depois: a pauta passa a marcar o processo, e a
        # regra vale na hora (o que se guardou da leitura anterior não serve).
        pauta_com_sigiloso(self.pauta, y)
        self.assertIn(y.nome_arquivo, sigilo.chaves_sigilosas(self.sigilosos))
        self.assertTrue(sigilo.processo_sigiloso(self.cfg, y))

    def test_pauta_ilegivel_vale_o_que_ja_se_sabia(self):
        x = _numero("0700777")
        self.pauta.parent.mkdir(parents=True)
        pauta_com_sigiloso(self.pauta, x)
        self.assertTrue(sigilo.na_pauta(x))
        pauta_com_sigiloso(self.pauta, _numero("0700778"))      # o banco mudou
        with mock.patch.object(sigilo, "_ler_banco", side_effect=OSError("ocupado")), \
                self.assertLogs("nucleo.sigilo", "WARNING"):
            self.assertTrue(sigilo.na_pauta(x))
        sigilo.esquecer_pauta()
        # o que a pauta já apurou fica também no registro ao lado do banco
        with mock.patch.object(sigilo, "_ler_banco", side_effect=OSError("ocupado")), \
                self.assertLogs("nucleo.sigilo", "WARNING"):
            self.assertTrue(sigilo.na_pauta(x))
        sigilo.arquivo_apurado().unlink()
        sigilo.esquecer_pauta()
        with mock.patch.object(sigilo, "_ler_banco", side_effect=OSError("ocupado")), \
                self.assertLogs("nucleo.sigilo", "WARNING"):
            self.assertFalse(sigilo.na_pauta(x))                 # nunca levanta


if __name__ == "__main__":
    unittest.main()
