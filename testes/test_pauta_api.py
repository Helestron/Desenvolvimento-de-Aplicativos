"""Pauta pela API com o ServicoPauta DE VERDADE (o banco, a planilha, a senha da sessão).

O resto da API da pauta é testado com um serviço simulado em
test_servidor_pauta.py; aqui ficam os casos em que só o serviço real mostra o
defeito: a escolha "sem as partes dos sigilosos" feita na janela de
exportação, e a senha digitada com "Lembrar neste computador" desligado.
"""

from __future__ import annotations

import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from helestron.pauta import modelos, navegacao
from helestron.pauta.navegacao import Leitura
from helestron.pauta.tabelas import Reconhecimento

from testes import apoio_download as apoio
from testes import apoio_pauta as ap
from testes.test_servidor_base import ServidorDeTeste

DIA = date(2026, 10, 6)
PARTES = "Maria A. S. x João R. S."


class TestExportarPelaApi(ServidorDeTeste):
    def setUp(self):
        super().setUp()
        self.servico = self.app.pauta()
        self.servico.armazem.gravar([modelos.nova(
            sistema="esaj", tribunal="TJAL", data_=DIA, processo=ap.numero("0700105"),
            hora="09:30", tipo_original="Conciliação", partes=PARTES, sigiloso=True)],
            "esaj-tjal", None, registrar_novas=False)

    def partes_na_planilha(self, corpo: dict) -> str:
        from openpyxl import load_workbook

        dados = self.cliente.dados("POST", "/api/pauta/exportar",
                                   {"de": DIA.isoformat(), "ate": DIA.isoformat(), **corpo})
        return load_workbook(Path(dados["arquivo"]))["Pauta"].cell(5, 6).value

    def test_a_escolha_da_janela_vale_mesmo_com_o_ajuste_ligado(self):
        self.cfg.definir("pauta", "incluir_partes_sigilosos", True)
        self.assertEqual(self.partes_na_planilha({"incluir_partes_sigilosos": False}),
                         modelos.MASCARA_SIGILO, "desligado na janela: mascarado")
        self.assertEqual(self.partes_na_planilha({"incluir_partes_sigilosos": True}), PARTES)
        self.assertEqual(self.partes_na_planilha({}), PARTES, "sem o campo, vale o Ajuste")
        self.cfg.definir("pauta", "incluir_partes_sigilosos", False)
        self.assertEqual(self.partes_na_planilha({}), modelos.MASCARA_SIGILO)
        self.assertEqual(self.partes_na_planilha({"incluir_partes_sigilosos": "false"}),
                         modelos.MASCARA_SIGILO)


class _Portal:
    instancias: list = []

    def __init__(self, nav, tribunal, opcoes, ctx, credenciais):
        self.opcoes, self.credenciais = opcoes, credenciais
        self.sistema, self.base = tribunal.sistema, "https://portal.invalid"
        self.nome = f"{tribunal.nome_sistema} do {tribunal.sigla}"
        _Portal.instancias.append(self)

    def entrar(self):
        pass


class _Extrator:
    def __init__(self, *a, **k):
        pass

    def extrair(self, de, ate):
        return Leitura(reconhecimento=Reconhecimento(tabelas=1), paginas=1,
                       periodo_aplicado=True, url="https://portal.invalid/pauta")


class TestSenhaDaSessao(ServidorDeTeste):
    def setUp(self):
        super().setUp()
        _Portal.instancias = []
        self.servico = self.app.pauta()
        self.servico.fabrica_navegador = apoio.NavegadorFalso
        self.servico.fabrica_portal = _Portal
        p = mock.patch.object(navegacao, "ExtratorAutomatico", _Extrator)
        p.start()
        self.addCleanup(p.stop)

    def test_senha_so_por_agora_vale_para_sincronizar(self):
        acesso = self.cliente.dados("POST", "/api/acessos", {
            "portal": "esaj:TJAL", "usuario": "12345678900", "senha": "segredo",
            "lembrar": False})
        self.assertTrue(acesso.get("so_agora"))
        self.cliente.dados("POST", "/api/pauta/fontes", {"tribunal": "TJAL", "sistema": "esaj"})
        tarefa = self.cliente.dados("POST", "/api/pauta/sincronizar", {})["tarefa"]
        fim = self.esperar_tarefa(tarefa)
        self.assertNotEqual(fim["estado"], "falhou", fim)
        portal = _Portal.instancias[-1]
        self.assertEqual(portal.credenciais, ("12345678900", "segredo"))
        self.assertEqual(portal.opcoes.modo_login("esaj"), "senha",
                         "a tela de entrada do portal não abre: a senha foi informada")


if __name__ == "__main__":
    unittest.main()
