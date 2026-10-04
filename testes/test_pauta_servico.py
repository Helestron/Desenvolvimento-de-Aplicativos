"""Pauta: a fachada ServicoPauta (seção 8.10) - consulta, fontes, monitoramento e
sincronização (com o navegador e a leitura do portal substituídos por dublês)."""

from __future__ import annotations

import unittest
from datetime import date, datetime, timedelta
from unittest import mock

from helestron.download import modelos as mdl
from helestron.pauta import modelos, navegacao
from helestron.pauta.navegacao import Leitura, PautaNaoEncontrada
from helestron.pauta.servico import ErroPauta, ServicoPauta
from helestron.pauta.tabelas import Reconhecimento

from testes import apoio_download as apoio
from testes import apoio_pauta as ap

AGORA = datetime(2026, 10, 5, 11, 0)


def aud(seq, data_, hora, tipo="Conciliação", situacao="Designada", sistema="esaj",
        tribunal="TJAL", **extra):
    tr = "21" if tribunal == "TJRS" else "02"
    return modelos.nova(sistema=sistema, tribunal=tribunal, data_=data_,
                        processo=ap.numero(seq, tr), hora=hora, tipo_original=tipo,
                        situacao_original=situacao, **extra)


class PortalFalso:
    instancias: list = []

    def __init__(self, nav, tribunal, opcoes, ctx, credenciais, falhas=None):
        self.nav, self.tribunal, self.opcoes = nav, tribunal, opcoes
        self.ctx, self.credenciais = ctx, credenciais
        self.sistema = tribunal.sistema
        self.base = "https://portal.invalid"
        self.nome = f"{tribunal.nome_sistema} do {tribunal.sigla}"
        self.entradas = 0
        self.falhas = list(falhas or [])
        PortalFalso.instancias.append(self)

    def entrar(self):
        self.entradas += 1
        if self.falhas:
            erro = self.falhas.pop(0)
            if erro is not None:
                raise erro


class ExtratorFalso:
    """Faz as vezes do ExtratorAutomatico: segue um roteiro por fonte."""

    roteiro: dict[str, list] = {}
    chamadas: list = []

    def __init__(self, nav, portal, regras, reconhecedor, ctx=None, fonte=None, espera_s=0):
        self.fonte = dict(fonte or {})
        self.reconhecedor = reconhecedor

    def extrair(self, de, ate):
        ExtratorFalso.chamadas.append((self.fonte["id"], de, ate, self.fonte.get("url")))
        passo = ExtratorFalso.roteiro[self.fonte["id"]].pop(0)
        if isinstance(passo, BaseException):
            raise passo
        audiencias, aplicado, url = passo
        leitura = Leitura(reconhecimento=Reconhecimento(audiencias=list(audiencias), tabelas=1),
                          paginas=1, periodo_aplicado=aplicado, url=url, menu="Pauta de Audiências")
        return leitura


class Base(apoio.PastaTemporaria):
    def setUp(self):
        super().setUp()
        self.amb = ap.config_temporaria(self.tmp)
        self.eventos = []
        self.cofre = apoio.CofreFalso({"esaj:TJAL": ("u", "s"), "eproc:TJRS": ("u", "s")})
        self.falhas_portal: dict[str, list] = {}
        PortalFalso.instancias = []
        apoio.NavegadorFalso.instancias = []
        ExtratorFalso.roteiro, ExtratorFalso.chamadas = {}, []

        def fabrica_portal(nav, tribunal, opcoes, ctx, credenciais):
            return PortalFalso(nav, tribunal, opcoes, ctx, credenciais,
                               self.falhas_portal.get(tribunal.sigla))

        self.agora = AGORA
        self.servico = ServicoPauta(self.amb.cfg, self.tmp / "local" / "pauta.sqlite3",
                                    eventos=lambda t, d: self.eventos.append((t, d)),
                                    fabrica_navegador=apoio.NavegadorFalso,
                                    fabrica_portal=fabrica_portal, cofre=self.cofre,
                                    relogio=lambda: self.agora)
        self.addCleanup(self.servico.fechar)
        p = mock.patch.object(navegacao, "ExtratorAutomatico", ExtratorFalso)
        p.start()
        self.addCleanup(p.stop)


class TestConsulta(Base):
    def setUp(self):
        super().setUp()
        d = AGORA.date()
        self.servico.armazem.gravar([
            aud("0700101", d, "09:00", partes="Maria José x Banco Bradesco"),     # já passou
            aud("0700102", d, "14:00", tipo="Una", partes="João x Município"),    # a próxima
            aud("0700103", d, "15:00", situacao="Cancelada"),
            aud("0700104", d + timedelta(days=2), "10:00", sistema="eproc", tribunal="TJRS",
                tipo="Custódia", local="Central de Custódia"),
            aud("0700105", d + timedelta(days=20), "10:00", situacao="Redesignada"),
        ], "x", None, registrar_novas=False)

    def test_listar_e_resumo(self):
        d = AGORA.date()
        dados = self.servico.listar(d, d + timedelta(days=30))
        self.assertEqual(len(dados["audiencias"]), 5)
        self.assertEqual(dados["resumo"], {
            "total": 5, "hoje": 2, "semana": 3,
            "por_situacao": {"Designada": 3, "Cancelada": 1, "Redesignada": 1},
            "por_tipo": {"Conciliação": 3, "Una": 1, "Custódia": 1}})
        filtrar = lambda **f: [a["processo"][:7] for a in  # noqa: E731
                               self.servico.listar(d, d + timedelta(days=30), **f)["audiencias"]]
        self.assertEqual(filtrar(sistema="eproc"), ["0700104"])
        self.assertEqual(filtrar(sistema="Todos"), filtrar())
        self.assertEqual(filtrar(situacao="cancelada"), ["0700103"])
        self.assertEqual(filtrar(situacao="Redesignada"), ["0700105"])
        self.assertEqual(filtrar(busca="maria jose"), ["0700101"])
        self.assertEqual(filtrar(busca="BRADESCO"), ["0700101"])
        self.assertEqual(filtrar(busca="0700104"), ["0700104"])
        self.assertEqual(filtrar(busca="custodia central"), ["0700104"])
        self.assertEqual(filtrar(busca="nada disso"), [])

    def test_resumo_inicio(self):
        r = self.servico.resumo_inicio()
        self.assertEqual((r["hoje"], r["semana"]), (2, 3))
        self.assertEqual(r["proxima"]["processo"][:7], "0700102", "a das 9h já passou; a das 15h "
                                                                  "foi cancelada")
        self.assertIsNone(r["ultima_sincronizacao"])
        self.assertEqual(r["alteracoes_nao_vistas"], 0)
        self.agora = AGORA.replace(hour=23)
        self.assertEqual(self.servico.resumo_inicio()["proxima"]["processo"][:7], "0700104")

    def test_sigilo_pela_pasta(self):
        self.amb.sigilosos.mkdir(parents=True)
        (self.amb.sigilosos / f"{ap.numero('0700102')}.pdf").write_bytes(apoio.pdf_bytes(1))
        d = AGORA.date()
        sig = {a["processo"][:7]: a["sigiloso"] for a in self.servico.listar(d, d)["audiencias"]}
        self.assertEqual(sig, {"0700101": False, "0700102": True, "0700103": False})
        self.assertTrue(self.servico.resumo_inicio()["proxima"]["sigiloso"])


class TestFontesEMonitoramento(Base):
    def test_fontes(self):
        f = self.servico.salvar_fonte("tjal", "esaj", "")
        self.assertEqual((f["id"], f["rotulo"], f["modo"]), ("esaj-tjal", "TJAL · e-SAJ",
                                                              "automatico"))
        for args, trecho in ((("TJAL", "pje", ""), "e-SAJ ou eProc"), (("", "esaj", ""), "tribunal"),
                             (("TJAL", "esaj", "", "javascript:alert(1)"), "https://")):
            with self.subTest(args=args):
                with self.assertRaises(ValueError) as erro:
                    self.servico.salvar_fonte(*args)
                self.assertIn(trecho, str(erro.exception))
        self.servico.remover_fonte("esaj-tjal")
        self.assertEqual(self.servico.fontes(), [])

    def test_monitoramento_com_padroes_e_gravacao(self):
        m = self.servico.monitoramento()
        self.assertEqual((m["ativo"], m["intervalo_horas"], m["dias_atras"], m["dias_a_frente"],
                          m["proxima"], m["fontes_monitoradas"]), (True, 6, 7, 60, None, 0))
        self.servico.salvar_fonte("TJAL", "esaj", "", "https://portal/pauta")
        self.assertEqual(self.servico.monitoramento()["proxima"], AGORA.isoformat(),
                         "nunca sincronizou: a próxima é já")
        self.servico.armazem.definir_meta("ultima_sincronizacao", AGORA - timedelta(hours=2))
        self.assertEqual(self.servico.monitoramento()["proxima"],
                         (AGORA + timedelta(hours=4)).isoformat())
        m = self.servico.configurar_monitoramento(False, 12)
        self.assertEqual((m["ativo"], m["intervalo_horas"], m["proxima"]), (False, 12, None))
        texto = self.amb.cfg.arquivo.read_text(encoding="utf-8")
        self.assertIn("monitorar = false", texto)
        self.assertIn("intervalo_horas = 12", texto)
        for valor in (0, 73, "x"):
            with self.subTest(valor=valor):
                with self.assertRaises(ValueError):
                    self.servico.configurar_monitoramento(True, valor)

    def test_fonte_que_exige_a_pessoa_nao_e_monitorada(self):
        """Certificado, entrada manual ou senha não guardada: o login só acontece
        com a pessoa à frente, e o monitor não sincroniza a fonte. Antes ela
        aparecia como “Monitorada” só por ter a rota da pauta."""
        from helestron.pauta import servico as mod

        self.servico.salvar_fonte("TJAL", "esaj", "", "https://portal/pauta")
        self.servico.salvar_fonte("TJRS", "eproc", "")            # sem rota ainda
        self.servico.credenciais_sessao = {"esaj:TJAL": ("u", "digitada-agora")}

        def fonte(id_):
            return {f["id"]: f for f in self.servico.fontes()}[id_]

        casos = (("senha", {"esaj:TJAL": ("u", "s")}, ""),
                 # a senha digitada "só por agora" não serve ao monitor
                 ("senha", {}, mod.PRESENCA_SEM_SENHA),
                 ("certificado", {"esaj:TJAL": ("u", "s")}, mod.PRESENCA_CERTIFICADO),
                 ("manual", {"esaj:TJAL": ("u", "s")}, mod.PRESENCA_MANUAL))
        for modo, cofre, motivo in casos:
            with self.subTest(modo=modo, cofre=cofre):
                self.amb.cfg.definir("esaj", "login", modo)
                self.cofre.dados = dict(cofre)
                f = fonte("esaj-tjal")
                self.assertEqual((f["exige_presenca"], f["motivo_presenca"], f["monitorada"]),
                                 (bool(motivo), motivo, not motivo))
                self.assertEqual(self.servico.exige_presenca(f), bool(motivo))
                self.assertEqual(self.servico.motivo_presenca("esaj-tjal"), motivo)
                m = self.servico.monitoramento()
                self.assertEqual(m["fontes_monitoradas"], 0 if motivo else 1)
                # só fontes que exigem a pessoa: o monitor não tem próxima vez
                self.assertEqual(m["proxima"], None if motivo else AGORA.isoformat())
        # a fonte sem rota continua "sem rota" (entra sozinha, mas não sabe onde)
        self.cofre.dados["eproc:TJRS"] = ("u", "s")
        tjrs = fonte("eproc-tjrs")
        self.assertEqual((tjrs["exige_presenca"], tjrs["monitorada"]), (False, False))
        # na dúvida (tribunal que não usa o sistema), não exige: a sincronização decide
        self.servico.armazem.salvar_fonte("TJRS", "esaj", "TJRS · e-SAJ")
        self.assertFalse(self.servico.exige_presenca("esaj-tjrs"))

    def test_monitor_marca_o_motivo_presenca(self):
        # O monitor distingue "o portal pediu o código" de "o login exige você".
        self.servico.salvar_fonte("TJAL", "esaj", "", "https://portal.invalid/pauta")
        self.amb.cfg.definir("esaj", "login", "certificado")
        ctx = ap.ContextoDeFundo()
        ctx.motivo = None
        with self.assertRaises(ErroPauta):
            self.servico.sincronizar(ctx, None, date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual((ctx.pediu_login, ctx.motivo), (True, "presenca"))
        ctx.motivo = "codigo"            # o código pedido antes, por outra fonte, prevalece
        with self.assertRaises(ErroPauta):
            self.servico.sincronizar(ctx, None, date(2026, 10, 1), date(2026, 10, 31))
        self.assertEqual(ctx.motivo, "codigo")

    def test_configuracao_ausente_ou_estragada(self):
        class CfgVelho:
            """Um config de antes da pauta: sem a seção, ou com lixo."""
            pasta_sigilosos = pasta_acervo = self.tmp

            def __init__(self, valores):
                self.valores = valores

            def texto(self, secao, chave):
                return self.valores.get(chave, "")

        for valores, esperado in (({}, (True, 6, 7, 60)),
                                  ({"monitorar": "talvez", "intervalo_horas": "seis",
                                    "dias_atras": "-3", "dias_a_frente": "9999"},
                                   (True, 6, 0, 3 * 366)),
                                  ({"monitorar": "não", "intervalo_horas": "200"},
                                   (False, 72, 7, 60))):
            with self.subTest(valores=valores):
                s = ServicoPauta(CfgVelho(valores), self.tmp / "outro.sqlite3")
                self.addCleanup(s.fechar)
                m = s.monitoramento()
                self.assertEqual((m["ativo"], m["intervalo_horas"], m["dias_atras"],
                                  m["dias_a_frente"]), esperado)


class TestSincronizar(Base):
    PERIODO = (date(2026, 10, 1), date(2026, 10, 31))

    def test_grava_lembra_a_rota_e_avisa_as_mudancas(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        d = date(2026, 10, 6)
        url = "https://portal.invalid/pauta"
        ExtratorFalso.roteiro["esaj-tjal"] = [
            ([aud("0700101", d, "09:00"), aud("0700102", d, "10:00"), aud("0700103", d, "11:00")],
             True, url),
            ([aud("0700101", d, "14:00"), aud("0700102", d, "10:00", situacao="Cancelada"),
              aud("0700104", d, "16:00")], True, url),
        ]
        ctx = apoio.ContextoGravador()
        r = self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertEqual((r["novas"], r["atualizadas"], r["removidas"], r["alteracoes"]),
                         (3, 0, 0, 0), "a primeira é a base: nada de 'alterações'")
        self.assertEqual(ctx.status_[-1], "3 audiências conferidas · 3 novas.")
        fonte = self.servico.fontes()[0]
        self.assertEqual((fonte["url"], fonte["menu"], fonte["ultima_sincronizacao"]),
                         (url, "Pauta de Audiências", AGORA.isoformat()))
        self.assertEqual(self.servico.ultima_sincronizacao(), AGORA)
        self.assertEqual([t for t, _ in self.eventos], ["pauta"])
        self.assertEqual(self.eventos[0][1]["tipo"], "atualizada")

        self.agora = AGORA + timedelta(hours=6)
        self.eventos.clear()
        r = self.servico.sincronizar(ctx, ["esaj-tjal"], *self.PERIODO)
        # 0700101 remarcada (alterada), 0700102 cancelada, 0700103 saiu, 0700104 nova
        self.assertEqual((r["novas"], r["atualizadas"], r["canceladas"], r["removidas"]),
                         (1, 2, 1, 1))
        self.assertEqual(ExtratorFalso.chamadas[-1][3], url, "a rota lembrada vai ao extrator")
        tipos = sorted(a["tipo"] for a in self.servico.alteracoes())
        self.assertEqual(tipos, ["alterada", "cancelada", "nova", "removida"])
        self.assertEqual([e[1]["tipo"] for e in self.eventos], ["atualizada", "alteracoes"])
        self.assertEqual(self.eventos[1][1]["dados"]["nao_vistas"], 4)
        self.assertEqual(self.servico.resumo_inicio()["alteracoes_nao_vistas"], 4)
        self.servico.marcar_vistas()
        self.assertEqual(self.servico.resumo_inicio()["alteracoes_nao_vistas"], 0)
        self.assertEqual(ctx.status_[-1], "3 audiências conferidas · 1 nova, 2 alteradas, "
                                          "1 saiu da pauta.")

    def test_periodo_nao_aplicado_so_remove_entre_as_datas_que_vieram(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        ExtratorFalso.roteiro["esaj-tjal"] = [
            ([aud("0700101", date(2026, 10, 5), "09:00"), aud("0700102", date(2026, 10, 20), "09:00"),
              aud("0700103", date(2026, 10, 25), "09:00")], True, ""),
            # o portal mostrou só os dias 5 a 10 (sem campo de período)
            ([aud("0700101", date(2026, 10, 5), "09:00"), aud("0700105", date(2026, 10, 10), "09:00")],
             False, "")]
        ctx = apoio.ContextoGravador()
        self.servico.sincronizar(ctx, None, *self.PERIODO)
        r = self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertEqual(r["removidas"], 0, "o dia 20 e o dia 25 não foram conferidos")
        self.assertEqual(len(self.servico.listar(*self.PERIODO)["audiencias"]), 4)

    def test_uma_fonte_falha_e_a_outra_segue(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        self.servico.salvar_fonte("TJRS", "eproc", "")
        self.falhas_portal["TJAL"] = [mdl.LoginFalhou("o e-SAJ do TJAL recusou a senha")]
        ExtratorFalso.roteiro["eproc-tjrs"] = [
            ([aud("5000001", date(2026, 10, 7), "08:30", sistema="eproc", tribunal="TJRS")],
             True, "https://eproc.invalid/x")]
        ctx = apoio.ContextoGravador()
        r = self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertEqual(r["novas"], 1)
        self.assertEqual(r["erros"], [{"fonte": "esaj-tjal", "rotulo": "TJAL · e-SAJ",
                                       "mensagem": "O e-SAJ do TJAL recusou a senha"}])
        self.assertEqual(ctx.avisos, [("Pauta: TJAL · e-SAJ", "O e-SAJ do TJAL recusou a senha")])
        fontes = {f["id"]: f for f in self.servico.fontes()}
        self.assertEqual(fontes["esaj-tjal"]["ultimo_erro"], "O e-SAJ do TJAL recusou a senha")
        self.assertEqual(fontes["eproc-tjrs"]["ultimo_erro"], "")
        self.assertIn("1 fonte com problema", ctx.status_[-1])
        self.assertTrue(all(n.fechado for n in apoio.NavegadorFalso.instancias),
                        "o navegador fecha mesmo quando o login falha")

    def test_todas_falham(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        ExtratorFalso.roteiro["esaj-tjal"] = [PautaNaoEncontrada(
            "não encontrei a pauta de audiências no e-SAJ do TJAL")]
        ctx = apoio.ContextoGravador()
        with self.assertRaises(ErroPauta) as erro:
            self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertEqual(str(erro.exception), "Não consegui ler a pauta: TJAL · e-SAJ — Não "
                                              "encontrei a pauta de audiências no e-SAJ do TJAL.")
        # o erro da tarefa já diz qual fonte e por quê: sem o mesmo texto num aviso à parte
        self.assertEqual(ctx.avisos, [])
        self.assertIsNone(self.servico.ultima_sincronizacao())
        self.assertEqual(self.eventos, [])

    def test_sem_fontes_ou_fontes_inexistentes(self):
        with self.assertRaises(ErroPauta) as erro:
            self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        self.assertIn("Nenhuma fonte", str(erro.exception))
        self.servico.salvar_fonte("TJAL", "esaj", "")
        with self.assertRaises(ErroPauta) as erro:
            self.servico.sincronizar(apoio.ContextoGravador(), ["nao-existe"], *self.PERIODO)
        self.assertIn("não existem mais", str(erro.exception))

    def test_tribunal_que_nao_usa_o_sistema(self):
        self.servico.armazem.salvar_fonte("TJRS", "esaj", "TJRS · e-SAJ")
        with self.assertRaises(ErroPauta) as erro:
            self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        self.assertIn("O TJRS não usa o e-SAJ", str(erro.exception))

    def test_monitor_sem_senha_nao_abre_navegador(self):
        self.servico.salvar_fonte("TJAL", "esaj", "", "https://portal.invalid/pauta")
        self.cofre.dados.clear()
        ctx = ap.ContextoDeFundo()
        with self.assertRaises(ErroPauta) as erro:
            self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertIn("Entre no portal para continuar o monitoramento", str(erro.exception))
        self.assertTrue(ctx.pediu_login)
        self.assertEqual(apoio.NavegadorFalso.instancias, [], "nenhuma janela aparece do nada")

    def test_usuario_sem_senha_entra_a_mao(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        self.cofre.dados.clear()
        ExtratorFalso.roteiro["esaj-tjal"] = [([], True, "")]
        self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        portal = PortalFalso.instancias[0]
        self.assertIsNone(portal.credenciais)
        self.assertEqual(portal.opcoes.modo_login("esaj"), "manual",
                         "como no download: o navegador abre na tela de entrada")
        self.assertFalse(portal.opcoes.atualizar_ia)

    def test_sessao_caiu_entra_de_novo(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        ExtratorFalso.roteiro["esaj-tjal"] = [mdl.SessaoPerdida("a sessão caiu"),
                                              ([aud("0700101", date(2026, 10, 6), "09:00")], True,
                                               "")]
        ctx = apoio.ContextoGravador()
        r = self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertEqual(r["novas"], 1)
        self.assertEqual(PortalFalso.instancias[0].entradas, 2)
        self.assertIn("A sessão do e-SAJ do TJAL caiu; entrando de novo…", ctx.status_)

    def test_eproc_com_endereco_por_secao(self):
        """TRF4: o login espera um processo para saber a seção; a rota lembrada diz qual é."""

        class PortalPorSecao(PortalFalso):
            def __init__(self, *a, **k):
                super().__init__(*a, **k)
                self.base = ""
                self._candidatos_atuais = []

            def entrar(self):
                super().entrar()
                if self._candidatos_atuais:
                    self.base = self._candidatos_atuais[0]

        self.servico.fabrica_portal = PortalPorSecao
        self.cofre.dados["eproc:TRF4"] = ("u", "s")
        self.servico.salvar_fonte("TRF4", "eproc", "")
        with self.assertRaises(ErroPauta) as erro:
            self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        self.assertIn("seção judiciária", str(erro.exception))
        self.servico.salvar_fonte("TRF4", "eproc", "", "https://eproc.jfrs.jus.br/eprocV2/"
                                                       "controlador.php?acao=audiencia_listar&hash=x")
        ExtratorFalso.roteiro["eproc-trf4"] = [([], True, "")]
        self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        portal = PortalFalso.instancias[-1]
        self.assertEqual(portal.base, "https://eproc.jfrs.jus.br/eprocV2/")
        self.assertEqual(portal.entradas, 2)

    def test_parar(self):
        self.servico.salvar_fonte("TJAL", "esaj", "")
        ctx = apoio.ContextoGravador()
        ctx.cancelar()
        with self.assertRaises(mdl.Cancelado):
            self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.servico.salvar_fonte("TJRS", "eproc", "")
        ctx = apoio.ContextoGravador()
        ExtratorFalso.roteiro["eproc-tjrs"] = [mdl.Cancelado()]
        ExtratorFalso.roteiro["esaj-tjal"] = [([], True, "")]
        with self.assertRaises(mdl.Cancelado):
            self.servico.sincronizar(ctx, ["eproc-tjrs", "esaj-tjal"], *self.PERIODO)


class TestSemanaDeSeteDias(Base):
    def test_proximos_7_dias_sao_os_da_visao_semana(self):
        """'Nos próximos 7 dias' = hoje e os 6 seguintes (a visão Semana da tela)."""
        d = AGORA.date()
        self.servico.armazem.gravar([aud(f"07009{i:02d}", d + timedelta(days=i), "08:00")
                                     for i in range(0, 9)], "x", None, registrar_novas=False)
        self.agora = AGORA.replace(hour=7)
        r = self.servico.resumo_inicio()
        self.assertEqual((r["hoje"], r["semana"]), (1, 7))
        self.assertEqual(self.servico.listar(d, d + timedelta(days=30))["resumo"]["semana"], 7)


class TestPrimeirosPassos(Base):
    def test_fontes_e_configurada(self):
        r = self.servico.resumo_inicio()
        self.assertEqual((r["fontes"], r["configurada"]), (0, False),
                         "abrir a tela da pauta não a configura")
        self.servico.salvar_fonte("TJAL", "esaj", "")
        r = self.servico.resumo_inicio()
        self.assertEqual((r["fontes"], r["configurada"]), (1, True))
        self.servico.remover_fonte("esaj-tjal")
        self.assertFalse(self.servico.resumo_inicio()["configurada"])
        (self.tmp / "rel.csv").write_text(
            f"Data;Processo;Tipo\n05/10/2026;{ap.numero('0700101')};Una\n", encoding="utf-8")
        self.servico.importar(self.tmp / "rel.csv")
        r = self.servico.resumo_inicio()
        self.assertEqual((r["fontes"], r["configurada"]), (0, True), "a importação configura")


class TestSigiloPorProcesso(Base):
    """O sigilo é do processo: vale para todos os registros dele, para o histórico
    gravado antes e para quem pergunta (a transcrição)."""

    D = date(2026, 10, 6)

    def setUp(self):
        super().setUp()
        self.n = ap.numero("0700124")

    def planilha(self, **filtros):
        from openpyxl import load_workbook

        arquivo = self.servico.exportar(date(2026, 10, 1), date(2026, 10, 31), self.amb.pauta,
                                        **filtros)
        return load_workbook(arquivo)

    def test_relatorio_importado_antes_e_o_portal_depois(self):
        (self.tmp / "rel.csv").write_text(
            f"Data;Hora;Processo;Partes\n06/10/2026;09:00;{self.n};Maria Souza x José Souza\n",
            encoding="utf-8")
        self.servico.importar(self.tmp / "rel.csv")
        self.assertFalse(self.servico.processo_sigiloso(self.n))
        # outra audiência do MESMO processo, que o portal marca como segredo de justiça
        self.servico.armazem.gravar([aud("0700124", self.D + timedelta(days=7), "10:00",
                                         sigiloso=True)], "esaj-tjal", None,
                                    registrar_novas=False)
        self.assertTrue(self.servico.processo_sigiloso(self.n))
        self.assertTrue(self.servico.processo_sigiloso(self.n.replace("-", "").replace(".", "")))
        self.assertFalse(self.servico.processo_sigiloso(ap.numero("0700999")))
        self.assertFalse(self.servico.processo_sigiloso("não é número"))
        lista = self.servico.listar(date(2026, 10, 1), date(2026, 10, 31))["audiencias"]
        self.assertEqual([(a["sistema"], a["sigiloso"]) for a in lista],
                         [("arquivo", True), ("esaj", True)])
        aba = self.planilha()["Pauta"]
        self.assertEqual({aba.cell(r, 6).value for r in (5, 6)}, {"(segredo de justiça)"})

    def test_historico_de_antes_do_sigilo(self):
        """As partes mudaram quando o processo era público; depois o portal passou a
        mostrar "(Segredo de Justiça)": a aba Alterações também mascara."""
        a = lambda **x: aud("0700124", self.D, "09:00", **x)  # noqa: E731
        self.servico.armazem.gravar([a(partes="Maria Souza x José Souza")], "esaj-tjal", None,
                                    registrar_novas=False)
        self.servico.armazem.gravar([a(partes="Maria Souza x Pedro Lima")], "esaj-tjal", None)
        self.servico.armazem.gravar([a(partes="", sigiloso=True)], "esaj-tjal", None)
        hist = self.planilha()["Alterações"]
        textos = " ".join(str(c.value or "") for linha in hist.iter_rows() for c in linha)
        self.assertIn("Partes: (segredo de justiça) → (segredo de justiça)", textos)
        self.assertNotIn("Pedro", textos)
        self.assertNotIn("Maria", textos)
        self.assertTrue(all(x["audiencia"]["sigiloso"] for x in self.servico.alteracoes()))
        # com as partes incluídas (escolha do usuário), aparecem
        hist = self.planilha(incluir_partes_sigilosos=True)["Alterações"]
        self.assertIn("Pedro", " ".join(str(c.value or "") for linha in hist.iter_rows()
                                        for c in linha))


class TestSigiloRevelado(Base):
    """O processo que a pauta revela sigiloso (e que o programa ainda não tratava
    como tal) é informado na hora - no resultado e a quem pediu (o servidor, que
    tira do acervo o que houver dele) -, uma vez só."""

    PERIODO = (date(2026, 10, 1), date(2026, 10, 31))

    def setUp(self):
        super().setUp()
        self.revelados: list[list[str]] = []
        self.n1, self.n2 = ap.numero("0700131"), ap.numero("0700132")

    def relatorio(self, partes1="(Segredo de Justiça)", partes2="João x Município"):
        (self.tmp / "rel.csv").write_text(
            "Data;Hora;Processo;Tipo;Partes\n"
            f"06/10/2026;09:00;{self.n1};Conciliação;{partes1}\n"
            f"07/10/2026;10:00;{self.n2};Una;{partes2}\n", encoding="utf-8")
        return self.tmp / "rel.csv"

    def test_importar(self):
        self.servico.quando_revelar_sigilo(self.revelados.append)
        r = self.servico.importar(self.relatorio())
        self.assertEqual(r["sigilosos_novos"], [self.n1])
        self.assertEqual(self.revelados, [[self.n1]])
        # importar de novo: nada de novo
        r = self.servico.importar(self.relatorio())
        self.assertNotIn("sigilosos_novos", r)
        self.assertEqual(self.revelados, [[self.n1]])
        # a negação não revela nada
        r = self.servico.importar(self.relatorio(partes2="Segredo de justiça: não"))
        self.assertNotIn("sigilosos_novos", r)
        self.assertFalse(self.servico.processo_sigiloso(self.n2))

    def test_ja_na_pasta_dos_sigilosos_nao_e_novo(self):
        self.amb.sigilosos.mkdir(parents=True)
        (self.amb.sigilosos / f"{self.n1}.pdf").write_bytes(apoio.pdf_bytes(1))
        self.servico.quando_revelar_sigilo(self.revelados.append)
        r = self.servico.importar(self.relatorio())
        self.assertNotIn("sigilosos_novos", r)
        self.assertEqual(self.revelados, [])

    def test_sincronizar_e_o_que_ninguem_ouviu_ainda(self):
        """O monitor sincroniza antes de o servidor pedir o aviso: o revelado fica
        guardado e é entregue quando o pedido chega."""
        self.servico.salvar_fonte("TJAL", "esaj", "")
        ExtratorFalso.roteiro["esaj-tjal"] = [
            ([aud("0700131", date(2026, 10, 6), "09:00", sigiloso=True),
              aud("0700132", date(2026, 10, 7), "10:00")], True, "https://portal.invalid/p")]
        r = self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        self.assertEqual(r["sigilosos_novos"], [self.n1])
        self.servico.quando_revelar_sigilo(self.revelados.append)
        self.assertEqual(self.revelados, [[self.n1]])
        self.servico.quando_revelar_sigilo(self.revelados.append)
        self.assertEqual(self.revelados, [[self.n1]], "entregue uma vez só")

    def test_quem_ouve_e_falha_nao_derruba_a_importacao(self):
        def falha(numeros):
            raise RuntimeError("acervo fora do ar")

        self.servico.quando_revelar_sigilo(falha)
        with self.assertLogs("pauta.servico", "WARNING") as registro:
            r = self.servico.importar(self.relatorio())
        self.assertEqual(r["novas"], 2)
        self.assertIn("acervo fora do ar", "\n".join(registro.output))


class TestRelatorioDepoisSincronizacao(Base):
    """Importar o relatório (a sincronização falhou) e depois sincronizar: a
    audiência não fica em dobro."""

    PERIODO = (date(2026, 10, 1), date(2026, 10, 31))

    def test_portal_absorve_a_importada(self):
        n1, n2 = ap.numero("0700101"), ap.numero("0700102")
        (self.tmp / "rel.csv").write_text(
            "Data;Hora;Processo;Tipo;Local;Partes\n"
            f"06/10/2026;09:00;{n1};Conciliação;Sala 9;Maria x Banco\n"
            f"07/10/2026;10:00;{n2};Una;;João x Município\n", encoding="utf-8")
        self.servico.importar(self.tmp / "rel.csv")
        self.servico.salvar_fonte("TJAL", "esaj", "")
        ExtratorFalso.roteiro["esaj-tjal"] = [
            ([aud("0700101", date(2026, 10, 6), "09:00", local=""),
              aud("0700102", date(2026, 10, 7), "10:00", situacao="Cancelada", sigiloso=True)],
             True, "https://portal.invalid/pauta")]
        self.servico.armazem.atualizar_fonte("esaj-tjal", ultima_sincronizacao=AGORA)
        r = self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        lista = self.servico.listar(*self.PERIODO)["audiencias"]
        self.assertEqual([(a["sistema"], a["processo"]) for a in lista],
                         [("esaj", n1), ("esaj", n2)], "nenhuma cópia do relatório sobrou")
        self.assertEqual(self.servico.resumo(lista)["total"], 2)
        um = lista[0]
        self.assertEqual((um["local"], um["partes"]), ("Sala 9", "Maria x Banco"),
                         "o que só o relatório sabia completa o registro do portal")
        self.assertTrue(lista[1]["sigiloso"])
        # a que só mudou de origem não é "nova"; a que o portal cancelou avisa
        tipos = [(x["tipo"], x["audiencia"]["processo"]) for x in self.servico.alteracoes()]
        self.assertEqual(tipos, [("cancelada", n2)])
        self.assertEqual((r["novas"], r["atualizadas"], r["canceladas"]), (0, 1, 1))
        # sincronizar de novo não muda nada
        ExtratorFalso.roteiro["esaj-tjal"] = [
            ([aud("0700101", date(2026, 10, 6), "09:00"),
              aud("0700102", date(2026, 10, 7), "10:00", situacao="Cancelada")], True, "")]
        r = self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        self.assertEqual((r["novas"], r["atualizadas"], r["removidas"]), (0, 0, 0))
        self.assertEqual(len(self.servico.listar(*self.PERIODO)["audiencias"]), 2)


class TestSenhaSoPorAgora(Base):
    """A senha cadastrada com "Lembrar neste computador" desligado vale para
    Sincronizar e Capturar (como no download), mas não para o monitoramento."""

    PERIODO = (date(2026, 10, 1), date(2026, 10, 31))

    def test_sessao(self):
        self.servico.salvar_fonte("TJAL", "esaj", "", "https://portal.invalid/pauta")
        self.cofre.dados.clear()
        self.servico.credenciais_sessao = {"esaj:TJAL": ("12345678900", "segredo")}
        ExtratorFalso.roteiro["esaj-tjal"] = [([], True, "")]
        self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        portal = PortalFalso.instancias[-1]
        self.assertEqual(portal.credenciais, ("12345678900", "segredo"))
        self.assertEqual(portal.opcoes.modo_login("esaj"), "senha",
                         "a tela de entrada não abre: a senha foi informada")
        # no monitoramento, só a senha guardada (o manual promete isso)
        ctx = ap.ContextoDeFundo()
        with self.assertRaises(ErroPauta):
            self.servico.sincronizar(ctx, None, *self.PERIODO)
        self.assertTrue(ctx.pediu_login)
        # a guardada continua valendo, e a da sessão tem a vez quando há as duas
        self.cofre.dados["esaj:TJAL"] = ("u", "s")
        ExtratorFalso.roteiro["esaj-tjal"] = [([], True, ""), ([], True, "")]
        self.servico.sincronizar(apoio.ContextoGravador(), None, *self.PERIODO)
        self.assertEqual(PortalFalso.instancias[-1].credenciais, ("12345678900", "segredo"))
        self.servico.sincronizar(ap.ContextoDeFundo(), None, *self.PERIODO)
        self.assertEqual(PortalFalso.instancias[-1].credenciais, ("u", "s"))


if __name__ == "__main__":
    unittest.main()
