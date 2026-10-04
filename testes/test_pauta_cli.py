"""Pauta: a linha de comando (python -m helestron pauta ...)."""

from __future__ import annotations

import contextlib
import io
import json
import unittest
from unittest import mock

from helestron.nucleo import caminhos
from helestron.pauta import cli

from testes import apoio_download as apoio
from testes import apoio_pauta as ap

N1, N2 = ap.numero("0700101"), ap.numero("0700102")


class TestCLI(apoio.PastaTemporaria):
    def setUp(self):
        super().setUp()
        self.amb = ap.config_temporaria(self.tmp)
        local = self.tmp / "local"
        for nome, valor in (("ARQUIVO_CONFIG", self.amb.cfg.arquivo), ("LOCAL", local),
                            ("ARQUIVO_PAUTA", local / "pauta.sqlite3"), ("LOGS", local / "Logs"),
                            ("BASE_USUARIO", self.tmp)):
            p = mock.patch.object(caminhos, nome, valor, create=True)
            p.start()
            self.addCleanup(p.stop)
        p = mock.patch.object(cli, "_configurar_registro", lambda: None)
        p.start()
        self.addCleanup(p.stop)
        (self.tmp / "rel.csv").write_text(
            "Data;Hora;Processo;Tipo;Situação;Partes\n"
            f"05/10/2026;09:00;{N1};Conciliação;Designada;Maria x Banco\n"
            f"06/10/2026;14:30;{N2};Una;Cancelada;João x Município\n", encoding="utf-8")

    def rodar(self, *argv) -> tuple[int, str, str]:
        saida, erro = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(saida), contextlib.redirect_stderr(erro):
            codigo = cli.main(list(argv))
        return codigo, saida.getvalue(), erro.getvalue()

    def test_uso(self):
        codigo, saida, _ = self.rodar()
        self.assertEqual(codigo, 2)
        self.assertIn("sincronizar", saida)
        codigo, _, erro = self.rodar("voar")
        self.assertEqual(codigo, 2)
        self.assertIn("escolha inválida", erro)
        codigo, _, erro = self.rodar("exportar", "--de", "31/02/2026")
        self.assertEqual(codigo, 2)
        self.assertIn("data inválida", erro)
        self.assertEqual(self.rodar("listar", "--help")[0], 0)

    def test_importar_listar_e_exportar(self):
        codigo, saida, _ = self.rodar("importar", str(self.tmp / "rel.csv"))
        self.assertEqual(codigo, 0, saida)
        self.assertIn("rel.csv: 2 audiências lidas · 2 novas · 0 atualizadas · 0 linhas ignoradas.", saida)
        self.assertTrue((self.tmp / "local" / "pauta.sqlite3").exists(), "o banco de LOCAL")

        codigo, saida, _ = self.rodar("listar", "--de", "2026-10-01", "--ate", "2026-10-31",
                                      "--json")
        self.assertEqual(codigo, 0)
        dados = json.loads(saida)
        self.assertEqual([a["processo"] for a in dados["audiencias"]], [N1, N2])
        codigo, saida, _ = self.rodar("listar", "--de", "05/10/2026", "--ate", "2026-10-31",
                                      "--situacao", "Cancelada")
        self.assertIn(f"06/10/2026 14:30  {N2}", saida)
        self.assertNotIn(N1, saida)
        self.assertIn("1 audiência no período", saida)
        codigo, saida, _ = self.rodar("listar", "--de", "2027-01-01")
        self.assertIn("Nenhuma audiência de 01/01/2027 a 08/01/2027.", saida)

        codigo, saida, _ = self.rodar("exportar", "--de", "2026-10-01", "--ate", "2026-10-31")
        self.assertEqual(codigo, 0, saida)
        arquivo = self.amb.pauta / "Pauta de audiências 2026-10-01 a 2026-10-31.xlsx"
        self.assertEqual(saida.strip(), f"Planilha gravada: {arquivo}")
        self.assertTrue(arquivo.exists())
        codigo, _, erro = self.rodar("exportar", "--pasta", str(self.amb.acervo / "x"))
        self.assertEqual(codigo, 1)
        self.assertIn("dentro do acervo", erro)


    def test_exportar_mascarando_mesmo_com_o_ajuste_ligado(self):
        from openpyxl import load_workbook

        (self.tmp / "sig.csv").write_text(
            "Data;Hora;Processo;Partes;Sigilo\n"
            f"05/10/2026;09:00;{N1};Maria x José;Segredo de Justiça\n", encoding="utf-8")
        self.assertEqual(self.rodar("importar", str(self.tmp / "sig.csv"))[0], 0)
        self.amb.cfg.definir("pauta", "incluir_partes_sigilosos", True)
        partes = {}
        for opcao in ("", "--sem-partes-sigilosos", "--incluir-partes-sigilosos"):
            argv = ["exportar", "--de", "2026-10-05", "--ate", "2026-10-05",
                    "--pasta", str(self.tmp / f"saida{opcao}")] + ([opcao] if opcao else [])
            codigo, saida, erro = self.rodar(*argv)
            self.assertEqual(codigo, 0, erro)
            arquivo = saida.strip().removeprefix("Planilha gravada: ")
            partes[opcao] = load_workbook(arquivo)["Pauta"].cell(5, 6).value
        self.assertEqual(partes, {"": "Maria x José", "--sem-partes-sigilosos":
                                  "(segredo de justiça)",
                                  "--incluir-partes-sigilosos": "Maria x José"})
        codigo, _, erro = self.rodar("exportar", "--sem-partes-sigilosos",
                                     "--incluir-partes-sigilosos")
        self.assertEqual(codigo, 2)

    def test_importar_com_erro(self):
        codigo, saida, erro = self.rodar("importar", str(self.tmp / "rel.csv"),
                                         str(self.tmp / "nao-existe.pdf"))
        self.assertEqual(codigo, 1)
        self.assertIn("rel.csv: 2 audiências lidas · 2 novas", saida)
        self.assertIn("nao-existe.pdf: Não encontrei o arquivo nao-existe.pdf.", erro)

    def test_fontes_e_sincronizar(self):
        codigo, saida, _ = self.rodar("fontes")
        self.assertIn("Nenhuma fonte configurada", saida)
        codigo, saida, _ = self.rodar("sincronizar")
        self.assertEqual(codigo, 1)
        codigo, _, erro = self.rodar("sincronizar", "--de", "2026-10-10", "--ate", "2026-10-01")
        self.assertEqual(codigo, 2)
        self.assertIn("A data final vem antes da inicial.", erro)
        codigo, saida, _ = self.rodar("fontes", "--adicionar", "tjal", "ESAJ", "--url",
                                      "https://www2.tjal.jus.br/pauta")
        self.assertEqual(codigo, 0)
        self.assertIn("Fonte esaj-tjal salva: TJAL · e-SAJ.", saida)
        codigo, _, erro = self.rodar("fontes", "--adicionar", "TJAL", "pje")
        self.assertEqual(codigo, 2)
        self.assertIn("e-SAJ ou eProc", erro)
        codigo, saida, _ = self.rodar("fontes")
        self.assertIn("esaj-tjal", saida)
        self.assertIn("ainda não sincronizada", saida)
        self.assertIn("https://www2.tjal.jus.br/pauta", saida)
        codigo, saida, _ = self.rodar("fontes", "--remover", "esaj-tjal")
        self.assertIn("removida", saida)

    def test_sincronizar_falha_com_frase(self):
        self.rodar("fontes", "--adicionar", "TJRS", "esaj")
        codigo, _, erro = self.rodar("sincronizar", "--de", "2026-10-01", "--ate", "2026-10-31")
        self.assertEqual(codigo, 1)
        self.assertIn("Falhou: Não consegui ler a pauta", erro)
        self.assertIn("O TJRS não usa o e-SAJ", erro)

    def test_pelo_modulo_principal(self):
        from helestron import __main__ as principal

        codigo = None
        with contextlib.redirect_stdout(io.StringIO()) as saida:
            try:
                codigo = principal.main(["pauta", "listar", "--de", "2026-10-01"])
            except SystemExit as fim:          # o __main__ pode encerrar com o código
                codigo = fim.code
        self.assertEqual(codigo, 0, saida.getvalue())


if __name__ == "__main__":
    unittest.main()
