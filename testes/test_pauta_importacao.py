"""Pauta: importação de relatórios (xlsx, "xls" que é HTML, ods, csv, html, docx, pdf)."""

from __future__ import annotations

import unittest
import zipfile
from unittest import mock
from datetime import date, datetime, time

from helestron.pauta import importacao, regras
from helestron.pauta.importacao import RelatorioInvalido, ler_relatorio
from helestron.pauta.servico import ServicoPauta

from testes import apoio_download as apoio
from testes import apoio_pauta as ap

R = regras.carregar()
N1, N2, N3 = ap.numero("0700101"), ap.numero("0700102"), ap.numero("0700103")
# Um CNJ que sobrevive como número na planilha: os 19 dígitos dão um múltiplo
# de 1024, que o double guarda exato (os outros o Excel corrompe).
N_EXATO = ap.numero("0700102", origem="0992")


def digitos(n: str) -> str:
    return n.replace("-", "").replace(".", "")


class TestFormatos(apoio.PastaTemporaria):
    def ler(self, nome: str):
        return ler_relatorio(self.tmp / nome, R, agora=datetime(2026, 10, 3, 10, 0))

    def test_xlsx_com_celulas_tipadas_aba_oculta_e_numero_corrompido(self):
        from openpyxl import Workbook

        livro = Workbook()
        aba = livro.active
        aba.title = "Pauta"
        aba.append(["Pauta de audiências — 2ª Vara Cível"])
        aba.append([])
        aba.append(["Data", "Hora", "Processo", "Tipo de audiência", "Situação", "Partes"])
        aba.append([datetime(2026, 10, 5), time(14, 30), N1, "Conciliação", "Designada", "A x B"])
        # o CNJ guardado como número: o zero da frente some (e volta), o dígito confere
        aba.append([date(2026, 10, 6), "09h00", int(digitos(N_EXATO)), "Instrução", "Cancelada",
                    "C x D"])
        # número corrompido pelo Excel (os últimos dígitos trocados)
        aba.append([date(2026, 10, 7), "10:00", float(digitos(N3)[:-3] + "999"), "Una", "", ""])
        aba.append([None, None, None, None, None, None])
        oculta = livro.create_sheet("Rascunho")
        oculta.sheet_state = "hidden"
        oculta.append(["Data", "Processo", "Tipo"])
        oculta.append([datetime(2026, 10, 8), N3, "Una"])
        resumo = livro.create_sheet("Resumo")
        resumo.append(["Tipo", "Quantidade"])
        resumo.append(["Conciliação", 1])
        livro.save(self.tmp / "pauta.xlsx")

        r = self.ler("pauta.xlsx")
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo, a.situacao) for a in r.audiencias], [
            (date(2026, 10, 5), "14:30", N1, "Conciliação", "Designada"),
            (date(2026, 10, 6), "09:00", N_EXATO, "Instrução e julgamento", "Cancelada")])
        self.assertEqual(r.ignoradas, 1)
        self.assertIn("corrompeu", r.avisos[0])
        a = r.audiencias[0]
        self.assertEqual((a.sistema, a.tribunal, a.fonte, a.origem),
                         ("arquivo", "TJAL", "arquivo", "pauta.xlsx"))

    def test_xls_de_verdade(self):
        """O .xls binário (OLE) vai pelo xlrd, com a data do Excel convertida."""
        import xlrd

        class Celula:
            def __init__(self, ctype, value):
                self.ctype, self.value = ctype, value

        class Aba:
            def __init__(self, nome, linhas, visibility=0, merged_cells=()):
                self.name, self._linhas, self.visibility = nome, linhas, visibility
                self.nrows = len(linhas)
                self.merged_cells = list(merged_cells)

            def row(self, i):
                return self._linhas[i]

        texto = lambda v: Celula(xlrd.XL_CELL_TEXT, v)  # noqa: E731
        data = Celula(xlrd.XL_CELL_DATE, 46300 + 14.5 / 24)       # 05/10/2026 14:30
        livro = mock.Mock(datemode=0)
        livro.sheets.return_value = [
            Aba("Oculta", [[texto("Data"), texto("Processo"), texto("Tipo")],
                           [Celula(xlrd.XL_CELL_DATE, 46301), texto(N3), texto("Una")]], 1),
            # "Data/Hora" mesclada na vertical (linhas 2 e 3): vale para as duas audiências
            Aba("Pauta", [[texto("Data/Hora"), texto("Processo"), texto("Tipo")],
                          [data, texto(N1), texto("Conciliação")],
                          [Celula(xlrd.XL_CELL_EMPTY, ""), texto(N2), texto("Una")],
                          [Celula(xlrd.XL_CELL_EMPTY, ""), Celula(xlrd.XL_CELL_EMPTY, ""),
                           Celula(xlrd.XL_CELL_EMPTY, "")]], merged_cells=[(1, 3, 0, 1)])]
        (self.tmp / "antigo.xls").write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 600)
        with mock.patch.object(xlrd, "open_workbook", return_value=livro) as abrir:
            r = self.ler("antigo.xls")
        abrir.assert_called_once_with(str(self.tmp / "antigo.xls"), formatting_info=True)
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo) for a in r.audiencias],
                         [(date(2026, 10, 5), "14:30", N1, "Conciliação"),
                          (date(2026, 10, 5), "14:30", N2, "Una")])

    def test_xls_que_e_html_e_html(self):
        lista = ap.no_periodo(ap.audiencias_padrao(), ap.INICIO, ap.FIM)
        html = ap.pagina(ap.html_esaj(lista))
        (self.tmp / "Agenda de audiências.xls").write_text(html, encoding="cp1252",
                                                           errors="xmlcharrefreplace")
        (self.tmp / "pauta.html").write_text(ap.pagina(ap.html_eproc(lista, total=8)),
                                             encoding="utf-8")
        for nome in ("Agenda de audiências.xls", "pauta.html"):
            with self.subTest(nome=nome):
                r = self.ler(nome)
                self.assertEqual(len(r.audiencias), 8)
                self.assertEqual({a.tribunal for a in r.audiencias}, {"TJAL"})
                sig = next(a for a in r.audiencias if a.processo == lista[4].processo)
                self.assertTrue(sig.sigiloso)

    def test_csv_ponto_e_virgula_em_ansi(self):
        texto = ("Data;Hora;Processo;Tipo;Situação;Local\r\n"
                 f"05/10/2026;09:00;{N1};Conciliação;Designada;Sala 1\r\n"
                 f"06/10/2026;14h;{N2};Audiência Una;Redesignada;Sala 2\r\n")
        (self.tmp / "pauta.csv").write_bytes(texto.encode("cp1252"))
        r = self.ler("pauta.csv")
        self.assertEqual([(a.processo, a.hora, a.tipo, a.situacao, a.local) for a in r.audiencias],
                         [(N1, "09:00", "Conciliação", "Designada", "Sala 1"),
                          (N2, "14:00", "Una", "Redesignada", "Sala 2")])

    def test_ods(self):
        ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
              'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
              'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')

        def texto(t):
            return f'<table:table-cell office:value-type="string"><text:p>{t}</text:p></table:table-cell>'

        linhas = [
            "".join(texto(t) for t in ("Data", "Hora", "Processo", "Tipo")),
            ('<table:table-cell office:value-type="date" office:date-value="2026-10-05">'
             '<text:p>05/10/26</text:p></table:table-cell>'
             '<table:table-cell office:value-type="time" office:time-value="PT14H30M00S">'
             '<text:p>14:30</text:p></table:table-cell>' + texto(N1) + texto("Custódia")
             + '<table:table-cell table:number-columns-repeated="1020"/>'),
        ]
        conteudo = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns}>'
                    '<office:body><office:spreadsheet><table:table table:name="Pauta">'
                    + "".join(f"<table:table-row>{x}</table:table-row>" for x in linhas)
                    + '<table:table-row table:number-rows-repeated="1000"><table:table-cell/>'
                      '</table:table-row></table:table></office:spreadsheet></office:body>'
                      '</office:document-content>')
        with zipfile.ZipFile(self.tmp / "pauta.ods", "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            z.writestr("content.xml", conteudo)
        r = self.ler("pauta.ods")
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo) for a in r.audiencias],
                         [(date(2026, 10, 5), "14:30", N1, "Custódia")])

    def test_docx(self):
        import docx

        documento = docx.Document()
        documento.add_heading("Pauta de audiências da semana", 1)
        tabela = documento.add_table(rows=1, cols=5)
        for celula, titulo in zip(tabela.rows[0].cells,
                                  ("Data/Hora", "Processo", "Tipo", "Situação", "Partes")):
            celula.text = titulo
        for valores in (("05/10/2026 09:00", N1, "Conciliação", "Designada", "A x B"),
                        ("06/10/2026 10:30", N2, "Justificação", "Realizada", "C x D")):
            linha = tabela.add_row().cells
            for celula, valor in zip(linha, valores):
                celula.text = valor
        documento.save(self.tmp / "pauta.docx")
        r = self.ler("pauta.docx")
        self.assertEqual([(a.data, a.hora, a.tipo, a.situacao) for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", "Conciliação", "Designada"),
            (date(2026, 10, 6), "10:30", "Justificação", "Realizada")])

    def test_docx_sem_tabela_por_linhas(self):
        import docx

        documento = docx.Document()
        documento.add_paragraph("Segunda-feira, 5 de outubro de 2026")
        documento.add_paragraph(f"09:00 — {N1} — Audiência de Conciliação — Maria x João")
        documento.add_paragraph(f"10:00 — {N2} — Audiência de Instrução e Julgamento")
        documento.add_paragraph("Terça-feira, 6 de outubro de 2026")
        documento.add_paragraph(f"08:30 — {N3} — Custódia — Cancelada")
        documento.save(self.tmp / "lista.docx")
        r = self.ler("lista.docx")
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo, a.situacao) for a in r.audiencias],
                         [(date(2026, 10, 5), "09:00", N1, "Conciliação", "Designada"),
                          (date(2026, 10, 5), "10:00", N2, "Instrução e julgamento", "Designada"),
                          (date(2026, 10, 6), "08:30", N3, "Custódia", "Designada")])

    def _pdf_com_tabela(self, nome: str) -> None:
        import pymupdf

        documento = pymupdf.open()
        colunas = (30, 100, 150, 330, 470, 560)
        cabecalho = ("Data", "Hora", "Processo", "Tipo de audiência", "Situação", "Partes")

        def nova_pagina(titulo: str):
            pagina = documento.new_page(width=842, height=595)
            pagina.insert_text((30, 50), titulo, fontsize=13)
            for x, t in zip(colunas, cabecalho):
                pagina.insert_text((x, 90), t, fontsize=9)
            pagina.insert_text((30, 570), "Página 1 de 2", fontsize=8)
            pagina.insert_text((600, 570), "Gerado em 03/10/2026 10:00", fontsize=8)
            return pagina

        p1 = nova_pagina("Pauta de audiências - 05/10/2026 a 09/10/2026")
        linhas = [(110, ("05/10/2026", "09:00", N1, "Conciliação", "Designada", "Maria x Banco")),
                  (125, ("05/10/2026", "14:30", N2, "Instrução e julgamento", "Cancelada",
                         "José da Silva"))]
        for y, valores in linhas:
            for x, t in zip(colunas, valores):
                p1.insert_text((x, y), t, fontsize=9)
        p1.insert_text((560, 137), "x Equatorial S.A.", fontsize=9)   # partes que quebraram
        p2 = nova_pagina("Pauta de audiências (continuação)")
        for x, t in zip(colunas, ("06/10/2026", "10:00", N3, "Custódia", "Designada", "")):
            p2.insert_text((x, 110), t, fontsize=9)
        documento.save(self.tmp / nome)
        documento.close()

    def test_pdf_tabela_por_texto(self):
        self._pdf_com_tabela("pauta.pdf")
        r = self.ler("pauta.pdf")
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo, a.situacao, a.partes)
                          for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", N1, "Conciliação", "Designada", "Maria x Banco"),
            (date(2026, 10, 5), "14:30", N2, "Instrução e julgamento", "Cancelada",
             "José da Silva x Equatorial S.A."),
            (date(2026, 10, 6), "10:00", N3, "Custódia", "Designada", "")])

    def test_pdf_sem_cabecalho(self):
        import pymupdf

        documento = pymupdf.open()
        pagina = documento.new_page()
        pagina.insert_text((40, 60), "Audiências designadas", fontsize=12)
        pagina.insert_text((40, 90), f"05/10/2026 09:00 {N1} Audiência de Conciliação", fontsize=9)
        pagina.insert_text((40, 105), f"05/10/2026 11:00 {N2} Audiência Una Redesignada",
                           fontsize=9)
        documento.save(self.tmp / "solto.pdf")
        documento.close()
        r = self.ler("solto.pdf")
        self.assertEqual([(a.hora, a.processo, a.tipo) for a in r.audiencias],
                         [("09:00", N1, "Conciliação"), ("11:00", N2, "Una")])

    def test_erros_com_frase(self):
        (self.tmp / "vazio.xlsx").write_bytes(b"")
        (self.tmp / "texto.txt").write_text("nada de pauta aqui\n", encoding="utf-8")
        (self.tmp / "quebrado.xlsx").write_bytes(b"PK\x03\x04" + b"\x00" * 100)
        (self.tmp / "velho.doc").write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 600)
        casos = {"nao-existe.xlsx": "não encontrei", "vazio.xlsx": "está vazio",
                 "texto.txt": "não encontrei uma tabela de audiências",
                 "quebrado.xlsx": "danificado", "velho.doc": ".docx"}
        for nome, trecho in casos.items():
            with self.subTest(nome=nome):
                with self.assertRaises(RelatorioInvalido) as erro:
                    self.ler(nome)
                self.assertIn(trecho, str(erro.exception))
        self.assertIn("*.xlsx", importacao.TIPOS_ARQUIVO[0][1])


class TestImportarNoServico(apoio.PastaTemporaria):
    def setUp(self):
        super().setUp()
        self.amb = ap.config_temporaria(self.tmp)
        self.servico = ServicoPauta(self.amb.cfg, self.tmp / "pauta.sqlite3")
        self.addCleanup(self.servico.fechar)

    def test_nao_duplica_a_do_portal_e_completa(self):
        from helestron.pauta import modelos

        do_portal = modelos.nova(sistema="esaj", tribunal="TJAL", data_=date(2026, 10, 5),
                                 processo=N1, hora="09:00", tipo_original="Conciliação",
                                 fonte="esaj-tjal")
        self.servico.armazem.gravar([do_portal], "esaj-tjal", None, registrar_novas=False)
        (self.tmp / "rel.csv").write_text(
            "Data;Hora;Processo;Tipo;Local\n"
            f"05/10/2026;09:00;{N1};Conciliação;Sala 9\n"
            f"06/10/2026;10:00;{N2};Una;Sala 2\n"
            "sem data;10:00;" + N3 + ";Una;\n", encoding="utf-8")
        eventos = []
        self.servico.eventos = lambda tipo, dados: eventos.append((tipo, dados))
        r = self.servico.importar(self.tmp / "rel.csv")
        self.assertEqual((r["novas"], r["atualizadas"], r["ignoradas"], r["total"]), (1, 1, 1, 2))
        self.assertIn(N3, r["avisos"][0])
        lista = self.servico.listar(date(2026, 10, 1), date(2026, 10, 31))["audiencias"]
        self.assertEqual([(a["sistema"], a["processo"], a["local"]) for a in lista],
                         [("esaj", N1, "Sala 9"), ("arquivo", N2, "Sala 2")])
        self.assertEqual(eventos[-1][0], "pauta")
        self.assertEqual(eventos[-1][1]["tipo"], "atualizada")
        # importar de novo não duplica nada
        r = self.servico.importar(self.tmp / "rel.csv")
        self.assertEqual((r["novas"], r["atualizadas"]), (0, 0))

    def test_duas_do_mesmo_processo_no_mesmo_horario_cada_uma_com_a_sua(self):
        """O portal tem uma Conciliação cancelada e uma Instrução designada do mesmo
        processo, às 09:00; o relatório traz as duas, com a sala de cada uma: cada
        registro do portal é completado pelo do seu tipo, e nada fica em dobro."""
        from helestron.pauta import modelos

        def do_portal(tipo, situacao):
            return modelos.nova(sistema="esaj", tribunal="TJAL", data_=date(2026, 10, 5),
                                processo=N1, hora="09:00", tipo_original=tipo,
                                situacao_original=situacao, fonte="esaj-tjal")

        for ordem, linhas in enumerate((
                (f"05/10/2026;09:00;{N1};Conciliação;Cancelada;Sala 1",
                 f"05/10/2026;09:00;{N1};Instrução;Designada;Sala 2"),
                (f"05/10/2026;09:00;{N1};Instrução;Designada;Sala 2",
                 f"05/10/2026;09:00;{N1};Conciliação;Cancelada;Sala 1"))):
            with self.subTest(ordem=ordem):
                servico = ServicoPauta(self.amb.cfg, self.tmp / f"ordem{ordem}.sqlite3")
                self.addCleanup(servico.fechar)
                servico.armazem.gravar([do_portal("Conciliação", "Cancelada"),
                                        do_portal("Instrução", "Designada")],
                                       "esaj-tjal", None, registrar_novas=False)
                (self.tmp / "rel.csv").write_text("Data;Hora;Processo;Tipo;Situação;Local\n"
                                                  + "\n".join(linhas) + "\n", encoding="utf-8")
                r = servico.importar(self.tmp / "rel.csv")
                self.assertEqual((r["novas"], r["atualizadas"]), (0, 2))
                lista = servico.listar(date(2026, 10, 1), date(2026, 10, 31))["audiencias"]
                self.assertEqual(sorted((a["sistema"], a["tipo"], a["local"]) for a in lista), [
                    ("esaj", "Conciliação", "Sala 1"),
                    ("esaj", "Instrução e julgamento", "Sala 2")])

    def test_as_duas_do_mesmo_horario_com_data_hora_e_processo_mesclados(self):
        """O relatório mescla Data, Hora e Processo das duas audiências do mesmo
        processo no mesmo horário: as duas entram (ou completam as do portal), e a
        de baixo não vira observação da de cima."""
        from openpyxl import Workbook
        from helestron.pauta import modelos

        livro = Workbook()
        aba = livro.active
        aba.append(["Data", "Hora", "Processo", "Tipo de audiência", "Situação", "Local"])
        aba.append([datetime(2026, 10, 5), "09:00", N1, "Conciliação", "Cancelada", "Sala 1"])
        aba.append([None, None, None, "Instrução", "Designada", "Sala 2"])
        aba.append([datetime(2026, 10, 5), "10:00", N2, "Instrução", "Designada", "Sala 1"])
        for coluna in "ABC":
            aba.merge_cells(f"{coluna}2:{coluna}3")
        livro.save(self.tmp / "mesclado.xlsx")

        r = self.servico.importar(self.tmp / "mesclado.xlsx")
        self.assertEqual((r["total"], r["novas"], r["atualizadas"]), (3, 3, 0))
        lista = self.servico.listar(date(2026, 10, 1), date(2026, 10, 31))["audiencias"]
        self.assertEqual(sorted((a["hora"], a["tipo"], a["situacao"], a["local"],
                                 a["observacoes"]) for a in lista), [
            ("09:00", "Conciliação", "Cancelada", "Sala 1", ""),
            ("09:00", "Instrução e julgamento", "Designada", "Sala 2", ""),
            ("10:00", "Instrução e julgamento", "Designada", "Sala 1", "")])

        # com as duas já trazidas do portal, cada uma é completada pela do seu tipo
        servico = ServicoPauta(self.amb.cfg, self.tmp / "com_portal.sqlite3")
        self.addCleanup(servico.fechar)
        servico.armazem.gravar([
            modelos.nova(sistema="esaj", tribunal="TJAL", data_=date(2026, 10, 5), processo=N1,
                         hora="09:00", tipo_original=tipo, situacao_original=situacao,
                         fonte="esaj-tjal")
            for tipo, situacao in (("Conciliação", "Cancelada"), ("Instrução", "Designada"))],
            "esaj-tjal", None, registrar_novas=False)
        r = servico.importar(self.tmp / "mesclado.xlsx")
        self.assertEqual((r["novas"], r["atualizadas"]), (1, 2))
        lista = servico.listar(date(2026, 10, 1), date(2026, 10, 31))["audiencias"]
        self.assertEqual(sorted((a["sistema"], a["hora"], a["tipo"], a["local"])
                                for a in lista), [
            ("arquivo", "10:00", "Instrução e julgamento", "Sala 1"),
            ("esaj", "09:00", "Conciliação", "Sala 1"),
            ("esaj", "09:00", "Instrução e julgamento", "Sala 2")])

    def test_a_do_relatorio_sem_par_no_portal_fica_na_pauta(self):
        """Pareada na importação com a do mesmo tipo, a outra audiência do relatório
        no mesmo horário fica como está - também na sincronização seguinte."""
        from helestron.pauta import modelos

        conciliacao = modelos.nova(sistema="esaj", tribunal="TJAL", data_=date(2026, 10, 5),
                                   processo=N1, hora="09:00", tipo_original="Conciliação",
                                   fonte="esaj-tjal")
        self.servico.armazem.gravar([conciliacao], "esaj-tjal", None, registrar_novas=False)
        (self.tmp / "rel.csv").write_text(
            "Data;Hora;Processo;Tipo;Local\n"
            f"05/10/2026;09:00;{N1};Conciliação;Sala 1\n"
            f"05/10/2026;09:00;{N1};Instrução;Sala 2\n", encoding="utf-8")
        r = self.servico.importar(self.tmp / "rel.csv")
        self.assertEqual((r["novas"], r["atualizadas"]), (1, 1))
        b = self.servico.armazem.gravar([modelos.de_dict(modelos.como_dict(conciliacao))],
                                        "esaj-tjal", (date(2026, 10, 1), date(2026, 10, 31)))
        self.assertEqual((b.novas, b.removidas), (0, 0))
        lista = self.servico.listar(date(2026, 10, 1), date(2026, 10, 31))["audiencias"]
        self.assertEqual(sorted((a["sistema"], a["tipo"], a["local"]) for a in lista), [
            ("arquivo", "Instrução e julgamento", "Sala 2"), ("esaj", "Conciliação", "Sala 1")])

    def test_relatorio_invalido_vira_valueerror(self):
        (self.tmp / "x.txt").write_text("nada", encoding="utf-8")
        with self.assertRaises(ValueError) as erro:
            self.servico.importar(self.tmp / "x.txt")
        self.assertTrue(str(erro.exception)[0].isupper(), "frase com maiúscula para a tela")

    def test_ida_e_volta_pela_planilha_exportada(self):
        lista = ap.no_periodo(ap.audiencias_padrao(), ap.INICIO, ap.FIM)
        (self.tmp / "p.html").write_text(ap.pagina(ap.html_esaj(lista)), encoding="utf-8")
        self.servico.importar(self.tmp / "p.html")
        arquivo = self.servico.exportar(ap.INICIO, ap.FIM, self.amb.pauta)
        outro = ServicoPauta(self.amb.cfg, self.tmp / "outro.sqlite3")
        self.addCleanup(outro.fechar)
        r = outro.importar(arquivo)
        self.assertEqual(r["novas"], 8)
        de_volta = outro.listar(ap.INICIO, ap.FIM)["audiencias"]
        sig = next(a for a in de_volta if a["processo"] == lista[4].processo)
        self.assertTrue(sig["sigiloso"])
        self.assertEqual(sig["partes"], "", "a máscara não vira nome de parte")
        virtual = next(a for a in de_volta if a["processo"] == lista[5].processo)
        self.assertEqual(virtual["link"], lista[5].link)


class TestRelatoriosDoMundoReal(apoio.PastaTemporaria):
    """Layouts comuns de relatório: data mesclada ou só na primeira linha do dia,
    preâmbulo com "Data:/Hora:" da emissão, título longo."""

    def ler(self, nome: str):
        return ler_relatorio(self.tmp / nome, R, agora=datetime(2026, 10, 3, 10, 0))

    def esperado(self):
        return [(date(2026, 10, 5), "09:00", N1, "Sala 1"),
                (date(2026, 10, 5), "10:00", N2, "Sala 1"),
                (date(2026, 10, 5), "11:00", N3, "Sala 2")]

    def test_xlsx_com_celulas_mescladas_e_preambulo_de_emissao(self):
        from openpyxl import Workbook

        livro = Workbook()
        aba = livro.active
        aba.append(["Relatório de audiências - 2ª Vara Cível"])
        aba.append(["Data:", datetime(2026, 10, 3), "Hora:", "10:15"])
        aba.append(["Vara:", "2ª Vara Cível", "Data de emissão:", "03/10/2026"])
        aba.append([])
        aba.append(["Data", "Hora", "Processo", "Tipo de audiência", "Local"])
        aba.append([datetime(2026, 10, 5), "09:00", N1, "Conciliação", "Sala 1"])
        aba.append([None, "10:00", N2, "Una", None])
        aba.append([None, "11:00", N3, "Instrução", "Sala 2"])
        aba.merge_cells("A6:A8")          # a data do dia, uma vez só
        aba.merge_cells("E6:E7")          # a mesma sala para as duas primeiras
        livro.save(self.tmp / "mesclada.xlsx")
        r = self.ler("mesclada.xlsx")
        self.assertEqual([(a.data, a.hora, a.processo, a.local) for a in r.audiencias],
                         self.esperado())
        self.assertEqual((r.ignoradas, r.avisos), (0, []))

    def test_ods_com_linhas_mescladas(self):
        ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
              'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
              'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')

        def texto(t, extra=""):
            return (f'<table:table-cell office:value-type="string"{extra}><text:p>{t}</text:p>'
                    '</table:table-cell>')

        coberta = "<table:covered-table-cell/>"
        linhas = ["".join(texto(t) for t in ("Data", "Hora", "Processo", "Local")),
                  texto("05/10/2026", ' table:number-rows-spanned="3"') + texto("09:00")
                  + texto(N1) + texto("Sala 1", ' table:number-rows-spanned="2"'),
                  coberta + texto("10:00") + texto(N2) + coberta,
                  coberta + texto("11:00") + texto(N3) + texto("Sala 2")]
        conteudo = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns}>'
                    '<office:body><office:spreadsheet><table:table table:name="Pauta">'
                    + "".join(f"<table:table-row>{x}</table:table-row>" for x in linhas)
                    + "</table:table></office:spreadsheet></office:body></office:document-content>")
        with zipfile.ZipFile(self.tmp / "mesclada.ods", "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            z.writestr("content.xml", conteudo)
        r = self.ler("mesclada.ods")
        self.assertEqual([(a.data, a.hora, a.processo, a.local) for a in r.audiencias],
                         self.esperado())

    # cada audiência em duas linhas: Data, Hora e Processo mesclados na vertical e,
    # embaixo, as partes (mescladas na horizontal sob o Tipo e a Situação)
    def confere_duas_linhas(self, r):
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo, a.partes) for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", N1, "Conciliação", "Fulano de Tal x Banco XYZ S/A"),
            (date(2026, 10, 5), "10:00", N2, "Instrução e julgamento", "Cicrano x Empresa W")])
        self.assertEqual((r.ignoradas, r.avisos), (0, []))

    def test_xlsx_com_a_audiencia_em_duas_linhas(self):
        from openpyxl import Workbook

        livro = Workbook()
        aba = livro.active
        aba.append(["Data", "Hora", "Processo", "Tipo de audiência", "Situação"])
        aba.append([datetime(2026, 10, 5), "09:00", N1, "Conciliação", "Designada"])
        aba.append([None, None, None, "Partes: Fulano de Tal x Banco XYZ S/A", None])
        aba.append([datetime(2026, 10, 5), "10:00", N2, "Instrução", "Designada"])
        aba.append([None, None, None, "Partes: Cicrano x Empresa W", None])
        for linha in (2, 4):
            for coluna in "ABC":
                aba.merge_cells(f"{coluna}{linha}:{coluna}{linha + 1}")
            aba.merge_cells(f"D{linha + 1}:E{linha + 1}")
        livro.save(self.tmp / "duas.xlsx")
        self.confere_duas_linhas(self.ler("duas.xlsx"))

    # Data, Hora e Processo mesclados porque são iguais, e embaixo o tipo e a
    # situação de OUTRA audiência do mesmo processo no mesmo horário
    def confere_mesmo_horario(self, r):
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo, a.situacao, a.local,
                           a.observacoes) for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", N1, "Conciliação", "Cancelada", "Sala 1", ""),
            (date(2026, 10, 5), "09:00", N1, "Instrução e julgamento", "Designada", "Sala 2", ""),
            (date(2026, 10, 5), "10:00", N2, "Instrução e julgamento", "Designada", "Sala 1", "")])
        self.assertEqual((r.ignoradas, r.avisos), (0, []))

    def test_xlsx_com_duas_audiencias_do_mesmo_processo_no_mesmo_horario(self):
        from openpyxl import Workbook

        livro = Workbook()
        aba = livro.active
        aba.append(["Data", "Hora", "Processo", "Tipo de audiência", "Situação", "Local"])
        aba.append([datetime(2026, 10, 5), "09:00", N1, "Conciliação", "Cancelada", "Sala 1"])
        aba.append([None, None, None, "Instrução", "Designada", "Sala 2"])
        aba.append([datetime(2026, 10, 5), "10:00", N2, "Instrução", "Designada", "Sala 1"])
        for coluna in "ABC":
            aba.merge_cells(f"{coluna}2:{coluna}3")
        livro.save(self.tmp / "mesmo_horario.xlsx")
        self.confere_mesmo_horario(self.ler("mesmo_horario.xlsx"))

    def test_html_com_duas_audiencias_do_mesmo_processo_no_mesmo_horario(self):
        (self.tmp / "mesmo_horario.html").write_text(
            "<html><head><title>Pauta de audiências</title></head><body><table>"
            "<tr><th>Data</th><th>Hora</th><th>Processo</th><th>Tipo</th><th>Situação</th>"
            "<th>Local</th></tr>"
            f"<tr><td rowspan=2>05/10/2026</td><td rowspan=2>09:00</td><td rowspan=2>{N1}</td>"
            "<td>Conciliação</td><td>Cancelada</td><td>Sala 1</td></tr>"
            "<tr><td>Instrução</td><td>Designada</td><td>Sala 2</td></tr>"
            f"<tr><td>05/10/2026</td><td>10:00</td><td>{N2}</td><td>Instrução</td>"
            "<td>Designada</td><td>Sala 1</td></tr></table></body></html>", encoding="utf-8")
        self.confere_mesmo_horario(self.ler("mesmo_horario.html"))

    def test_xls_com_a_audiencia_em_duas_linhas(self):
        import xlrd

        class Celula:
            def __init__(self, ctype, value):
                self.ctype, self.value = ctype, value

        class Aba:
            name, visibility = "Pauta", 0

            def __init__(self, linhas, merged_cells):
                self._linhas, self.nrows, self.merged_cells = linhas, len(linhas), merged_cells

            def row(self, i):
                return self._linhas[i]

        def linha(*valores):
            return [Celula(xlrd.XL_CELL_EMPTY, "") if v is None else Celula(xlrd.XL_CELL_TEXT, v)
                    for v in valores]

        aba = Aba([linha("Data", "Hora", "Processo", "Tipo", "Situação"),
                   linha("05/10/2026", "09:00", N1, "Conciliação", "Designada"),
                   linha(None, None, None, "Partes: Fulano de Tal x Banco XYZ S/A", None),
                   linha("05/10/2026", "10:00", N2, "Instrução", "Designada"),
                   linha(None, None, None, "Partes: Cicrano x Empresa W", None)],
                  [(l0, l0 + 2, c, c + 1) for l0 in (1, 3) for c in range(3)]
                  + [(l0 + 1, l0 + 2, 3, 5) for l0 in (1, 3)])
        livro = mock.Mock(datemode=0)
        livro.sheets.return_value = [aba]
        (self.tmp / "duas.xls").write_bytes(b"\xd0\xcf\x11\xe0" + b"\x00" * 600)
        with mock.patch.object(xlrd, "open_workbook", return_value=livro):
            self.confere_duas_linhas(self.ler("duas.xls"))

    def _ods(self, nome: str, linhas: list[str]) -> None:
        ns = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
              'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
              'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')
        conteudo = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {ns}>'
                    '<office:body><office:spreadsheet><table:table table:name="Pauta">'
                    + "".join(f"<table:table-row>{x}</table:table-row>" for x in linhas)
                    + "</table:table></office:spreadsheet></office:body></office:document-content>")
        with zipfile.ZipFile(self.tmp / nome, "w") as z:
            z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
            z.writestr("content.xml", conteudo)

    @staticmethod
    def _texto_ods(t, extra=""):
        return (f'<table:table-cell office:value-type="string"{extra}><text:p>{t}</text:p>'
                '</table:table-cell>')

    def test_ods_com_a_audiencia_em_duas_linhas(self):
        texto = self._texto_ods
        mescla = ' table:number-rows-spanned="2"'
        # o LibreOffice junta as cobertas vizinhas numa só, com number-columns-repeated
        cobertas = '<table:covered-table-cell table:number-columns-repeated="3"/>'
        largura = ' table:number-columns-spanned="2"'
        linhas = ["".join(texto(t) for t in ("Data", "Hora", "Processo", "Tipo", "Situação"))]
        for processo, hora, tipo, partes in ((N1, "09:00", "Conciliação",
                                              "Partes: Fulano de Tal x Banco XYZ S/A"),
                                             (N2, "10:00", "Instrução",
                                              "Partes: Cicrano x Empresa W")):
            linhas.append(texto("05/10/2026", mescla) + texto(hora, mescla)
                          + texto(processo, mescla) + texto(tipo) + texto("Designada"))
            linhas.append(cobertas + texto(partes, largura) + "<table:covered-table-cell/>")
        self._ods("duas.ods", linhas)
        self.confere_duas_linhas(self.ler("duas.ods"))

    def test_ods_cobertas_de_mesclas_diferentes_numa_so(self):
        """Data e Hora mescladas lado a lado: a coberta comprimida da linha de baixo
        devolve a cada coluna o valor da sua mescla (a hora não vira a data)."""
        texto = self._texto_ods
        mescla = ' table:number-rows-spanned="2"'
        linhas = ["".join(texto(t) for t in ("Data", "Hora", "Processo", "Tipo")),
                  texto("05/10/2026", mescla) + texto("09:00", mescla) + texto(N1)
                  + texto("Conciliação"),
                  '<table:covered-table-cell table:number-columns-repeated="2"/>' + texto(N2)
                  + texto("Conciliação")]
        self._ods("concentrada.ods", linhas)
        r = self.ler("concentrada.ods")
        self.assertEqual([(a.data, a.hora, a.processo) for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", N1), (date(2026, 10, 5), "09:00", N2)])

    def test_csv_com_a_data_so_na_primeira_linha_do_dia(self):
        (self.tmp / "rel.csv").write_text(
            "Data;Hora;Processo;Local\n"
            f"05/10/2026;09:00;{N1};Sala 1\n;10:00;{N2};Sala 1\n;11:00;{N3};Sala 2\n",
            encoding="utf-8")
        r = self.ler("rel.csv")
        self.assertEqual([(a.data, a.hora, a.processo, a.local) for a in r.audiencias],
                         self.esperado())

    def test_html_com_rowspan_e_grupo_com_contagem(self):
        html = ap.pagina(
            "<table><tr><th>Data</th><th>Hora</th><th>Processo</th><th>Partes</th></tr>"
            "<tr><td colspan='2'>05/10/2026 - Segunda-feira</td><td></td><td>3 audiências</td>"
            "</tr>"
            f"<tr><td rowspan='3'></td><td>09:00</td><td>{N1}</td><td>A x B</td></tr>"
            f"<tr><td>10:00</td><td>{N2}</td><td>C x D</td></tr>"
            f"<tr><td>11:00</td><td>{N3}</td><td>E x F</td></tr></table>")
        (self.tmp / "agenda.xls").write_text(html, encoding="utf-8")
        r = self.ler("agenda.xls")
        self.assertEqual([(a.data, a.hora, a.processo, a.partes) for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", N1, "A x B"), (date(2026, 10, 5), "10:00", N2, "C x D"),
            (date(2026, 10, 5), "11:00", N3, "E x F")])

    def test_pdf_com_data_e_hora_da_emissao_no_topo(self):
        import pymupdf

        documento = pymupdf.open()
        pagina = documento.new_page(width=842, height=595)
        colunas = (30, 100, 150, 330, 470, 560)
        pagina.insert_text((30, 40), "Pauta de audiências", fontsize=13)
        for x, t in ((30, "Data:"), (100, "03/10/2026"), (330, "Hora:"), (470, "10:15")):
            pagina.insert_text((x, 60), t, fontsize=9)
        for x, t in zip(colunas, ("Data", "Hora", "Processo", "Tipo de audiência", "Situação",
                                  "Partes")):
            pagina.insert_text((x, 90), t, fontsize=9)
        for x, t in zip(colunas, ("05/10/2026", "09:00", N1, "Conciliação", "Designada",
                                  "Maria x Banco")):
            pagina.insert_text((x, 110), t, fontsize=9)
        for x, t in zip(colunas, ("05/10/2026", "14:30", N2, "Una", "Cancelada", "José")):
            pagina.insert_text((x, 125), t, fontsize=9)
        pagina.insert_text((560, 137), "x Equatorial S.A.", fontsize=9)
        documento.save(self.tmp / "emitido.pdf")
        documento.close()
        r = self.ler("emitido.pdf")
        self.assertEqual([(a.data, a.hora, a.processo, a.tipo, a.situacao, a.partes)
                          for a in r.audiencias], [
            (date(2026, 10, 5), "09:00", N1, "Conciliação", "Designada", "Maria x Banco"),
            (date(2026, 10, 5), "14:30", N2, "Una", "Cancelada", "José x Equatorial S.A.")])


class TestSigiloDoRelatorioNaPlanilha(apoio.PastaTemporaria):
    """"Ministério Público" nas partes não desfaz o selo de segredo de justiça: o
    relatório importado sai mascarado na planilha."""

    def test_ministerio_publico_ao_lado_do_selo(self):
        from openpyxl import load_workbook

        amb = ap.config_temporaria(self.tmp)
        servico = ServicoPauta(amb.cfg, self.tmp / "pauta.sqlite3")
        self.addCleanup(servico.fechar)
        icone = "<img src='data:,' title='Segredo de Justiça'>"
        html = ap.pagina(
            "<table><tr><th>Data</th><th>Hora</th><th>Processo</th><th>Classe</th>"
            "<th>Partes</th></tr>"
            f"<tr><td>05/10/2026</td><td>09:00</td><td>{N1}</td><td>Ato Infracional</td>"
            f"<td>Ministério Público do Estado de Alagoas x Adolescente J.S. {icone}</td></tr>"
            f"<tr><td>05/10/2026</td><td>10:00</td><td>{N2}</td>"
            "<td>Ação Civil Pública - Segredo de Justiça</td>"
            "<td>Associação X x Município Y</td></tr></table>")
        (self.tmp / "pauta.html").write_text(html, encoding="utf-8")
        servico.importar(self.tmp / "pauta.html")
        lista = servico.listar(date(2026, 10, 5), date(2026, 10, 5))["audiencias"]
        self.assertEqual([a["sigiloso"] for a in lista], [True, True])
        arquivo = servico.exportar(date(2026, 10, 5), date(2026, 10, 5), amb.pauta)
        aba = load_workbook(arquivo)["Pauta"]
        self.assertEqual([aba.cell(r, 6).value for r in (5, 6)],
                         ["(segredo de justiça)", "(segredo de justiça)"])


class TestRelatorioSoMarcaSigiloComSelo(apoio.PastaTemporaria):
    """O relatório importado vira a regra única do sigilo (nucleo/sigilo.py: a
    pauta marca, o processo sai do acervo e da IA). Textos comuns do Local e das
    Observações não podem marcar ninguém; o selo de verdade continua marcando."""

    CASOS = [  # (local, observações, partes, sigiloso?)
        ("Fórum Des. Jairon Maia, Nível 1, Sala 3", "", "Fulano x Banco", False),
        ("Sala 5 - Bloco Nível 2", "", "Empresa x Empresa", False),
        ("Sala 2 - Térreo", "Segredo de justiça: não", "Beltrano x Estado", False),
        ("Sala 2", "Sem segredo de justiça", "Cicrano x Município", False),
        ("Sala 2", "Não sigiloso", "A x B", False),
        ("Sala 2", "Nível de sigilo: 0 (público)", "Fulana x Banco", False),
        ("Sala 2", "Oitiva de testemunha sigilosa (Prov. 32)", "MP x Réu", False),
        ("Sala 1", "Retirado o sigilo dos autos", "A x B", False),
        ("Sala 1", "", "A x B", False),
        ("Sala 1", "", "(Segredo de Justiça)", True),
        ("Sala 1", "Processo em segredo de justiça", "A x B", True),
        ("Sala 1", "Segredo de Justiça", "Ministério Público x Adolescente", True),
    ]

    def test_importacao_e_regra_unica(self):
        from helestron.nucleo import sigilo

        amb = ap.config_temporaria(self.tmp)
        banco = self.tmp / "pauta.sqlite3"
        servico = ServicoPauta(amb.cfg, banco)
        self.addCleanup(servico.fechar)
        numeros = [ap.numero(f"07009{i:02d}") for i in range(len(self.CASOS))]
        linhas = ["Data;Hora;Processo;Tipo;Situação;Local;Observações;Partes"]
        for n, (local, obs, partes, _) in zip(numeros, self.CASOS):
            linhas.append(f"20/10/2026;09:00;{n};Instrução;Designada;{local};{obs};{partes}")
        (self.tmp / "rel.csv").write_text("\n".join(linhas) + "\n", encoding="utf-8")
        r = servico.importar(self.tmp / "rel.csv")
        self.assertEqual(r["total"], len(self.CASOS))
        sigilo.esquecer_pauta()
        marcados = sigilo.chaves_da_pauta(banco)
        for n, (local, obs, partes, esperado) in zip(numeros, self.CASOS):
            with self.subTest(local=local, observacoes=obs, partes=partes):
                self.assertEqual(n in marcados, esperado)
        self.assertEqual(r["sigilosos_novos"],
                         sorted(n for n, caso in zip(numeros, self.CASOS) if caso[3]))


if __name__ == "__main__":
    unittest.main()
