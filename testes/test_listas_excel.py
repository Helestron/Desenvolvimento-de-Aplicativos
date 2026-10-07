"""Relação de processos em planilha do Excel: as variantes que o mundo real
manda (exportação de sistemas, arquivo que o openpyxl recusa, .xlsb, planilha
com senha) - e a que deixava o 'Baixar processos' recusar o Excel.

E, da revisão da 1.0.2: a célula fora do padrão em outra coluna (vírgula
decimal, texto sem o tipo, data não ISO) não derruba a leitura; a aba oculta
não entra no lote pela porta dos fundos (o XML cru, os bytes crus); a coluna
"Processo" com o CNJ corrompido não faz ler a de "Processo de origem"; o
título do relatório não passa por cabeçalho; o .txt "Texto Unicode" não
inventa senha; o .ods com células repetidas não desloca as colunas."""

import re
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import openpyxl

from helestron.nucleo import listas

N1 = "0700001-70.2024.8.02.0058"
N2 = "0700002-55.2024.8.02.0058"
ORIGEM1 = "0800123-45.2019.8.02.0001"
ORIGEM2 = "0800999-11.2018.8.02.0001"
# O CNJ digitado numa célula do Excel sem o formato Texto: um double, que já
# não guarda os 20 dígitos (os de N1 e N2, como o Excel os grava).
COMO_NUMERO1 = 7.00001702024802e18
COMO_NUMERO2 = 7.00002552024802e18

_NS = ('xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')


def xlsx_a_mao(destino: Path, abas, compartilhadas=None, ocultas=()) -> Path:
    """Planilha como a dos exportadores de sistemas (XML escrito à mão).

    abas: [(nome, conteúdo do sheetData)]; ocultas: os nomes das abas ocultas."""
    tipos = ['<Default Extension="rels" ContentType="application/vnd.openxmlformats-'
             'package.relationships+xml"/>',
             '<Default Extension="xml" ContentType="application/xml"/>',
             '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
             'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>']
    folhas, relacoes = [], []
    for i, (nome, _) in enumerate(abas, 1):
        tipos.append(f'<Override PartName="/xl/worksheets/sheet{i}.xml" ContentType="application/'
                     'vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>')
        estado = ' state="hidden"' if nome in ocultas else ""
        folhas.append(f'<sheet name="{nome}" sheetId="{i}"{estado} r:id="rId{i}"/>')
        relacoes.append(f'<Relationship Id="rId{i}" Type="http://schemas.openxmlformats.org/'
                        'officeDocument/2006/relationships/worksheet" '
                        f'Target="worksheets/sheet{i}.xml"/>')
    if compartilhadas is not None:
        relacoes.append(f'<Relationship Id="rId{len(abas) + 1}" Type="http://schemas.'
                        'openxmlformats.org/officeDocument/2006/relationships/sharedStrings" '
                        'Target="sharedStrings.xml"/>')
    with zipfile.ZipFile(destino, "w") as z:
        z.writestr("[Content_Types].xml", '<?xml version="1.0" encoding="UTF-8"?><Types xmlns='
                   '"http://schemas.openxmlformats.org/package/2006/content-types">'
                   + "".join(tipos) + "</Types>")
        z.writestr("_rels/.rels", '<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns='
                   '"http://schemas.openxmlformats.org/package/2006/relationships"><Relationship '
                   'Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                   'relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr("xl/workbook.xml", f'<?xml version="1.0" encoding="UTF-8"?><workbook {_NS}>'
                   f'<sheets>{"".join(folhas)}</sheets></workbook>')
        z.writestr("xl/_rels/workbook.xml.rels", '<?xml version="1.0" encoding="UTF-8"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/'
                   f'relationships">{"".join(relacoes)}</Relationships>')
        for i, (_, dados) in enumerate(abas, 1):
            z.writestr(f"xl/worksheets/sheet{i}.xml", f'<?xml version="1.0" encoding="UTF-8"?>'
                       f'<worksheet {_NS}><dimension ref="A1"/><sheetData>{dados}</sheetData>'
                       "</worksheet>")
        if compartilhadas is not None:
            si = "".join(f"<si><t>{s}</t></si>" for s in compartilhadas)
            z.writestr("xl/sharedStrings.xml", f'<?xml version="1.0" encoding="UTF-8"?><sst {_NS}>'
                       f"{si}</sst>")
    return destino


def texto(ref: str, valor: str) -> str:
    return f'<c r="{ref}" t="inlineStr"><is><t>{valor}</t></is></c>'


def linha(n: int, *celulas: str) -> str:
    return f'<row r="{n}">{"".join(celulas)}</row>'


class TestRelacaoNoExcel(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        livro = openpyxl.Workbook()
        aba = livro.active
        aba.append(["Processo", "Classe"])
        aba.append([N1, "Procedimento Comum"])
        aba.append([N2, "Execução"])
        self.base = self.tmp / "base.xlsx"
        livro.save(self.base)

    def _variante(self, nome, trocar, alvo="xl/worksheets/sheet"):
        destino = self.tmp / nome
        with zipfile.ZipFile(self.base) as zi, zipfile.ZipFile(destino, "w") as zo:
            for item in zi.infolist():
                dados = zi.read(item.filename)
                if item.filename.startswith(alvo):
                    dados = trocar(dados)
                zo.writestr(item, dados)
        return destino

    def _numeros(self, caminho):
        return [n.formatado for n in listas.ler_arquivo(caminho).processos]

    def test_planilha_comum(self):
        self.assertEqual(self._numeros(self.base), [N1, N2])

    def test_dimensao_declarada_errada_como_nas_exportacoes(self):
        # O defeito relatado: a planilha exportada declara a dimensão "A1",
        # e o modo read_only lia só a primeira linha (o cabeçalho).
        arq = self._variante("exportada.xlsx", lambda d: re.sub(
            rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', d))
        self.assertEqual(self._numeros(arq), [N1, N2])

    def test_sem_dimensao(self):
        arq = self._variante("sem-dimensao.xlsx", lambda d: re.sub(
            rb'<dimension ref="[^"]*"/>', b"", d))
        self.assertEqual(self._numeros(arq), [N1, N2])

    def test_openpyxl_recusa_e_le_pelo_xml(self):
        arq = self._variante("variante.xlsx", lambda d: b"<lixo>", alvo="xl/workbook.xml")
        self.assertEqual(self._numeros(arq), [N1, N2])

    def test_aba_de_grafico(self):
        # A aba de gráfico (relatório com gráfico em aba própria) não tem
        # células: o openpyxl a devolve sem iter_rows, e a leitura caía com
        # AttributeError (erro 500 na tela, traceback na linha de comando).
        from openpyxl.chart import BarChart, Reference

        livro = openpyxl.load_workbook(self.base)
        grafico = BarChart()
        grafico.add_data(Reference(livro.active, min_col=2, min_row=1, max_row=3))
        livro.create_chartsheet("Gráfico", 0).add_chart(grafico)
        arq = self.tmp / "com-grafico.xlsx"
        livro.save(arq)
        self.assertEqual(self._numeros(arq), [N1, N2])

    def test_xlsb_explica_o_que_fazer(self):
        arq = self.tmp / "binaria.xlsb"
        with zipfile.ZipFile(arq, "w") as z:
            z.writestr("[Content_Types].xml", "<x/>")
            z.writestr("xl/workbook.bin", b"\x00")
        with self.assertRaises(listas.ListaInvalida) as ctx:
            listas.ler_arquivo(arq)
        self.assertIn(".xlsx", str(ctx.exception))

    def test_planilha_com_senha_de_abertura(self):
        arq = self.tmp / "senha.xlsx"
        arq.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512
                        + "EncryptedPackage".encode("utf-16-le") + b"\x00" * 512)
        with self.assertRaises(listas.ListaInvalida) as ctx:
            listas.ler_arquivo(arq)
        self.assertIn("senha", str(ctx.exception))


class _ComPasta(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.dir.cleanup)
        self.tmp = Path(self.dir.name)

    def numeros(self, caminho):
        return [n.formatado for n in listas.ler_arquivo(caminho).processos]

    def recusada_pelos_corrompidos(self, caminho):
        with self.assertRaises(listas.ListaInvalida) as ctx:
            listas.ler_arquivo(caminho)
        self.assertIn("como NÚMERO", str(ctx.exception))
        return str(ctx.exception)


class TestCelulaForaDoPadraoEmOutraColuna(_ComPasta):
    """Achado R25: no modo read_only, o openpyxl só converte a célula ao
    percorrer a aba, e o ValueError dele derrubava a leitura inteira - com a
    coluna "Processo" perfeita."""

    def _planilha(self, nome, segunda_coluna, celulas):
        return xlsx_a_mao(self.tmp / nome, [("Plan1",
            linha(1, texto("A1", "Processo"), texto("B1", segunda_coluna))
            + linha(2, texto("A2", N1), celulas[0])
            + linha(3, texto("A3", N2), celulas[1]))])

    def test_virgula_decimal(self):
        # exportador .NET com o Windows em português
        arq = self._planilha("valores.xlsx", "Valor da causa",
                             ['<c r="B2"><v>1500,50</v></c>', '<c r="B3"><v>200,00</v></c>'])
        self.assertEqual(self.numeros(arq), [N1, N2])

    def test_data_que_nao_e_iso(self):
        arq = self._planilha("datas.xlsx", "Distribuição",
                             ['<c r="B2" t="d"><v>15/01/2024</v></c>',
                              '<c r="B3" t="d"><v>16/01/2024</v></c>'])
        self.assertEqual(self.numeros(arq), [N1, N2])

    def test_numero_do_processo_gravado_sem_o_tipo(self):
        arq = xlsx_a_mao(self.tmp / "sem-tipo.xlsx", [("Plan1",
            linha(1, texto("A1", "Processo"))
            + linha(2, f'<c r="A2"><v>{N1}</v></c>')
            + linha(3, f'<c r="A3"><v>{N2}</v></c>'))])
        self.assertEqual(self.numeros(arq), [N1, N2])

    def test_leitura_pelo_xml_segue_as_regras_da_leitura_normal(self):
        # A célula estranha manda a planilha para a leitura direta do XML: ela
        # também deixa de fora a aba oculta e a coluna de outro processo, e
        # lê os textos compartilhados.
        arq = xlsx_a_mao(self.tmp / "regras.xlsx", [
            ("Relação", linha(1, '<c r="A1" t="s"><v>0</v></c>',
                              texto("B1", "Processo de origem"), texto("C1", "Valor"))
             + linha(2, '<c r="A2" t="s"><v>1</v></c>', texto("B2", ORIGEM1),
                     '<c r="C2"><v>1500,50</v></c>')),
            ("Antigos", linha(1, texto("A1", "Processo")) + linha(2, texto("A2", ORIGEM2)))],
            compartilhadas=["Processo", N1], ocultas={"Antigos"})
        leitura = listas.ler_arquivo(arq)
        self.assertEqual([n.formatado for n in leitura.processos], [N1])
        self.assertEqual(leitura.avisos, ["aba oculta 'Antigos' ignorada"])

    def test_erro_inesperado_de_um_leitor_vira_frase_em_portugues(self):
        # Nada de 'could not convert...' na tela, nem de traceback e causa
        # 'inesperado' no 'baixar --lista'.
        arq = self.tmp / "relacao.html"
        arq.write_text(f"<html><table><tr><td>{N1}</td></tr></table></html>", encoding="utf-8")

        def quebra(caminho, col):
            col.texto(N2)             # o que leu antes de parar não fica
            raise ValueError("could not convert string to float: 'x'")

        with mock.patch.dict(listas._LEITORES, {"html": quebra}):
            leitura = listas.ler_arquivo(arq)
        # a última tentativa (o arquivo como texto) ainda acha o número
        self.assertEqual([n.formatado for n in leitura.processos], [N1])
        arq.write_text("<html><table><tr><td>nada</td></tr></table></html>", encoding="utf-8")
        with mock.patch.dict(listas._LEITORES, {"html": quebra}):
            with self.assertRaises(listas.ListaInvalida) as ctx:
                listas.ler_arquivo(arq)
        self.assertIn("não consegui ler o arquivo relacao.html", str(ctx.exception))
        self.assertNotIn("could not convert", str(ctx.exception))


class TestAbaOcultaNaoEntraNoLote(_ComPasta):
    """Achado R24: a aba visível com o CNJ como NÚMERO não achava processo, e
    a leitura do XML cru (ou dos bytes crus) trazia a aba OCULTA para o lote,
    ao lado do aviso de que ela fora ignorada."""

    def _livro(self):
        livro = openpyxl.Workbook()
        rel = livro.active
        rel.title = "Relação"
        rel.append(["Processo", "Classe"])
        rel.append([COMO_NUMERO1, "Execução"])
        rel.append([COMO_NUMERO2, "Comum"])
        base = livro.create_sheet("Base antiga")
        base.sheet_state = "hidden"
        base.append(["Processo"])
        base.append([ORIGEM1])
        return livro

    def test_xlsx(self):
        arq = self.tmp / "relacao.xlsx"
        self._livro().save(arq)
        self.recusada_pelos_corrompidos(arq)

    def test_xlsx_sem_compressao(self):
        # O padrão do SheetJS ("Exportar para Excel" de sistemas web): o texto
        # das abas, ocultas inclusive, aparece nos bytes do arquivo.
        comprimido = self.tmp / "comprimido.xlsx"
        self._livro().save(comprimido)
        arq = self.tmp / "sem-compressao.xlsx"
        with zipfile.ZipFile(comprimido) as zi, zipfile.ZipFile(arq, "w", zipfile.ZIP_STORED) as zo:
            for item in zi.infolist():
                zo.writestr(item.filename, zi.read(item.filename))
        self.assertIn(ORIGEM1.encode(), arq.read_bytes())
        self.recusada_pelos_corrompidos(arq)

    def _xls(self, abas):
        """.xls (OLE) com o texto da aba oculta nos bytes, como no BIFF, e o
        xlrd trocado por um livro com estas abas."""
        import xlrd

        class Aba:
            def __init__(self, nome, oculta, linhas):
                self.name, self.visibility, self._linhas = nome, int(oculta), linhas
                self.nrows = len(linhas)

            def row_values(self, i):
                return self._linhas[i]

        class Livro:
            def sheets(self):
                return [Aba(*a) for a in abas]

        arq = self.tmp / "relacao.xls"
        arq.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 600
                        + f"Processo {ORIGEM1} Processo de origem {ORIGEM2}".encode("utf-16-le")
                        + b"\x00" * 600)
        return arq, mock.patch.object(xlrd, "open_workbook", return_value=Livro())

    def test_xls_aba_oculta(self):
        arq, xlrd_falso = self._xls([("Relação", False, [["Processo"], [COMO_NUMERO1]]),
                                     ("Base", True, [["Processo"], [ORIGEM1]])])
        with xlrd_falso:
            self.recusada_pelos_corrompidos(arq)

    def test_xls_coluna_de_origem(self):
        arq, xlrd_falso = self._xls([
            ("Relação", False, [["Processo", "Processo de origem"], [COMO_NUMERO1, ORIGEM2]])])
        with xlrd_falso:
            self.recusada_pelos_corrompidos(arq)

    def test_xls_que_o_xlrd_nao_abre_ainda_e_lido_pelo_texto(self):
        # A última tentativa continua valendo para o arquivo que o leitor
        # dele não abriu (danificado, ou um .doc com a extensão trocada).
        import xlrd

        arq, _ = self._xls([])
        with mock.patch.object(xlrd, "open_workbook", side_effect=xlrd.XLRDError("BOF")):
            self.assertEqual(self.numeros(arq), [ORIGEM1, ORIGEM2])

    def test_ods_com_aba_oculta(self):
        arq = _ods(self.tmp / "relacao.ods", [
            ("Relação", False, [_celula("Processo")], [_celula(N1)]),
            ("Base", True, [_celula("Processo")], [_celula(ORIGEM1)])])
        leitura = listas.ler_arquivo(arq)
        self.assertEqual([n.formatado for n in leitura.processos], [N1])
        self.assertEqual(leitura.avisos, ["aba oculta 'Base' ignorada"])

    def test_frase_nao_manda_salvar_como_csv(self):
        # Achado R30: o .csv grava o que a célula mostra (notação científica,
        # ou o número com os últimos algarismos zerados): não recupera nada.
        arq = self.tmp / "relacao.xlsx"
        self._livro().save(arq)
        frase = self.recusada_pelos_corrompidos(arq)
        self.assertNotIn("salve a relação como .csv", frase)
        self.assertIn("digite-os de novo", frase)
        self.assertIn(".csv não adianta", frase)
        self.assertIn("com 2 linhas deste arquivo", frase)
        self.assertNotIn("(s)", frase)


class TestColunaDoNumero(_ComPasta):
    """Achado R26: a coluna "Processo" só com o CNJ corrompido fazia ler a
    linha inteira - e o lote virava o dos processos de origem."""

    def _salvar(self, *linhas):
        livro = openpyxl.Workbook()
        for valores in linhas:
            livro.active.append(list(valores))
        arq = self.tmp / "relacao.xlsx"
        livro.save(arq)
        return arq

    def test_processo_corrompido_nao_le_o_processo_de_origem(self):
        arq = self._salvar(["Processo", "Processo de origem", "Classe"],
                           [COMO_NUMERO1, ORIGEM1, "Cumprimento de sentença"],
                           [COMO_NUMERO2, ORIGEM2, "Cumprimento de sentença"])
        frase = self.recusada_pelos_corrompidos(arq)
        self.assertIn("com 2 linhas", frase)        # contados uma vez só

    def test_coluna_de_ordem_vazia_nao_le_o_processo_de_origem(self):
        # "Nº" é a primeira coluna com rótulo de número, mas traz a ordem
        # (1, 2...): lida a linha inteira, a de origem fica de fora.
        arq = self._salvar(["Nº", "Processo", "Processo de origem"],
                           [1, N1, ORIGEM1], [2, N2, ORIGEM2])
        self.assertEqual(self.numeros(arq), [N1, N2])

    def test_titulo_do_relatorio_nao_passa_por_cabecalho(self):
        # O título ("Relação de processos", numa célula só) também tem
        # "processo": tomado por cabeçalho, a coluna lida era a primeira - a
        # do processo de origem.
        arq = self._salvar(["Relação de processos - 1ª Vara Cível"], [],
                           ["Processo de origem", "Processo", "Classe"],
                           [ORIGEM1, N1, "Cumprimento de sentença"],
                           [ORIGEM2, N2, "Cumprimento de sentença"])
        self.assertEqual(self.numeros(arq), [N1, N2])

    def test_lista_de_uma_coluna_continua_lida(self):
        arq = self._salvar(["Processos conclusos"], ["Processo"], [N1], [N2])
        self.assertEqual(self.numeros(arq), [N1, N2])


class TestTextoUnicodeDoExcel(_ComPasta):
    """Achado R28: o "Salvar como > Texto Unicode" do Excel (.txt, UTF-16,
    tabulação) fazia da coluna vizinha a senha do processo sigiloso."""

    def _txt(self, conteudo):
        arq = self.tmp / "relacao.txt"
        arq.write_bytes(b"\xff\xfe" + conteudo.encode("utf-16-le"))
        return listas.ler_arquivo(arq)

    def test_coluna_vizinha_nao_vira_senha(self):
        leitura = self._txt(f"Processo\tSituação\r\n{N1}\tAtivo\r\n{N2}\tSuspenso\r\n")
        self.assertEqual([n.formatado for n in leitura.processos], [N1, N2])
        self.assertEqual(leitura.senhas, {})
        self.assertEqual(listas.ler_texto(f"{N1}\tAtivo\n{N2}\tSuspenso\n").senhas, {})

    def test_coluna_senha_continua_valendo(self):
        leitura = self._txt(f"Processo\tSenha\r\n{N1}\tk7Q2\r\n")
        self.assertEqual(leitura.senhas, {N1: "k7Q2"})


_NS_ODS = ('xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
           'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
           'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"')


def _celula(valor: str = "", repetida: int = 1, coberta: bool = False) -> str:
    tag = "table:covered-table-cell" if coberta else "table:table-cell"
    rep = f' table:number-columns-repeated="{repetida}"' if repetida > 1 else ""
    if not valor:
        return f"<{tag}{rep}/>"
    return f'<{tag}{rep} office:value-type="string"><text:p>{valor}</text:p></{tag}>'


def _ods(destino: Path, tabelas) -> Path:
    """tabelas: [(nome, oculta, linha 1, linha 2...)], cada linha uma lista
    de células já em XML."""
    corpo = []
    for nome, oculta, *linhas in tabelas:
        estilo = "taOculta" if oculta else "taVisivel"
        corpo.append(f'<table:table table:name="{nome}" table:style-name="{estilo}">'
                     + "".join("<table:table-row>" + "".join(l) + "</table:table-row>"
                               for l in linhas)
                     + "</table:table>")
    xml = (f'<?xml version="1.0"?><office:document-content {_NS_ODS}><office:automatic-styles>'
           '<style:style style:name="taVisivel" style:family="table"><style:table-properties '
           'table:display="true"/></style:style><style:style style:name="taOculta" '
           'style:family="table"><style:table-properties table:display="false"/></style:style>'
           "</office:automatic-styles><office:body><office:spreadsheet>" + "".join(corpo)
           + "</office:spreadsheet></office:body></office:document-content>")
    with zipfile.ZipFile(destino, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet")
        z.writestr("content.xml", xml)
    return destino


class TestOdsComCelulasRepetidas(_ComPasta):
    """Achado R29: o LibreOffice grava as células vazias seguidas como uma
    só (number-columns-repeated); contadas uma vez, deslocavam as colunas, e
    a "Senha" lia a coluna vizinha."""

    def test_vazias_repetidas_nao_deslocam_a_senha(self):
        arq = _ods(self.tmp / "relacao.ods", [("P", False,
            [_celula("Processo"), _celula("Classe"), _celula("Assunto"), _celula("Senha"),
             _celula("Observação"), _celula(repetida=1019)],
            [_celula(N1), _celula(repetida=2), _celula("k7Q2"), _celula("Urgente"),
             _celula(repetida=1019)])])
        self.assertEqual(listas.ler_arquivo(arq).senhas, {N1: "k7Q2"})

    def test_celula_coberta_pela_mescla_ocupa_o_lugar_dela(self):
        arq = _ods(self.tmp / "mescla.ods", [("P", False,
            [_celula("Processo"), _celula("Partes"), _celula(coberta=True), _celula("Senha")],
            [_celula(N1), _celula("Autor"), _celula("Réu"), _celula("XY99")])])
        self.assertEqual(listas.ler_arquivo(arq).senhas, {N1: "XY99"})


if __name__ == "__main__":
    unittest.main()
