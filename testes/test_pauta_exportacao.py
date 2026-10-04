"""Pauta: a planilha Excel - valores, formatos, cores, máscara dos sigilosos e abas."""

from __future__ import annotations

import unittest
from datetime import date, datetime, time

from helestron.pauta import exportacao, modelos
from helestron.pauta.servico import ServicoPauta

from testes import apoio_download as apoio
from testes import apoio_pauta as ap

AGORA = datetime(2026, 10, 5, 10, 0)
DE, ATE = ap.INICIO, ap.FIM


def _cor(celula_cor) -> str:
    return (getattr(celula_cor, "rgb", "") or "")[-6:].upper()


class TestExportacao(apoio.PastaTemporaria):
    def setUp(self):
        super().setUp()
        self.amb = ap.config_temporaria(self.tmp)
        self.servico = ServicoPauta(self.amb.cfg, self.tmp / "local" / "pauta.sqlite3",
                                    relogio=lambda: AGORA)
        self.addCleanup(self.servico.fechar)
        self.lista = ap.no_periodo(ap.audiencias_padrao(), DE, ATE)
        audiencias = []
        for i, f in enumerate(self.lista):
            audiencias.append(modelos.nova(
                sistema="eproc" if i == 3 else "esaj", tribunal="TJAL", data_=f.data,
                processo=f.processo, hora=f.hora, tipo_original=f.tipo,
                situacao_original=f.situacao, local=f.local, classe=f.classe, partes=f.partes,
                magistrado=f.magistrado, sigiloso=f.sigiloso, link=f.link))
        self.servico.armazem.gravar(audiencias, "esaj-tjal", None, registrar_novas=False,
                                    agora=datetime(2026, 10, 3, 9, 0))
        # sigiloso sem o portal dizer: os autos estão na pasta de sigilosos
        self.sigiloso_na_pasta = self.lista[0]
        lote = self.amb.sigilosos / "Lote de setembro"
        lote.mkdir(parents=True)
        (lote / f"{self.sigiloso_na_pasta.processo}.pdf").write_bytes(apoio.pdf_bytes(1))

    def abrir(self, arquivo):
        from openpyxl import load_workbook

        return load_workbook(arquivo)

    def linhas_por_processo(self, aba) -> dict[str, int]:
        return {aba.cell(row=r, column=4).value: r for r in range(5, aba.max_row + 1)
                if aba.cell(row=r, column=4).value}

    def test_aba_pauta(self):
        arquivo = self.servico.exportar(DE, ATE, self.amb.pauta)
        self.assertEqual(arquivo, self.amb.pauta / "Pauta de audiências 2026-10-05 a 2026-10-16.xlsx")
        livro = self.abrir(arquivo)
        self.assertEqual(livro.sheetnames, ["Pauta", "Resumo", "Alterações"])
        aba = livro["Pauta"]
        self.assertEqual(aba["A1"].value, "Pauta de audiências")
        self.assertEqual(aba["A2"].value, "De 05/10/2026 a 16/10/2026 · 8 audiências")
        self.assertEqual([c.value for c in aba[4]], [
            "Data", "Dia da semana", "Hora", "Processo", "Classe", "Partes", "Tipo de audiência",
            "Situação", "Local", "Magistrado/Conciliador", "Sistema", "Tribunal", "Link",
            "Observações"])
        cab = aba["A4"]
        self.assertTrue(cab.font.bold)
        self.assertEqual(_cor(cab.font.color), "FFFFFF")
        self.assertEqual(_cor(cab.fill.fgColor), exportacao.NAVY)
        self.assertEqual(aba.auto_filter.ref, "A4:N12")
        self.assertEqual(aba.freeze_panes, "E5")
        self.assertGreaterEqual(aba.column_dimensions["D"].width, 26)
        self.assertGreater(aba.column_dimensions["F"].width, aba.column_dimensions["C"].width)

        linhas = self.linhas_por_processo(aba)
        primeira = linhas[self.lista[1].processo]
        self.assertEqual(aba.cell(primeira, 1).value, datetime(2026, 10, 5))
        self.assertEqual(aba.cell(primeira, 1).number_format, "dd/mm/yyyy")
        self.assertEqual(aba.cell(primeira, 2).value, "Segunda-feira")
        self.assertEqual(aba.cell(primeira, 3).value, time(14, 30))
        self.assertEqual(aba.cell(primeira, 3).number_format, "hh:mm")
        self.assertEqual(aba.cell(primeira, 7).value, "Una")
        self.assertEqual(aba.cell(primeira, 11).value, "e-SAJ")
        self.assertEqual(aba.cell(primeira, 12).value, "TJAL")
        self.assertEqual(_cor(aba.cell(primeira, 1).fill.fgColor), exportacao.AZUL_HOJE,
                         "a audiência de hoje em azul-claro")

        # ordem: data e hora
        datas = [(aba.cell(r, 1).value, aba.cell(r, 3).value) for r in range(5, 13)]
        self.assertEqual(datas, sorted(datas))
        # sigilo: o portal disse, ou os autos estão na pasta de sigilosos
        for f in (self.lista[4], self.sigiloso_na_pasta):
            self.assertEqual(aba.cell(linhas[f.processo], 6).value, "(segredo de justiça)")
        # cancelada: cinza e tachada; redesignada: âmbar
        canc = linhas[self.lista[6].processo]
        self.assertTrue(aba.cell(canc, 1).font.strike)
        self.assertTrue(aba.cell(canc, 6).font.strike)
        self.assertEqual(_cor(aba.cell(canc, 8).font.color), exportacao.CINZA_TEXTO)
        redes = linhas[self.lista[2].processo]
        self.assertEqual(_cor(aba.cell(redes, 8).font.color), exportacao.AMBAR)
        self.assertFalse(aba.cell(redes, 1).font.strike)
        # zebra leve nas linhas que não são de hoje
        fundos = {_cor(aba.cell(r, 2).fill.fgColor) for r in range(5, 13)}
        self.assertIn(exportacao.CINZA_ZEBRA, fundos)
        # link clicável
        virtual = linhas[self.lista[5].processo]
        self.assertEqual(aba.cell(virtual, 13).hyperlink.target, self.lista[5].link)
        self.assertEqual(aba.cell(virtual, 13).value, self.lista[5].link)
        self.assertEqual(aba.cell(virtual, 11).value, "e-SAJ")
        self.assertEqual(aba.cell(linhas[self.lista[3].processo], 11).value, "eProc")
        # rodapé
        rodapes = [aba.cell(r, 1).value for r in range(13, aba.max_row + 1)]
        self.assertIn("Gerado pelo Helestron em 05/10/2026 10:00", rodapes)
        self.assertIn("Gerado pelo Helestron em 05/10/2026 10:00", aba.oddFooter.left.text)
        self.assertEqual(aba.page_setup.orientation, "landscape")

    def test_incluir_partes_dos_sigilosos(self):
        arquivo = self.servico.exportar(DE, ATE, self.amb.pauta, incluir_partes_sigilosos=True)
        aba = self.abrir(arquivo)["Pauta"]
        linhas = self.linhas_por_processo(aba)
        self.assertEqual(aba.cell(linhas[self.lista[4].processo], 6).value, "M. A. S. x J. R. S.")
        # e a configuração também liga (padrão: mascarar)
        self.amb.cfg.definir("pauta", "incluir_partes_sigilosos", True)
        aba = self.abrir(self.servico.exportar(DE, ATE, self.amb.pauta))["Pauta"]
        self.assertEqual(aba.cell(self.linhas_por_processo(aba)[self.lista[4].processo], 6).value,
                         "M. A. S. x J. R. S.")

    def test_resumo(self):
        livro = self.abrir(self.servico.exportar(DE, ATE, self.amb.pauta))
        r = livro["Resumo"]
        self.assertEqual([r.cell(4, c).value for c in (1, 2, 3)], ["Data", "Dia da semana",
                                                                    "Quantidade"])
        self.assertEqual((r.cell(5, 1).value, r.cell(5, 2).value, r.cell(5, 3).value),
                         (datetime(2026, 10, 5), "Segunda-feira", 2))
        self.assertEqual(r.cell(5, 1).number_format, "dd/mm/yyyy")
        por_dia = {r.cell(x, 1).value: r.cell(x, 3).value for x in range(5, 14)}
        self.assertEqual(por_dia["Total"], 8)
        por_tipo = {r.cell(x, 5).value: r.cell(x, 6).value for x in range(5, 13)
                    if r.cell(x, 5).value}
        self.assertEqual(por_tipo["Conciliação"], 2)
        self.assertEqual(por_tipo["Total"], 8)
        por_situacao = {r.cell(x, 8).value: r.cell(x, 9).value for x in range(5, 10)
                        if r.cell(x, 8).value}
        self.assertEqual(por_situacao, {"Designada": 6, "Cancelada": 1, "Redesignada": 1,
                                        "Total": 8})

    def test_alteracoes(self):
        f = self.lista[4]       # sigiloso: as partes que mudaram também são mascaradas
        mudada = modelos.nova(sistema="esaj", tribunal="TJAL", data_=f.data, processo=f.processo,
                              hora=f.hora, tipo_original=f.tipo, situacao_original="Cancelada",
                              partes="Outras Partes", local="Sala 9", sigiloso=True)
        self.servico.armazem.gravar([mudada], "esaj-tjal", None, agora=datetime(2026, 10, 4, 18, 30))
        livro = self.abrir(self.servico.exportar(DE, ATE, self.amb.pauta))
        h = livro["Alterações"]
        self.assertEqual([h.cell(4, c).value for c in range(1, 9)], [
            "Quando", "Alteração", "Processo", "Data da audiência", "Hora", "O que mudou",
            "Sistema", "Tribunal"])
        self.assertEqual(h.cell(5, 1).value, datetime(2026, 10, 4, 18, 30))
        self.assertEqual(h.cell(5, 1).number_format, "dd/mm/yyyy hh:mm")
        self.assertEqual(h.cell(5, 2).value, "Cancelada")
        self.assertEqual(h.cell(5, 3).value, f.processo)
        self.assertEqual(h.cell(5, 4).value, datetime(2026, 10, 8))
        self.assertEqual(h.cell(5, 5).value, time(9, 30))
        mudou = h.cell(5, 6).value
        self.assertIn("Situação: Designada → Cancelada", mudou)
        self.assertIn("Local: Sala de audiências da 2ª Vara Cível → Sala 9", mudou)
        self.assertIn("Partes: (segredo de justiça) → (segredo de justiça)", mudou)
        self.assertNotIn("Outras Partes", mudou)
        self.assertEqual(h.cell(5, 7).value, "e-SAJ")
        # sem alteração no período: a aba explica
        vazio = self.abrir(self.servico.exportar(date(2026, 11, 1), date(2026, 11, 2),
                                                 self.amb.pauta))
        self.assertEqual(vazio["Alterações"].cell(5, 1).value,
                         "Nenhuma alteração registrada para as audiências do período.")
        self.assertEqual(vazio["Pauta"].cell(5, 1).value, "Nenhuma audiência no período.")
        self.assertEqual(vazio["Pauta"]["A2"].value, "De 01/11/2026 a 02/11/2026 · 0 audiências")

    def test_filtros_e_nome_livre(self):
        a1 = self.servico.exportar(DE, ATE, self.amb.pauta, sistema="eproc")
        aba = self.abrir(a1)["Pauta"]
        self.assertEqual(list(self.linhas_por_processo(aba)), [self.lista[3].processo])
        self.assertIn("filtros: sistema eProc", aba["A2"].value)
        a2 = self.servico.exportar(DE, ATE, self.amb.pauta, situacao="cancelada")
        self.assertEqual(a2.name, "Pauta de audiências 2026-10-05 a 2026-10-16 (2).xlsx")
        aba = self.abrir(a2)["Pauta"]
        self.assertEqual(list(self.linhas_por_processo(aba)), [self.lista[6].processo])
        a3 = self.servico.exportar(DE, ATE, self.amb.pauta, busca="banco do brasil")
        self.assertEqual(a3.name, "Pauta de audiências 2026-10-05 a 2026-10-16 (3).xlsx")
        aba = self.abrir(a3)["Pauta"]
        self.assertEqual(list(self.linhas_por_processo(aba)), [self.lista[0].processo])
        self.assertIn("busca “banco do brasil”", aba["A2"].value)
        self.assertEqual(sorted(p.name for p in self.amb.pauta.iterdir()),
                         sorted([a1.name, a2.name, a3.name]), "nenhum temporário largado")

    def test_nunca_dentro_do_acervo(self):
        with self.assertRaises(ValueError) as erro:
            self.servico.exportar(DE, ATE, self.amb.acervo / "Pauta")
        self.assertIn("acervo", str(erro.exception))
        self.assertFalse((self.amb.acervo / "Pauta").exists())


if __name__ == "__main__":
    unittest.main()
