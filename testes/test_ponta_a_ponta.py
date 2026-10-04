"""Ponta a ponta DE VERDADE: o programa real e a interface real, sem demonstração.

Sobe ``python -m helestron --servidor --sem-janela`` com as pastas de dados
numa pasta temporária, abre a URL que ele imprime no Chromium (Playwright) e
percorre o que o usuário faz, pela interface:

* Início;
* Ajustes: grava um campo (vai para o config.ini), troca uma pasta, tenta uma
  pasta que poria os sigilosos dentro do acervo (recusada, com a frase) e
  guarda o acesso a um portal com senha;
* Processos: cola uma relação com números CNJ válidos do TJAL, do TJSP e do
  TRF4, revisa e baixa com o MOTOR REAL. Os portais são falsos e locais
  (enderecos-locais.json): o do TJAL recusa a conexão, o do TJSP nunca
  responde. A falha do TJAL tem de aparecer clara, a tela não pode travar
  enquanto o TJSP não responde, e Parar tem de interromper o lote;
* Audiências: sem microfone, a gravação não começa e a tela diz por quê;
* Pauta: importa um relatório (CSV gerado aqui) pelo seletor de arquivos do
  navegador, mostra as audiências e exporta o Excel, conferido com openpyxl;
* Compartilhar: prepara o acervo (vazio) para a IA;
* Ajuda: a busca acha o assunto.

Nenhum erro de console, exceção na página ou resposta 4xx/5xx além das
esperadas (o "sem diálogo nativo" 409, que faz a interface usar o seletor do
navegador, e a pasta recusada, 400). Nada vai à internet: o motor só fala
com os portais locais.

Pulado sem o Playwright ou sem o Chromium. As capturas vão para
HELESTRON_CAPTURAS/ponta-a-ponta (o CI as guarda) ou para uma pasta
temporária.
"""

from __future__ import annotations

import http.server
import importlib.util
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import tempfile
import threading
import time
import unittest
import urllib.request
from collections import Counter
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urlparse

RAIZ_REPO = Path(__file__).resolve().parents[1]
# Pedidos ao 127.0.0.1 nunca passam por proxy do sistema.
_SEM_PROXY = urllib.request.build_opener(urllib.request.ProxyHandler({}))
# Números CNJ com o dígito verificador certo, um por portal.
TJAL = "0700123-83.2024.8.02.0001"
TJSP = "1001234-20.2025.8.26.0100"
TRF4 = "5001234-46.2025.4.04.7100"
PAUTA = ("0700123-83.2024.8.02.0001", "0700456-35.2024.8.02.0001", "0700789-84.2024.8.02.0001")


class _PortalQueNaoResponde(http.server.BaseHTTPRequestHandler):
    """Aceita a conexão e não responde (até o teste acabar): o portal lento."""

    liberar = threading.Event()

    def do_GET(self):  # noqa: N802 - nome do http.server
        self.liberar.wait(180)

    do_POST = do_GET

    def log_message(self, *args):
        pass


def _porta_fechada() -> int:
    s = socket.socket()
    s.bind(("127.0.0.1", 0))
    porta = s.getsockname()[1]
    s.close()
    return porta


def _executavel_chromium() -> str | None:
    from testes.test_web_interface import _executavel_chromium as achar

    return achar()


def _navegador_do_motor(raiz: Path) -> dict | None:
    """O ambiente para o MOTOR abrir o Chromium do Playwright.

    No CI, o Chromium da versão do pacote está instalado e nada é preciso.
    Aqui (pacote mais novo que o Chromium da máquina), o caminho que o
    pacote espera aponta para o Chromium que existe. None: sem navegador
    nenhum para o motor.
    """
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        esperado = Path(p.chromium.executable_path)
    if esperado.exists():
        return {}
    exe = _executavel_chromium()
    if exe is None or "headless_shell" in exe:
        return None
    pasta = raiz / "navegadores"
    alvo = pasta / esperado.parent.parent.name / esperado.parent.name
    alvo.parent.mkdir(parents=True, exist_ok=True)
    try:
        alvo.symlink_to(Path(exe).parent, target_is_directory=True)
    except OSError:
        return None
    return {"PLAYWRIGHT_BROWSERS_PATH": str(pasta)}


class _Vigia:
    """Junta os problemas da página; respostas de erro esperadas são declaradas antes."""

    def __init__(self):
        self.problemas: list[str] = []
        self.permitidas: Counter = Counter()
        self.status_permitidos: Counter = Counter()

    def permitir(self, status: int, caminho: str, vezes: int = 1) -> None:
        self.permitidas[(status, caminho)] += vezes

    def ligar(self, pagina) -> None:
        pagina.on("console", self._console)
        pagina.on("pageerror", lambda e: self.problemas.append(f"exceção: {e}"))
        pagina.on("response", self._resposta)
        pagina.on("requestfailed", lambda r: self.problemas.append(
            f"falhou: {r.method} {r.url} ({r.failure})")
            if "/api/eventos" not in r.url else None)

    def _resposta(self, r) -> None:
        if r.status < 400:
            return
        chave = (r.status, urlparse(r.url).path)
        if self.permitidas[chave] > 0:
            self.permitidas[chave] -= 1
            self.status_permitidos[r.status] += 1
            return
        self.problemas.append(f"HTTP {r.status}: {r.request.method} {r.url}")

    def _console(self, m) -> None:
        if m.type not in ("error", "warning"):
            return
        achado = re.search(r"Failed to load resource: the server responded with a status of (\d+)",
                           m.text)
        if achado and self.status_permitidos[int(achado.group(1))] > 0:
            self.status_permitidos[int(achado.group(1))] -= 1
            return
        self.problemas.append(f"console.{m.type}: {m.text}")


class PontaAPonta(unittest.TestCase):
    """O percurso inteiro, em ordem, contra o programa real."""

    @classmethod
    def setUpClass(cls):
        if importlib.util.find_spec("playwright") is None:
            raise unittest.SkipTest("Playwright não instalado")
        cls.raiz = Path(tempfile.mkdtemp(prefix="helestron-ponta-"))
        cls.local, cls.dados = cls.raiz / "local", cls.raiz / "documentos"
        cls.local.mkdir()
        cls.dados.mkdir()
        try:
            cls._preparar()
        except BaseException:
            cls._desmontar()
            raise

    @classmethod
    def _preparar(cls):
        from playwright.sync_api import sync_playwright

        motor = _navegador_do_motor(cls.raiz)
        if motor is None:
            raise unittest.SkipTest("sem Chromium para o motor de download")

        # Portais falsos e locais: nada vai à internet.
        _PortalQueNaoResponde.liberar = threading.Event()
        cls.portal_lento = http.server.ThreadingHTTPServer(("127.0.0.1", 0), _PortalQueNaoResponde)
        cls.portal_lento.daemon_threads = True
        threading.Thread(target=cls.portal_lento.serve_forever, daemon=True).start()
        lento = f"http://127.0.0.1:{cls.portal_lento.server_address[1]}"
        fechado = f"http://127.0.0.1:{_porta_fechada()}"
        enderecos = {
            "esaj:TJAL": {"base": fechado},
            "eproc:TJAL": {g: f"{fechado}/eproc/" for g in ("1g", "2g")},
            "esaj:TJSP": {"base": lento},
            "eproc:TJSP": {g: f"{fechado}/eproc/" for g in ("1g", "2g")},
            "eproc:TRF4": {g: f"{lento}/eprocV2/" for g in ("1g_70", "1g_71", "1g_72", "2g")},
        }
        (cls.local / "enderecos-locais.json").write_text(json.dumps(enderecos), encoding="utf-8")
        # Espera curta por página e uma tentativa só; o navegador do motor à
        # vista (o Chromium completo, também no xvfb).
        (cls.local / "config.ini").write_text(
            "[download]\nespera_segundos = 10\npausa_entre_processos = 0\ntentativas = 1\n"
            "mostrar_navegador = true\n", encoding="utf-8")

        env = dict(os.environ)
        env.update(motor)
        env.update({"HELESTRON_LOCAL": str(cls.local), "HELESTRON_DADOS": str(cls.dados),
                    "PYTHONPATH": str(RAIZ_REPO), "PYTHONIOENCODING": "utf-8"})
        env.pop("LOCALAPPDATA", None)
        cls.servidor = subprocess.Popen(
            [sys.executable, "-m", "helestron", "--servidor", "--sem-janela"],
            cwd=str(cls.raiz), env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
            text=True, encoding="utf-8")
        cls.url = cls._ler_url()
        cls.base = cls.url.split("?")[0]
        cls.token = cls.url.split("t=", 1)[1]
        # o resto da saída do servidor não pode encher o cano
        threading.Thread(target=lambda: [None for _ in cls.servidor.stdout], daemon=True).start()

        cls.pw = sync_playwright().start()
        try:
            cls.chromium = cls.pw.chromium.launch()
        except Exception as erro:
            exe = _executavel_chromium()
            if "Executable doesn't exist" not in str(erro) or exe is None:
                raise unittest.SkipTest("Chromium do Playwright não instalado") from erro
            cls.chromium = cls.pw.chromium.launch(executable_path=exe)
        destino = os.environ.get("HELESTRON_CAPTURAS")
        cls.capturas = (Path(destino) / "ponta-a-ponta") if destino else cls.raiz / "capturas"
        cls.capturas.mkdir(parents=True, exist_ok=True)

    @classmethod
    def _ler_url(cls) -> str:
        resultado: list[str] = []

        def ler():
            for linha in cls.servidor.stdout:
                if linha.startswith("URL="):
                    resultado.append(linha.strip()[4:])
                    return

        leitor = threading.Thread(target=ler, daemon=True)
        leitor.start()
        leitor.join(90)
        if not resultado:
            raise AssertionError("o servidor não imprimiu URL= em 90 s")
        return resultado[0]

    @classmethod
    def tearDownClass(cls):
        cls._desmontar()

    @classmethod
    def _desmontar(cls):
        for nome in ("chromium", "pw"):
            objeto = getattr(cls, nome, None)
            try:
                if objeto is not None:
                    objeto.close() if nome == "chromium" else objeto.stop()
            except Exception:
                pass
        servidor = getattr(cls, "servidor", None)
        if servidor is not None and servidor.poll() is None:
            try:
                pedido = urllib.request.Request(cls.base + "api/encerrar", data=b"{}", method="POST",
                                                headers={"X-Helestron-Token": cls.token,
                                                         "Content-Type": "application/json"})
                _SEM_PROXY.open(pedido, timeout=10).read()
                servidor.wait(30)
            except Exception:
                servidor.terminate()
                try:
                    servidor.wait(10)
                except subprocess.TimeoutExpired:
                    servidor.kill()
        if getattr(cls, "portal_lento", None) is not None:
            _PortalQueNaoResponde.liberar.set()
            cls.portal_lento.shutdown()
            cls.portal_lento.server_close()
        if servidor is not None and servidor.stdout is not None:
            servidor.stdout.close()
        shutil.rmtree(cls.raiz, ignore_errors=True)

    # ------------------------------------------------------------- apoio
    def esperar_secao(self, secao: str) -> None:
        self.pg.wait_for_function(
            "s => document.documentElement.dataset.secao === s && "
            "document.documentElement.dataset.pronta === '1'", arg=secao, timeout=30000)

    def ir(self, rota: str) -> None:
        self.pg.evaluate("r => { location.hash = '#/' + r; }", rota)
        self.esperar_secao(rota.split("/")[0])

    def capturar(self, nome: str) -> None:
        self.pg.wait_for_timeout(300)
        self.pg.screenshot(path=str(self.capturas / f"{nome}.png"))

    def ini(self) -> str:
        return (self.local / "config.ini").read_text(encoding="utf-8")

    def api(self, caminho: str):
        pedido = urllib.request.Request(self.base + caminho.lstrip("/"),
                                        headers={"X-Helestron-Token": self.token})
        with _SEM_PROXY.open(pedido, timeout=20) as r:
            return json.loads(r.read().decode("utf-8"))["dados"]

    # ------------------------------------------------------------- o percurso
    def test_percurso_completo_pela_interface(self):
        contexto = self.chromium.new_context(viewport={"width": 1280, "height": 820},
                                             locale="pt-BR")
        self.addCleanup(contexto.close)
        self.pg = contexto.new_page()
        self.vigia = _Vigia()
        self.vigia.ligar(self.pg)
        self.pg.goto(self.url)
        self.esperar_secao("inicio")

        with self.subTest(etapa="início"):
            self.etapa_inicio()
        with self.subTest(etapa="ajustes"):
            self.etapa_ajustes()
        with self.subTest(etapa="processos"):
            self.etapa_processos()
        with self.subTest(etapa="audiências"):
            self.etapa_audiencias()
        with self.subTest(etapa="pauta"):
            self.etapa_pauta()
        with self.subTest(etapa="compartilhar"):
            self.etapa_compartilhar()
        with self.subTest(etapa="ajuda"):
            self.etapa_ajuda()
        with self.subTest(etapa="diagnóstico"):
            self.etapa_diagnostico()
        self.assertEqual(self.vigia.problemas, [],
                         "erros no navegador:\n" + "\n".join(self.vigia.problemas))
        self.assertIsNone(self.servidor.poll(), "o servidor caiu durante o percurso")

    def etapa_inicio(self):
        pg = self.pg
        self.assertRegex(pg.locator(".titulo-grande").inner_text(), r"^(Bom dia|Boa tarde|Boa noite)")
        self.assertEqual(pg.locator(".cartao-funcao").count(), 4)
        self.assertIn("Primeiros passos", pg.locator("#pagina").inner_text())
        self.capturar("01-inicio")

    def etapa_ajustes(self):
        pg = self.pg
        # Um campo de texto: salvo na hora, no config.ini.
        self.ir("ajustes/unidade")
        pg.fill("#cfg-unidade-magistrado", "Camila Duarte Albuquerque")
        pg.press("#cfg-unidade-magistrado", "Enter")
        pg.wait_for_selector(".valor-salvo")
        self.assertIn("magistrado = Camila Duarte Albuquerque", self.ini())

        # Pastas: sem diálogo nativo (modo servidor), a folha pede o caminho.
        self.ir("ajustes/pastas")
        pg.wait_for_selector("text=Pasta dos processos sigilosos")
        planilhas = self.raiz / "Planilhas da pauta"
        self.vigia.permitir(409, "/api/dialogo/pasta")
        pg.click(".linha:has-text('Pasta da pauta exportada') button:has-text('Alterar')")
        pg.fill(".folha input", str(planilhas))
        pg.click(".folha button:has-text('Usar esta pasta')")
        pg.wait_for_selector(".aviso:has-text('Pasta alterada')")
        self.assertIn(f"pasta = {planilhas}", self.ini())

        # Sigilosos dentro do acervo: recusado, com a frase, e nada gravado.
        antes = self.ini()
        self.vigia.permitir(409, "/api/dialogo/pasta")
        self.vigia.permitir(400, "/api/config")
        pg.click(".linha:has-text('Pasta dos processos sigilosos') button:has-text('Alterar')")
        pg.fill(".folha input", str(self.dados / "Acervo" / "Sigilosos"))
        pg.click(".folha button:has-text('Usar esta pasta')")
        recusa = pg.locator(".folha:has-text('Não deu certo')")
        recusa.wait_for()
        self.assertIn("compartilhado com a IA", recusa.inner_text())
        self.capturar("02-ajustes-pasta-recusada")
        recusa.locator("button:has-text('OK')").click()
        pg.wait_for_selector(".folha", state="detached")
        self.assertEqual(self.ini(), antes)

        # Acesso a um portal, com senha guardada.
        self.ir("ajustes/acessos")
        pg.click("#adicionar-acesso")
        pg.select_option("#acesso-portal", "esaj:TJAL")
        pg.fill("#acesso-usuario", "camila.albuquerque")
        pg.fill("#acesso-senha", "s3nh@ de teste")
        pg.click(".folha button:has-text('Salvar')")
        pg.wait_for_selector("[data-portal='esaj:TJAL']:has-text('Senha guardada')")
        acessos = {a["portal"]: a for a in self.api("/api/acessos")}
        self.assertTrue(acessos["esaj:TJAL"]["tem_senha"])
        self.assertTrue((self.local / "credenciais.json").is_file())
        self.capturar("03-ajustes-acessos")

        # Endereço do portal: as correções (aqui, os portais falsos deste
        # teste) aparecem; a folha mostra o endereço de cada seção do TRF4.
        grupo = pg.locator("#grupo-enderecos")
        grupo.locator("text=TJSP · e-SAJ — Endereço do portal").wait_for()
        pg.click("#corrigir-endereco")
        pg.select_option("#endereco-portal", "eproc:TRF4")
        pg.wait_for_selector("label:has-text('1º grau — Rio Grande do Sul')")
        valores = pg.locator(".folha input[data-grau]").evaluate_all("es => es.map(e => e.value)")
        self.assertTrue(valores and all(v.startswith("http://127.0.0.1:") for v in valores), valores)
        self.capturar("03b-ajustes-endereco-do-portal")
        pg.click(".folha button:has-text('Cancelar')")
        pg.wait_for_selector(".folha", state="detached")

    def etapa_processos(self):
        pg = self.pg
        self.ir("processos")
        pg.click("#area-soltar button:has-text('Colar lista')")
        pg.fill(".folha textarea", f"Relação da semana\n{TJAL}\n{TJSP}; {TRF4}\n")
        pg.click(".folha button:has-text('Ler a lista')")
        pg.wait_for_selector(".revisao")
        self.assertEqual(pg.locator(".revisao tbody tr").count(), 3)
        texto = pg.locator(".revisao").inner_text()
        for sigla in ("TJAL", "TJSP", "TRF4"):
            self.assertIn(sigla, texto)
        self.capturar("04-processos-revisao")

        pg.click("#botao-baixar")
        pg.wait_for_selector(".andamento")
        # O TJAL recusa a conexão: a falha aparece na linha dele, com a causa.
        linha_tjal = pg.locator(f".itens-lote tr[data-numero='{TJAL}']")
        linha_tjal.locator(".pilula:has-text('Falhou')").wait_for(timeout=60000)
        self.assertIn("recusou a conexão", linha_tjal.inner_text())
        # O TJSP não responde: a tela continua viva, com o lote em andamento.
        pg.wait_for_function(
            "() => /Entrando no e-SAJ do TJSP|Abrindo o e-SAJ do TJSP/.test("
            "document.querySelector('.andamento-status').textContent)", timeout=60000)
        self.capturar("05-processos-andamento")
        self.assertTrue(pg.locator(".andamento button:has-text('Parar')").is_enabled())
        # Parar interrompe o lote (o motor desiste no próximo ponto de parada).
        inicio = time.monotonic()
        pg.click(".andamento button:has-text('Parar')")
        pg.click(".folha button:has-text('Parar')")
        pg.wait_for_selector(".lote-concluido", timeout=90000)
        self.assertLess(time.monotonic() - inicio, 90)
        self.assertIn("Lote interrompido", pg.locator(".andamento").inner_text())
        self.capturar("06-processos-parado")
        tarefas = self.api("/api/tarefas")
        self.assertFalse([t for t in tarefas if t["estado"] == "rodando"], tarefas)

    def etapa_audiencias(self):
        pg = self.pg
        self.ir("audiencias")
        pg.fill("#processo-audiencia", TJAL.replace("-", "").replace(".", ""))
        self.assertIn("Número válido", pg.locator("#ajuda-processo").inner_text())
        if self.api("/api/transcricao/microfones"):
            self.skipTest("há microfone nesta máquina: o caso sem microfone não se aplica")
        pg.wait_for_selector("text=Nenhum microfone encontrado")
        pg.click("#botao-gravar")
        aviso = pg.locator(".folha:has-text('A gravação não começou')")
        aviso.wait_for(timeout=30000)
        self.assertIn("microfone", aviso.inner_text().lower())
        self.capturar("07-audiencias-sem-microfone")
        aviso.locator("button:has-text('OK')").click()
        pg.wait_for_selector(".folha", state="detached")
        # De volta à preparação, pronta para tentar de novo.
        pg.wait_for_selector("#botao-gravar")
        self.assertEqual(self.api("/api/transcricao/estado")["estado"], "erro")

    def etapa_pauta(self):
        pg = self.pg
        hoje = date.today()
        relatorio = self.raiz / "Relatório da pauta.csv"
        linhas = ["Data;Hora;Processo;Tipo de audiência;Situação;Local;Partes"]
        tipos = ("Conciliação", "Instrução e julgamento", "Una")
        situacoes = ("Designada", "Designada", "Cancelada")
        for i, numero in enumerate(PAUTA):
            dia = hoje + timedelta(days=i)
            linhas.append(f"{dia:%d/%m/%Y};{9 + i:02d}:30;{numero};{tipos[i]};{situacoes[i]};"
                          f"Sala {i + 1};Autor {i + 1} x Réu {i + 1}")
        relatorio.write_text("\n".join(linhas) + "\n", encoding="utf-8")

        self.ir("pauta")
        pg.wait_for_selector(".pauta-vazia-cartao")
        self.vigia.permitir(409, "/api/dialogo/arquivo")
        with pg.expect_file_chooser() as escolha:
            pg.click(".pauta-vazia-cartao button:has-text('Importar relatório')")
        escolha.value.set_files(str(relatorio))
        pg.wait_for_selector("#resultado-pauta .faixa", timeout=30000)
        resultado = pg.locator("#resultado-pauta").inner_text()
        self.assertIn("Relatório da pauta.csv", resultado)
        self.assertIn("3 audiências lidas", resultado)
        pg.wait_for_function("() => document.querySelectorAll('.audiencia').length === 3")
        lista = pg.locator("#lista-pauta").inner_text()
        for numero in PAUTA:
            self.assertIn(numero, lista)
        # Os números do resumo seguem a importação: hoje 1; nos próximos 7
        # dias 2 (a cancelada não conta); canceladas ou redesignadas 1.
        pg.wait_for_function(
            "() => ['chip-total', 'chip-hoje', 'chip-semana', 'chip-mudancas'].map("
            "id => document.querySelector('#' + id + ' strong').textContent).join() === '3,1,2,1'",
            timeout=15000)
        self.capturar("08-pauta-importada")

        pg.click("#botao-exportar")
        pg.wait_for_selector(".folha:has-text('Exportar a pauta para o Excel')")
        pg.click(".folha button:has-text('Exportar')")
        pg.wait_for_selector(".aviso:has-text('Planilha pronta')", timeout=30000)
        self.capturar("09-pauta-exportada")
        planilhas = list((self.raiz / "Planilhas da pauta").glob("*.xlsx"))
        self.assertEqual(len(planilhas), 1, planilhas)
        import openpyxl

        livro = openpyxl.load_workbook(planilhas[0])
        self.assertEqual(livro.sheetnames, ["Pauta", "Resumo", "Alterações"])
        valores = [str(c.value) for linha in livro["Pauta"].iter_rows() for c in linha
                   if c.value is not None]
        for numero in PAUTA:
            self.assertTrue(any(numero in v for v in valores), numero)
        self.assertIn("Autor 1 x Réu 1", valores)
        livro.close()

    def etapa_compartilhar(self):
        pg = self.pg
        self.ir("compartilhar")
        pg.click("#botao-preparar")
        pg.wait_for_selector(".aviso:has-text('0 processos, 0 transcrições')", timeout=60000)
        self.capturar("10-compartilhar")
        acervo = self.dados / "Acervo"
        self.assertTrue((acervo / "CLAUDE.md").is_file())
        self.assertTrue((acervo / "AGENTS.md").is_file())

    def etapa_ajuda(self):
        pg = self.pg
        self.ir("ajuda")
        pg.fill("#busca-ajuda", "pauta")
        pg.wait_for_function("() => document.querySelectorAll('mark').length > 0")
        self.capturar("11-ajuda")

    def etapa_diagnostico(self):
        """Ajustes › Sobre e diagnóstico: a verificação de verdade, legível."""
        pg = self.pg
        self.ir("ajustes/sobre")
        grupo = pg.locator("section[aria-label='Verificação da instalação']")
        grupo.wait_for(timeout=60000)
        nomes = grupo.locator(".linha-titulo").all_inner_texts()
        for esperado in ("Bibliotecas", "Pastas de trabalho", "Conector do acervo (MCP)"):
            self.assertIn(esperado, nomes)
        self.assertEqual(nomes.count("WebView2 (janela do programa)"), 1)
        # O "o que fazer" fica embaixo do detalhe; o título não é espremido.
        larguras = grupo.locator(".linha-titulo").evaluate_all(
            "ts => ts.map(t => t.getBoundingClientRect().width)")
        self.assertGreater(min(larguras), 150, larguras)
        self.assertGreater(grupo.locator(".verificacao-acao").count(), 0)
        self.capturar("12-diagnostico")


if __name__ == "__main__":
    unittest.main()
