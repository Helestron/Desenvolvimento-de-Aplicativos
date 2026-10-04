"""O documento da transcrição (DOCX), as falas e a ficha da audiência.

Porte do `saida.py` do Assessor SAJ, com o que faltava para uso processual:
número do processo no título e na ficha, tipo, data, horário, unidade,
magistrado(a), participantes e a forma da transcrição (ao vivo, revisão ou
gravação). O resto vem da base, que já tinha sido afinado com o uso:
Arial 11, ficha em tabela, aviso de conferência humana, legenda para
identificar os rótulos automáticos e numeração de páginas pelo Word.

A gravação é atômica (arquivo temporário + troca), e o documento aberto no
Word não derruba nada: o Word tranca o arquivo, então gravamos ao lado, em
"<nome> (cópia).docx", e devolvemos esse caminho.
"""

from __future__ import annotations

import io
import logging
import re
import time
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path

from ..nucleo import sigilo, sistema

log = logging.getLogger("transcricao.documento")

SUBPASTA_TRANSCRICOES = sigilo.SUBPASTA_TRANSCRICOES   # dentro da pasta de sigilosos
PAUSA_AGRUPAR_S = 3.0
RE_AUTOMATICO = re.compile(r"^FALANTE \d+$")

AVISO_AUTOMATICO = (
    "Transcrição produzida automaticamente por reconhecimento de fala. Os rótulos "
    "FALANTE 1, FALANTE 2 etc. são atribuídos por similaridade de voz e não "
    "identificam nominalmente os participantes. O conteúdo exige conferência humana "
    "contra a gravação original antes de qualquer uso processual."
)
AVISO_MANUAL = (
    "Transcrição produzida automaticamente por reconhecimento de fala. A indicação de "
    "quem fala foi marcada durante a audiência e pode falhar nos instantes de troca. "
    "O conteúdo exige conferência humana contra a gravação original antes de qualquer "
    "uso processual."
)
# Sem rótulo nenhum (gravação transcrita sem a separação de vozes, ou
# ninguém marcou quem falava): dizer que a indicação "foi marcada durante a
# audiência" seria falso num documento que pode ir para os autos.
AVISO_SEM_FALANTE = (
    "Transcrição produzida automaticamente por reconhecimento de fala, sem indicação "
    "de quem fala. O conteúdo exige conferência humana contra a gravação original "
    "antes de qualquer uso processual."
)

FORMAS = {
    "ao vivo": "Simultânea (ao vivo, pelo microfone)",
    "revisão": "Revisão da gravação, depois da audiência",
    "gravação": "A partir de arquivo de gravação",
}


@dataclass
class Fala:
    """Uma fala: tempos em segundos desde o início da gravação."""

    inicio: float
    fim: float
    falante: str
    texto: str
    # palavras com tempo [(início, fim, texto)], quando o modelo as deu:
    # permitem dividir a fala na troca de voz (separação de falantes)
    palavras: list | None = field(default=None, repr=False, compare=False)

    def como_dict(self) -> dict:
        return {"inicio": round(float(self.inicio), 2), "fim": round(float(self.fim), 2),
                "falante": self.falante or "", "texto": self.texto}

    @classmethod
    def de_dict(cls, d: dict) -> "Fala":
        return cls(float(d.get("inicio", 0.0)), float(d.get("fim", 0.0)),
                   str(d.get("falante") or ""), str(d.get("texto") or ""))


@dataclass
class MetaAudiencia:
    """O que vai na ficha do documento."""

    numero: str
    tipo: str = ""
    data: datetime | None = None
    unidade: str = ""
    magistrado: str = ""
    participantes: dict[str, str] = field(default_factory=dict)
    inicio: datetime | None = None
    fim: datetime | None = None
    modelo: str = ""
    origem: str = "ao vivo"          # "ao vivo" | "revisão" | "gravação"
    gravacao: str = ""               # nome do arquivo de áudio
    duracao: float = 0.0             # segundos de áudio gravado (sem as pausas)
    observacao: str = ""

    def como_dict(self) -> dict:
        d = asdict(self)
        for chave in ("data", "inicio", "fim"):
            valor = getattr(self, chave)
            d[chave] = valor.isoformat(timespec="seconds") if valor else None
        return d

    @classmethod
    def de_dict(cls, d: dict) -> "MetaAudiencia":
        dados = dict(d or {})

        def data(valor):
            if not valor:
                return None
            try:
                return datetime.fromisoformat(str(valor))
            except ValueError:
                return None

        conhecidos = {f for f in cls.__dataclass_fields__}  # type: ignore[attr-defined]
        dados = {k: v for k, v in dados.items() if k in conhecidos}
        for chave in ("data", "inicio", "fim"):
            dados[chave] = data(dados.get(chave))
        dados["participantes"] = dict(dados.get("participantes") or {})
        dados["numero"] = str(dados.get("numero") or "")
        try:
            dados["duracao"] = float(dados.get("duracao") or 0.0)
        except (TypeError, ValueError):
            dados["duracao"] = 0.0
        return cls(**dados)


# ------------------------------------------------------------- utilidades
def hms(segundos: float) -> str:
    s = max(0, int(segundos))
    return f"{s // 3600:02d}:{(s % 3600) // 60:02d}:{s % 60:02d}"


def automatico(rotulo: str) -> bool:
    """O rótulo foi dado pela separação automática (FALANTE n)?"""
    return bool(RE_AUTOMATICO.match((rotulo or "").strip()))


def e_marca(fala: Fala) -> bool:
    """Marcação da gravação (pausa), não fala de ninguém: sem falante e sem
    duração. Fala sem rótulo (gravação sem separação de vozes, ou ninguém
    marcou quem falava) TEM duração e é fala normal."""
    return not (fala.falante or "").strip() and (fala.fim - fala.inicio) < 0.01


def agrupar(falas: list[Fala], pausa: float = PAUSA_AGRUPAR_S) -> list[Fala]:
    """Une falas consecutivas do mesmo falante numa só.

    Não emenda por cima de uma pausa longa (provavelmente é outro assunto,
    ou outra pergunta), nem marcações sem falante (pausa da gravação).
    """
    saida: list[Fala] = []
    for f in falas:
        texto = (f.texto or "").strip()
        if not texto:
            continue
        ultimo = saida[-1] if saida else None
        if (ultimo is not None and f.falante and ultimo.falante == f.falante
                and (f.inicio - ultimo.fim) < pausa):
            ultimo.texto = f"{ultimo.texto} {texto}".strip()
            ultimo.fim = max(ultimo.fim, f.fim)
        else:
            saida.append(Fala(f.inicio, f.fim, f.falante, texto))
    return saida


def unidade_da_config(cfg) -> str:
    partes = [cfg.texto("unidade", c) for c in ("vara", "comarca", "tribunal")]
    return ", ".join(p for p in partes if p)


def magistrado_da_config(cfg) -> str:
    nome = cfg.texto("unidade", "magistrado")
    cargo = cfg.texto("unidade", "cargo")
    if nome and cargo:
        return f"{nome} ({cargo})"
    return nome or cargo


def processo_sigiloso(cfg, numero) -> bool:
    """O processo já está na pasta de sigilosos?

    Os autos (um PDF com o número em <sigilosos>/ ou <sigilosos>/<lote>/, o
    critério do compartilhamento) ou uma transcrição, gravação ou diário dele
    em <sigilosos>/Transcricoes: a audiência que já foi sigilosa uma vez
    continua sigilosa - na retranscrição, no envio da gravação pela página
    (que perde a pasta de origem) e numa nova audiência do mesmo processo. O
    que vale para os autos vale para a transcrição, a gravação e o diário.
    (É a parte "pasta" da regra única do sigilo, nucleo/sigilo.py; a pauta é
    a outra.)
    """
    try:
        pasta = cfg.pasta_sigilosos
    except Exception:
        return False
    return sigilo.na_pasta(pasta, numero)


def pasta_das_transcricoes(cfg, numero=None, sigiloso: bool = False) -> Path:
    """Onde vão o DOCX (e, em _audio, a gravação e o diário) do processo.

    Processo em segredo de justiça - informado (`sigiloso`) ou que o
    programa já sabe sigiloso pela regra única (autos, transcrição ou
    gravação na pasta de sigilosos, ou a pauta de audiências) - fica em
    <sigilosos>/Transcricoes, FORA do acervo: tudo o que está no acervo é
    lido pela IA (Cowork, Claude Code, ChatGPT) e espelhado na nuvem. A
    informação só acrescenta sigilo, nunca o tira.
    """
    if sigiloso or (numero is not None and sigilo.processo_sigiloso(cfg, numero)):
        return Path(cfg.pasta_sigilosos) / SUBPASTA_TRANSCRICOES
    return cfg.pasta_transcricoes


def pastas_das_transcricoes(cfg) -> list[Path]:
    """As duas pastas de transcrições: a do acervo e a dos sigilosos."""
    pastas = [cfg.pasta_transcricoes]
    try:
        sigilosas = Path(cfg.pasta_sigilosos) / SUBPASTA_TRANSCRICOES
    except (AttributeError, OSError, ValueError):
        return pastas
    if sigilosas not in pastas:
        pastas.append(sigilosas)
    return pastas


# ------------------------------------------------------------------- DOCX
def _sombrear(celula, cor: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), cor)
    celula._tc.get_or_add_tcPr().append(shd)


def _campo(paragrafo, instrucao: str):
    """Campo do Word (PAGE, NUMPAGES): o Word calcula ao abrir/imprimir."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragrafo.add_run()
    for tipo, texto in (("begin", None), (None, instrucao), ("separate", None), ("end", None)):
        if tipo:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipo)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instrucao
        run._r.append(el)
    return run


def _larguras(tabela, larguras_cm: tuple[float, ...]) -> None:
    """Largura em TODAS as células: só na primeira, o LibreOffice e o Word
    para a web ignoram (detalhe apontado na análise da base)."""
    from docx.shared import Cm

    tabela.autofit = False
    for linha in tabela.rows:
        for celula, largura in zip(linha.cells, larguras_cm):
            celula.width = Cm(largura)


def _linhas_da_ficha(meta: MetaAudiencia) -> list[tuple[str, str]]:
    traco = "—"
    data = meta.data or meta.inicio
    horario = traco
    if meta.inicio and meta.fim:
        horario = f"{meta.inicio:%H:%M} às {meta.fim:%H:%M}"
    elif meta.inicio:
        horario = f"início às {meta.inicio:%H:%M}"
    participantes = "\n".join(
        f"{papel}: {nome}" if nome else papel
        for papel, nome in (meta.participantes or {}).items() if papel or nome)
    linhas = [
        ("Processo nº", meta.numero or traco),
        ("Tipo de audiência", meta.tipo or traco),
        ("Data", f"{data:%d/%m/%Y}" if data else traco),
        ("Início e término", horario),
        ("Duração da gravação", hms(meta.duracao) if meta.duracao else traco),
        ("Unidade", meta.unidade or traco),
        ("Magistrado(a)", meta.magistrado or traco),
        ("Participantes", participantes or traco),
        ("Forma da transcrição", FORMAS.get(meta.origem, meta.origem or traco)),
        ("Modelo de transcrição", f"Whisper {meta.modelo} (português)" if meta.modelo else traco),
        ("Gravação", meta.gravacao or traco),
    ]
    if meta.observacao:
        linhas.append(("Observação", meta.observacao))
    return linhas


def montar_docx(falas: list[Fala], meta: MetaAudiencia, *, marcar_tempo: bool = True,
                legenda_automatica: bool = False) -> bytes:
    """O DOCX em memória (para gravar de uma vez)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    agrupadas = agrupar(falas)
    if legenda_automatica:
        rotulos = [f.falante for f in agrupadas if f.falante]
    else:
        rotulos = [f.falante for f in agrupadas if automatico(f.falante)]
    legenda = list(dict.fromkeys(rotulos))  # ordem de aparição, sem repetir
    if legenda and all(automatico(r) for r in legenda):
        legenda.sort(key=lambda r: int(r.split()[-1]))

    doc = Document()
    estilo = doc.styles["Normal"]
    estilo.font.name = "Arial"
    estilo.font.size = Pt(11)
    pf = estilo.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    try:  # sem isto, o Word usa outra fonte nos caracteres "asiáticos" (aspas, travessão)
        from docx.oxml.ns import qn

        rpr = estilo.element.get_or_add_rPr()
        fontes = rpr.find(qn("w:rFonts"))
        if fontes is not None:
            for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
                fontes.set(qn(atributo), "Arial")
    except Exception:  # pragma: no cover - cosmético
        pass

    for secao in doc.sections:
        secao.left_margin = secao.right_margin = Cm(2.5)
        secao.top_margin = secao.bottom_margin = Cm(2.0)

    # --- título
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("TRANSCRIÇÃO DE AUDIÊNCIA")
    r.bold = True
    r.font.size = Pt(14)
    if meta.numero:
        st = doc.add_paragraph()
        st.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rs = st.add_run(f"Processo nº {meta.numero}")
        rs.italic = True
        rs.font.size = Pt(10)
        rs.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    # --- ficha
    tab = doc.add_table(rows=0, cols=2)
    tab.style = "Table Grid"
    for rotulo, valor in _linhas_da_ficha(meta):
        celulas = tab.add_row().cells
        p = celulas[0].paragraphs[0]
        rr = p.add_run(rotulo)
        rr.bold = True
        rr.font.size = Pt(9)
        _sombrear(celulas[0], "F2F2F2")
        linhas_valor = str(valor).split("\n")
        pv = celulas[1].paragraphs[0]
        rv = pv.add_run(linhas_valor[0])
        rv.font.size = Pt(9)
        for extra in linhas_valor[1:]:
            rv.add_break()
            rv = pv.add_run(extra)
            rv.font.size = Pt(9)
    _larguras(tab, (5.0, 11.0))
    doc.add_paragraph()

    # --- aviso
    av = doc.add_paragraph()
    if legenda:
        aviso = AVISO_AUTOMATICO
    elif any((f.falante or "").strip() for f in agrupadas):
        aviso = AVISO_MANUAL
    else:
        aviso = AVISO_SEM_FALANTE
    ra = av.add_run(aviso)
    ra.italic = True
    ra.font.size = Pt(9)
    ra.font.color.rgb = RGBColor(0x80, 0x30, 0x00)

    # --- legenda dos rótulos automáticos (o usuário preenche os nomes)
    if legenda:
        doc.add_paragraph()
        h = doc.add_paragraph()
        rh = h.add_run("IDENTIFICAÇÃO DOS FALANTES")
        rh.bold = True
        rh.font.size = Pt(11)
        instr = doc.add_paragraph().add_run(
            "Preencha a coluna da direita e use Localizar e Substituir (Ctrl+U) "
            "para trocar os rótulos no corpo da transcrição.")
        instr.font.size = Pt(9)
        leg = doc.add_table(rows=1, cols=2)
        leg.style = "Table Grid"
        for i, texto in enumerate(("Rótulo", "Quem é (preencher)")):
            p = leg.rows[0].cells[i].paragraphs[0]
            rc = p.add_run(texto)
            rc.bold = True
            rc.font.size = Pt(9)
            _sombrear(leg.rows[0].cells[i], "E8E8E8")
        for rotulo in legenda:
            c = leg.add_row().cells
            c[0].paragraphs[0].add_run(rotulo).font.size = Pt(9)
            c[1].paragraphs[0].add_run("").font.size = Pt(9)
        _larguras(leg, (4.0, 12.0))

    doc.add_page_break()

    # --- corpo
    hb = doc.add_paragraph()
    rb = hb.add_run("TRANSCRIÇÃO")
    rb.bold = True
    rb.font.size = Pt(12)

    if not agrupadas:
        vazio = doc.add_paragraph().add_run("(Nenhuma fala transcrita até o momento.)")
        vazio.italic = True
    for fala in agrupadas:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        if fala.falante:
            p.add_run(fala.falante).bold = True
        if marcar_tempo:
            rt = p.add_run(f" [{hms(fala.inicio)}]" if fala.falante else f"[{hms(fala.inicio)}]")
            rt.font.size = Pt(9)
            rt.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if fala.falante or marcar_tempo:
            p.add_run(" — ").bold = True
        corpo = p.add_run(fala.texto)
        # Só a marcação da gravação (pausa) sai em itálico. Antes, qualquer
        # fala sem rótulo saía assim: a transcrição inteira de uma gravação
        # sem separação de vozes ficava em itálico, como se fosse anotação.
        if e_marca(fala):
            corpo.italic = True

    # --- rodapé: processo e "página N de M"
    rodape = doc.sections[0].footer.paragraphs[0]
    rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefixo = f"Processo nº {meta.numero} · " if meta.numero else ""
    rodape.add_run(f"{prefixo}página ")
    _campo(rodape, "PAGE")
    rodape.add_run(" de ")
    _campo(rodape, "NUMPAGES")
    for run in rodape.runs:
        run.font.size = Pt(8)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


ESPERAS_TRAVA_S = (0.2, 0.5)   # novas tentativas antes de desistir do nome


def _gravar(destino: Path, dados: bytes) -> None:
    # Antivírus corporativo, indexador e sincronizadores abrem o arquivo por
    # um instante logo depois de gravado; a troca (os.replace) falha com
    # PermissionError nesse intervalo. Sem repetir, o salvamento automático
    # caía na "(cópia)" sem o Word estar aberto, e o documento final podia
    # sair com outro nome que não o do processo.
    for espera in (*ESPERAS_TRAVA_S, None):
        try:
            sistema.gravar_atomico(destino, dados)
            return
        except PermissionError:
            _apagar_parcial(destino)
            if espera is None:
                raise
            time.sleep(espera)
        except BaseException:
            _apagar_parcial(destino)
            raise


def _apagar_parcial(destino: Path) -> None:
    # O próprio .parcial pode estar preso pelo antivírus: não apagá-lo não
    # pode esconder o erro original (nem fingir uma trava que não existe).
    try:
        destino.with_name(destino.name + ".parcial").unlink(missing_ok=True)
    except OSError:
        pass


def gerar_docx(destino: Path, falas: list[Fala], meta: MetaAudiencia, *,
               marcar_tempo: bool = True, legenda_automatica: bool = False) -> Path:
    """Grava o DOCX e devolve o caminho efetivamente gravado.

    Se o destino estiver aberto no Word (PermissionError), grava em
    "<nome> (cópia).docx" - sempre a mesma cópia, para o salvamento
    automático da audiência não espalhar dezenas de arquivos.
    """
    destino = Path(destino)
    dados = montar_docx(falas, meta, marcar_tempo=marcar_tempo,
                        legenda_automatica=legenda_automatica)
    try:
        _gravar(destino, dados)
        return destino
    except PermissionError:
        log.warning("%s está aberto em outro programa; gravando uma cópia ao lado.",
                    destino.name)
    copia = destino.with_name(f"{destino.stem} (cópia){destino.suffix}")
    try:
        _gravar(copia, dados)
        return copia
    except PermissionError:
        alternativa = sistema.destino_livre(destino.parent, f"{destino.stem} (cópia)",
                                            destino.suffix)
        _gravar(alternativa, dados)
        return alternativa
