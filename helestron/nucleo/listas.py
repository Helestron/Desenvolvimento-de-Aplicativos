"""Lê a relação de processos de quase qualquer arquivo: Excel, Word, PDF, texto.

Herdado do Assessor SAJ (app/esaj/planilha.py), que já tinha apanhado de
arquivo de verdade. O que se aprendeu lá continua valendo:

1. A EXTENSÃO MENTE. O ".xls" que os sistemas do Judiciário exportam é,
   quase sempre, uma tabela HTML. O arquivo é reconhecido pelos primeiros
   bytes, não pelo fim do nome.
2. O EXCEL COME O ZERO DA FRENTE. CNJ guardado como número fica com 19
   dígitos; a célula numérica de 19 dígitos recebe o zero de volta.
3. NÚMERO GRANDE VIRA NOTAÇÃO CIENTÍFICA, e str(7.0e18) não contém número
   de processo nenhum. A célula numérica vira inteiro antes de ser lida.
4. O EXCEL CORROMPE O CNJ GUARDADO COMO NÚMERO: o double de 64 bits não
   comporta 20 dígitos, e os últimos saem trocados - e o número corrompido
   ainda parece um número de processo. Por isso a célula numérica que já
   não guarda o inteiro exato (de 2**53 em diante - todo CNJ como número)
   nunca entra no lote, mesmo que o dígito verificador confira: ele é
   módulo 97, e cerca de 1 em 97 números corrompidos passaria por ele. O
   inteiro exato (raro: XML montado à mão) só entra se o dígito conferir.
   Sem isso, o programa baixaria processo alheio sem ninguém perceber.
5. PLANILHA COM FÓRMULA que o Excel nunca abriu não tem valor gravado:
   tenta-se de novo pelo texto das fórmulas.

E o que foi corrigido em relação à base (todos com teste):
  * o sufixo de dependente só é aceito com barra ("/01"): "... - 2ª Vara"
    deixava de ser lido como o incidente 02; e, mesmo com a barra, não
    quando o que vem depois é texto ("... / 1ª Vara", ".../2ª Vara",
    "... / 3 réus" continuam sendo os autos principais - cnj._PADRAO);
  * o Word é lido na ordem do documento (parágrafos e tabelas
    intercalados), e não todos os parágrafos antes de todas as tabelas;
  * várias colunas de número: se uma delas se chama "processo" (e não
    "processo de origem", "principal"...), só ela é lida;
  * abas ocultas da planilha são ignoradas;
  * arquivo em UTF-16 (BOM) é lido;
  * senha de processo sigiloso na lista ("número ; senha" ou coluna
    "senha") fica associada ao número, com a mesma grafia usada no
    download;
  * nada de estado global: tudo volta no objeto Leitura.
"""

from __future__ import annotations

import csv
import logging
import re
import urllib.parse
import urllib.request
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.etree import ElementTree as ET

from . import cnj, sistema

log = logging.getLogger("listas")

EXTENSOES = (".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".txt", ".htm",
             ".html", ".pdf", ".docx")
TIPOS_ARQUIVO = [
    ("Relação de processos", " ".join("*" + e for e in EXTENSOES)),
    ("Planilhas", "*.xlsx *.xlsm *.xls *.ods *.csv"),
    ("Documentos", "*.docx *.pdf *.txt"),
    ("Todos os arquivos", "*.*"),
]
CODIFICACOES = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
LNLN = "\n\n"

# Cabeçalho da coluna que traz o número do processo...
_COLUNA_NUMERO = re.compile(r"processo|n[uú]mero|autos|\bn[º°o]\b|cnj", re.I)
# ...e o que, mesmo com "processo" no nome, é OUTRO processo.
_COLUNA_OUTRO = re.compile(
    r"origem|principal|apens|refer[eê]ncia|vinculad|relacionad|anterior|"
    r"originári|piloto|paradigma", re.I)
_COLUNA_SENHA = re.compile(r"senha", re.I)

# "número ; senha" - a linha inteira tem de ser só isso, para a segunda
# coluna de uma tabela ("João da Silva", "Audiência") não virar senha.
_SENHA_DEPOIS = re.compile(
    r"^\s*[;|\t:]\s*(?:senhas?\s*[:=]\s*)?([^\s;|\t]{3,40})\s*$", re.I)
_ANTES_DO_NUMERO = re.compile(r"^\s*(?:[-•*]|\d{1,4}[.)º°]?)?\s*$")


class ListaInvalida(RuntimeError):
    """O arquivo não pode ser lido, ou não tem número de processo dentro."""


@dataclass
class Leitura:
    """O resultado de uma leitura: o que entra, e o que a tela deve avisar."""
    processos: list[cnj.Numero] = field(default_factory=list)
    senhas: dict[str, str] = field(default_factory=dict)   # Numero.formatado -> senha
    corrompidos: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)
    formato: str = ""
    origem: str = ""

    @property
    def digito_errado(self) -> list[cnj.Numero]:
        return [n for n in self.processos if not n.digito_confere]

    def juntar(self, outra: "Leitura") -> "Leitura":
        """Acrescenta outra leitura, sem repetir processo."""
        vistos = {cnj.chave(n) for n in self.processos}
        for n in outra.processos:
            if cnj.chave(n) not in vistos:
                vistos.add(cnj.chave(n))
                self.processos.append(n)
        self.senhas.update(outra.senhas)
        self.corrompidos += outra.corrompidos
        self.avisos += outra.avisos
        return self


# ----------------------------------------------------------------- utilidades
def _digitos(texto: str) -> str:
    return re.sub(r"\D", "", texto or "")


def _texto_da_celula(valor) -> str:
    """O conteúdo da célula em texto que o buscador consiga ler."""
    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, int):
        digitos = str(abs(valor))
    elif isinstance(valor, float):
        if valor != valor or valor in (float("inf"), float("-inf")):
            return ""
        inteiro = int(valor)
        if inteiro != valor:
            return repr(valor)
        digitos = str(abs(inteiro))
    else:
        return str(valor)
    if len(digitos) == 19:          # o Excel comeu o zero da frente
        digitos = "0" + digitos
    return digitos


# Célula numérica (float) que não guarda mais o inteiro exato: de 2**53 em
# diante, o espaçamento entre dois doubles vizinhos passa de 1. O CNJ como
# número tem de 18 a 20 dígitos (o zero da frente some) - sempre nessa faixa.
# Acima de 10**20 não há CNJ possível, e a célula segue o caminho comum.
_FLOAT_INEXATO = float(2 ** 53)
_ALEM_DO_CNJ = 1e20


class _Coletor:
    """Junta números e senhas, na ordem, sem repetir."""

    def __init__(self) -> None:
        self.leitura = Leitura()
        self._vistos: set[str] = set()

    def numero(self, n: cnj.Numero, senha: str = "") -> None:
        k = cnj.chave(n)
        if k not in self._vistos:
            self._vistos.add(k)
            self.leitura.processos.append(n)
        if senha:
            self.leitura.senhas[n.formatado] = senha

    def texto(self, texto: str, senha: str = "") -> int:
        achados = cnj.extrair_todos(texto)
        for n in achados:
            self.numero(n, senha if len(achados) == 1 else "")
        return len(achados)

    def linha_livre(self, linha: str) -> None:
        """Linha de texto colado: 'numero ; senha' associa a senha."""
        achados = cnj.extrair_todos(linha)
        senha = ""
        if len(achados) == 1:
            m = cnj._PADRAO.search(linha)
            depois = _SENHA_DEPOIS.match(linha[m.end():]) if m else None
            if m and depois and _ANTES_DO_NUMERO.match(linha[:m.start()]):
                candidata = depois.group(1).strip()
                if len(_digitos(candidata)) < 15:   # não é outro número CNJ
                    senha = candidata
        for n in achados:
            self.numero(n, senha)

    def celula(self, valor) -> None:
        texto = _texto_da_celula(valor)
        if not texto:
            return
        if isinstance(valor, float) and _FLOAT_INEXATO <= abs(valor) < _ALEM_DO_CNJ:
            # Item 4 do docstring. Acima de 2**53 o double não guarda mais o
            # inteiro exato (o Excel, aliás, só guarda 15 algarismos): o CNJ
            # lido daqui já não é o da relação. O dígito verificador não basta
            # como filtro - é módulo 97, e cerca de 1 em 97 números
            # corrompidos ainda confere (com o foro, ou até o tribunal,
            # trocado). Nada desta célula entra no lote; ela é contada entre
            # os corrompidos (inclusive o número com "00" na frente, que nem
            # chega a parecer um CNJ), para o aviso "formate como Texto".
            achados = cnj.extrair_todos(texto)
            self.leitura.corrompidos.extend(
                [n.formatado for n in achados] or [texto])
            return
        # Inteiro de verdade (XML montado à mão, com todos os dígitos) é
        # exato: vale o dígito verificador, como antes.
        numerico = isinstance(valor, (int, float)) and not isinstance(valor, bool)
        for n in cnj.extrair_todos(texto):
            if numerico and not n.digito_confere:
                self.leitura.corrompidos.append(n.formatado)
                continue
            self.numero(n)


# -------------------------------------------------------------- reconhecer
def formato(caminho: Path) -> str:
    """O que o arquivo É, pelos primeiros bytes - não pelo nome."""
    try:
        with open(caminho, "rb") as f:
            cabeca = f.read(4096)
    except OSError as erro:
        raise ListaInvalida(f"não consegui abrir o arquivo: {str(erro)[:120]}") from erro

    if cabeca.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(caminho) as z:
                nomes = set(z.namelist())
        except Exception:
            return "zip-quebrado"
        if "content.xml" in nomes:
            return "ods"
        if any(n.startswith("word/") for n in nomes):
            return "docx"
        if any(n.startswith("xl/") for n in nomes):
            return "xlsx"
        return "zip-desconhecido"
    if cabeca.startswith(b"\xd0\xcf\x11\xe0"):
        # OLE: .xls de verdade - ou um .doc antigo do Word, com a mesma cara.
        return "doc" if b"W\x00o\x00r\x00d\x00D\x00o\x00c" in cabeca or \
            caminho.suffix.lower() == ".doc" else "xls"
    if cabeca.startswith(b"%PDF"):
        return "pdf"
    texto = _decodificar(cabeca)[:800].lower()
    if "<html" in texto or "<table" in texto or "<!doctype html" in texto:
        return "html"
    return "texto"


def _decodificar(dados: bytes) -> str:
    if dados.startswith((b"\xff\xfe", b"\xfe\xff")):
        return dados.decode("utf-16", errors="replace")
    # UTF-16 sem BOM: metade dos bytes é zero.
    if len(dados) > 20 and dados[1:200:2].count(0) > 0.4 * len(dados[1:200:2]):
        return dados.decode("utf-16-le", errors="replace")
    for codificacao in CODIFICACOES:
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("latin-1", errors="replace")


def _sem_marcacao(bruto: str) -> str:
    # Etiqueta vira espaço: dois números vizinhos não podem se colar.
    sem = re.sub(r"<[^>]+>", " ", bruto)
    return (sem.replace("&nbsp;", " ").replace("&#160;", " ")
               .replace("&amp;", "&"))


# --------------------------------------------------------------- planilhas
def _escolher_coluna(linhas: list[tuple]) -> tuple[int | None, int | None]:
    """Nas 10 primeiras linhas, a coluna do número e a da senha, se houver."""
    for linha in linhas[:10]:
        num = senha = None
        for i, valor in enumerate(linha):
            if not isinstance(valor, str):
                continue
            rotulo = valor.strip()
            if not rotulo or cnj.extrair_todos(rotulo):
                continue
            if num is None and _COLUNA_NUMERO.search(rotulo) and not _COLUNA_OUTRO.search(rotulo):
                num = i
            elif senha is None and _COLUNA_SENHA.search(rotulo):
                senha = i
        if num is not None:
            return num, senha
    return None, None


def _ler_linhas(col: _Coletor, linhas: list[tuple], aviso_aba: str) -> None:
    coluna, coluna_senha = _escolher_coluna(linhas)
    if coluna is not None:
        antes = len(col.leitura.processos)
        corrompidos = len(col.leitura.corrompidos)
        for linha in linhas:
            valor = linha[coluna] if coluna < len(linha) else None
            senha = ""
            if coluna_senha is not None and coluna_senha < len(linha):
                senha = _texto_da_celula(linha[coluna_senha]).strip()
            n_antes = len(col.leitura.processos)
            col.celula(valor)
            if senha and len(col.leitura.processos) > n_antes:
                col.leitura.senhas[col.leitura.processos[-1].formatado] = senha
        if len(col.leitura.processos) > antes:
            return
        # A linha inteira vai ser lida: o que a coluna já contou como
        # corrompido não conta duas vezes (a tela diria o dobro de linhas).
        del col.leitura.corrompidos[corrompidos:]
    # Sem coluna reconhecida (ou ela veio vazia): varre a linha inteira.
    for linha in linhas:
        for valor in linha:
            col.celula(valor)


def _de_xlsx(caminho: Path, col: _Coletor) -> None:
    from openpyxl import load_workbook

    def tentar(valores: bool, alvo: _Coletor) -> None:
        # Em memória: o openpyxl recusa pela extensão, e aqui a extensão já
        # se mostrou mentirosa; e no modo read_only ele lê as abas depois.
        try:
            livro = load_workbook(BytesIO(caminho.read_bytes()),
                                  read_only=True, data_only=valores)
        except Exception as erro:
            raise ListaInvalida(
                f"não consegui abrir a planilha: {str(erro)[:120]}") from erro
        try:
            for nome in livro.sheetnames:
                aba = livro[nome]
                if getattr(aba, "sheet_state", "visible") != "visible":
                    alvo.leitura.avisos.append(f"aba oculta '{nome}' ignorada")
                    continue
                # No modo read_only, o openpyxl lê só até a "dimensão" que o
                # arquivo declara - e as planilhas exportadas por sistemas
                # (o próprio SAJ, relatórios, Apache POI) costumam declarar
                # "A1": lia-se a primeira linha e mais nada, e a relação
                # inteira era recusada ("nenhum número de processo").
                redefinir = getattr(aba, "reset_dimensions", None)
                if callable(redefinir):
                    redefinir()
                linhas = list(aba.iter_rows(values_only=True))
                _ler_linhas(alvo, linhas, nome)
        finally:
            try:
                livro.close()
            except Exception:
                pass

    erro_openpyxl: ListaInvalida | None = None
    try:
        tentar(True, col)
    except ListaInvalida as erro:
        erro_openpyxl = erro
    if not col.leitura.processos and erro_openpyxl is None:
        # De novo, pelo texto das fórmulas. As células numéricas e as abas
        # ocultas são as mesmas da primeira leitura: os corrompidos e os
        # avisos não são contados de novo (a tela diria o dobro de linhas).
        formulas = _Coletor()
        try:
            tentar(False, formulas)
        except ListaInvalida:
            pass
        for n in formulas.leitura.processos:
            col.numero(n, formulas.leitura.senhas.get(n.formatado, ""))
    if not col.leitura.processos:
        # O openpyxl recusa variantes legítimas do formato (Strict Open XML,
        # arquivos gerados por outros programas com estilos que ele não
        # entende): lê-se o texto das células direto do XML.
        _de_xlsx_bruto(caminho, col)
    if not col.leitura.processos and erro_openpyxl is not None:
        raise erro_openpyxl


_RE_CELULA_XML = re.compile(rb"<(?:\w+:)?(?:t|v)(?:\s[^>]*)?>([^<]*)</(?:\w+:)?(?:t|v)>")


def _de_xlsx_bruto(caminho: Path, col: _Coletor) -> None:
    """O texto das células (textos compartilhados e das abas), sem o
    openpyxl: só o que é texto entra, como na leitura normal."""
    try:
        with zipfile.ZipFile(caminho) as z:
            nomes = z.namelist()
            if "xl/workbook.bin" in nomes:
                raise ListaInvalida(
                    "esta é uma pasta de trabalho binária do Excel (.xlsb), que o "
                    "programa não lê. No Excel, use Salvar como > Pasta de Trabalho "
                    "do Excel (.xlsx), ou CSV, e escolha de novo.")
            partes = [n for n in nomes if n == "xl/sharedStrings.xml"
                      or (n.startswith("xl/worksheets/") and n.endswith(".xml"))]
            for nome in partes:
                dados = z.read(nome)
                for m in _RE_CELULA_XML.finditer(dados):
                    bruto = m.group(1).decode("utf-8", errors="ignore")
                    texto = (bruto.replace("&amp;", "&").replace("&lt;", "<")
                             .replace("&gt;", ">").replace("&quot;", '"')
                             .replace("&apos;", "'"))
                    # Só texto: o valor numérico de uma célula (<v> de célula
                    # numérica) não tem os 20 dígitos do CNJ de forma confiável.
                    if re.search(r"\D", texto.strip()):
                        col.texto(texto)
    except ListaInvalida:
        raise
    except Exception as erro:
        log.debug("leitura direta do XML da planilha falhou: %s", type(erro).__name__)


def _de_xls(caminho: Path, col: _Coletor) -> None:
    import xlrd

    try:
        cabeca = caminho.read_bytes()[:1_000_000]
    except OSError:
        cabeca = b""
    if "EncryptedPackage".encode("utf-16-le") in cabeca:
        # .xlsx com senha de abertura: o Excel o guarda dentro de um arquivo
        # OLE, com a mesma cara de um .xls antigo.
        raise ListaInvalida(
            "esta planilha está protegida por senha de abertura, e o programa não a "
            "abre. No Excel, abra-a, tire a senha (Arquivo > Informações > Proteger "
            "pasta de trabalho > Criptografar com senha, e apague a senha), salve e "
            "escolha de novo.")
    try:
        livro = xlrd.open_workbook(str(caminho))
    except Exception as erro:
        raise ListaInvalida(
            f"não consegui abrir a planilha antiga: {str(erro)[:120]}") from erro
    for aba in livro.sheets():
        if getattr(aba, "visibility", 0):
            col.leitura.avisos.append(f"aba oculta '{aba.name}' ignorada")
            continue
        linhas = [tuple(aba.row_values(i)) for i in range(aba.nrows)]
        _ler_linhas(col, linhas, aba.name)


_NS_ODS = {
    "table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
    "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0",
    "text": "urn:oasis:names:tc:opendocument:xmlns:text:1.0",
}


def _de_ods(caminho: Path, col: _Coletor) -> None:
    try:
        with zipfile.ZipFile(caminho) as z:
            raiz = ET.fromstring(z.read("content.xml"))
    except Exception as erro:
        raise ListaInvalida(f"não consegui abrir o arquivo .ods: {str(erro)[:120]}") from erro
    t, o = _NS_ODS["table"], _NS_ODS["office"]
    for tabela in raiz.iter(f"{{{t}}}table"):
        linhas = []
        for linha in tabela.iter(f"{{{t}}}table-row"):
            valores = []
            for celula in linha.iter(f"{{{t}}}table-cell"):
                tipo = celula.get(f"{{{o}}}value-type")
                if tipo == "float" and celula.get(f"{{{o}}}value"):
                    try:
                        valores.append(float(celula.get(f"{{{o}}}value")))
                        continue
                    except ValueError:
                        pass
                valores.append(" ".join("".join(p.itertext()) for p in celula))
            linhas.append(tuple(valores))
        _ler_linhas(col, linhas, tabela.get(f"{{{t}}}name", ""))


# --------------------------------------------------------------- documentos
_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _texto_do_elemento(el) -> str:
    partes = []
    for x in el.iter():
        if x.tag == f"{_W}t" and x.text:
            partes.append(x.text)
        elif x.tag in (f"{_W}tab", f"{_W}br", f"{_W}cr"):
            partes.append(" ")
    return "".join(partes)


def _blocos_word(el):
    """Parágrafos e tabelas do Word NA ORDEM DO DOCUMENTO.

    Devolve ('p', texto) ou ('tabela', [linhas]). Controles de conteúdo
    (w:sdt) e outros invólucros são abertos; a tabela vem por linha, para a
    coluna do número (e a da senha) ser reconhecida como numa planilha.
    """
    for filho in el:
        tag = filho.tag
        if tag == f"{_W}p":
            yield "p", _texto_do_elemento(filho)
        elif tag == f"{_W}tbl":
            linhas = []
            for tr in filho.iter(f"{_W}tr"):
                celulas = [_texto_do_elemento(tc).strip()
                           for tc in tr if tc.tag == f"{_W}tc"]
                linhas.append(tuple(celulas))
            yield "tabela", linhas
        else:
            yield from _blocos_word(filho)


def _de_docx(caminho: Path, col: _Coletor) -> None:
    try:
        with zipfile.ZipFile(caminho) as z:
            nomes = z.namelist()
            corpo = ET.fromstring(z.read("word/document.xml"))
            extras = [ET.fromstring(z.read(n)) for n in sorted(nomes)
                      if re.match(r"word/(header|footer)\d*\.xml$", n)]
    except Exception as erro:
        raise ListaInvalida(f"não consegui abrir o documento: {str(erro)[:120]}") from erro
    for tipo, conteudo in _blocos_word(corpo):
        if tipo == "p":
            col.linha_livre(conteudo)
        else:
            _ler_linhas(col, conteudo, "tabela")
    for raiz in extras:
        col.texto(_texto_do_elemento(raiz))


def _de_pdf(caminho: Path, col: _Coletor) -> None:
    try:
        try:
            import pymupdf
        except ImportError:  # pragma: no cover
            import fitz as pymupdf
        with pymupdf.open(str(caminho)) as doc:
            texto = "\n".join(p.get_text("text") for p in doc)
    except Exception as erro:
        raise ListaInvalida(f"não consegui abrir o PDF: {str(erro)[:120]}") from erro
    # Número quebrado no fim da linha do PDF: "8.02.\n0001" volta a ser um só.
    texto = re.sub(r"([.\-])\s*\n\s*(?=\d)", r"\1", texto)
    for linha in texto.splitlines():
        col.linha_livre(linha)


def _de_html(caminho: Path, col: _Coletor) -> None:
    bruto = _decodificar(caminho.read_bytes())
    linhas = re.split(r"(?i)</tr>|<br\s*/?>|</p>", bruto)
    for linha in linhas:
        col.texto(_sem_marcacao(linha))


def _de_texto(caminho: Path, col: _Coletor) -> None:
    bruto = _decodificar(caminho.read_bytes())
    ler_texto_em(col, bruto, csv_provavel=caminho.suffix.lower() == ".csv")


def ler_texto_em(col: _Coletor, bruto: str, csv_provavel: bool = False) -> None:
    if csv_provavel or (";" in bruto[:2000] and "\n" in bruto):
        try:
            dialeto = csv.Sniffer().sniff(bruto[:4000], delimiters=";,\t|")
            linhas = [tuple(l) for l in csv.reader(bruto.splitlines(), dialeto)]
            # Sem cabeçalho, "número ; senha" digitado linha a linha vai para a
            # leitura por linha, a que sabe associar a senha ao número. Tabela
            # colada do Excel (tabulação) não: sem o cabeçalho "Senha", a
            # segunda coluna é outra coisa (classe, vara, parte).
            if linhas and max(len(l) for l in linhas) > 1 and (
                    dialeto.delimiter == "\t" or _escolher_coluna(linhas)[0] is not None):
                _ler_linhas(col, linhas, "csv")
                return
        except csv.Error:
            pass
    for linha in bruto.splitlines():
        limpa = linha.strip()
        if limpa and limpa[0] in "#;":
            continue
        col.linha_livre(limpa)


_LEITORES = {"xlsx": _de_xlsx, "xls": _de_xls, "ods": _de_ods, "html": _de_html,
             "pdf": _de_pdf, "docx": _de_docx, "texto": _de_texto}
_NOMES = {"xlsx": "planilha do Excel", "xls": "planilha antiga do Excel",
          "ods": "planilha do LibreOffice", "html": "tabela em HTML",
          "pdf": "PDF", "docx": "documento do Word", "texto": "texto",
          "doc": "documento antigo do Word (.doc)"}


# ------------------------------------------------------------------- entrada
def ler_arquivo(caminho: Path | str) -> Leitura:
    """Os processos do arquivo, na ordem em que aparecem, sem repetir."""
    caminho = Path(caminho)
    if not caminho.exists():
        raise ListaInvalida(f"não encontrei o arquivo {caminho.name}")
    if caminho.stat().st_size == 0:
        raise ListaInvalida(f"o arquivo {caminho.name} está vazio.")

    tipo = formato(caminho)
    log.info("Relação %s: reconhecida como %s.", caminho.name, _NOMES.get(tipo, tipo))
    if tipo in ("zip-quebrado", "zip-desconhecido"):
        raise ListaInvalida(
            f"o arquivo {caminho.name} parece danificado. Abra-o no Excel ou "
            "no Word e salve de novo (.xlsx ou .docx).")

    col = _Coletor()
    col.leitura.formato = _NOMES.get(tipo, tipo)
    col.leitura.origem = str(caminho)
    erro_leitor: Exception | None = None
    if tipo == "doc":
        # .doc do Word 97: o texto fica em UTF-16 dentro do arquivo OLE.
        dados = caminho.read_bytes()
        col.texto(dados.decode("utf-16-le", errors="ignore"))
        col.texto(dados.decode("latin-1", errors="ignore"))
    else:
        try:
            _LEITORES.get(tipo, _de_texto)(caminho, col)
        except ListaInvalida as erro:
            erro_leitor = erro
        except ImportError as erro:  # pragma: no cover - instalação incompleta
            erro_leitor = ListaInvalida(
                f"falta uma biblioteca para ler {_NOMES.get(tipo, tipo)} "
                f"({erro.name}). {sistema.REINSTALAR}")

    # Última tentativa: texto com outra cara. Custa pouco e salva o arquivo
    # com extensão trocada.
    if not col.leitura.processos and tipo != "texto":
        try:
            _de_texto(caminho, col)
        except Exception:
            pass

    if not col.leitura.processos:
        if col.leitura.corrompidos:
            raise ListaInvalida(
                "esta planilha guarda os números de processo como NÚMERO, e "
                "não como texto." + LNLN
                + "O Excel não comporta os 20 dígitos do CNJ num campo "
                  "numérico: ele arredonda os últimos, e o número deixa de ser "
                  f"o do processo. Foi o que aconteceu com {len(col.leitura.corrompidos)} "
                  "linha(s) deste arquivo." + LNLN
                + "Para resolver: no Excel, formate a coluna dos números como "
                  "Texto e cole-os de novo; ou salve a relação como .csv." + LNLN
                + "Deixei de importá-los de propósito: importados assim, o "
                  "programa baixaria processo alheio.")
        if erro_leitor is not None:
            raise erro_leitor
        raise ListaInvalida(
            f"li o arquivo {caminho.name} ({_NOMES.get(tipo, tipo)}), mas não "
            "achei nenhum número de processo dentro." + LNLN
            + "Confira se os números estão no padrão CNJ "
              "(0000000-00.0000.0.00.0000). Número dentro de imagem colada "
              "não pode ser lido.")
    return col.leitura


def ler_texto(texto: str) -> Leitura:
    """Lista colada na tela (ou copiada do Excel, que chega em colunas)."""
    col = _Coletor()
    col.leitura.formato = "texto colado"
    ler_texto_em(col, texto or "", csv_provavel="\t" in (texto or ""))
    return col.leitura


# --------------------------------------------------------- link compartilhado
def url_de_download(url: str) -> str:
    """Converte o link de compartilhamento no link que entrega o arquivo."""
    url = url.strip()
    m = re.search(r"docs\.google\.com/spreadsheets/d/([\w-]+)", url)
    if m:
        return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=xlsx"
    m = re.search(r"docs\.google\.com/document/d/([\w-]+)", url)
    if m:
        return f"https://docs.google.com/document/d/{m.group(1)}/export?format=docx"
    m = re.search(r"drive\.google\.com/(?:file/d/|open\?id=)([\w-]+)", url)
    if m:
        return f"https://drive.google.com/uc?export=download&id={m.group(1)}"
    partes = urllib.parse.urlsplit(url)
    host = partes.netloc.lower()
    if "sharepoint.com" in host or "onedrive.live.com" in host or "1drv.ms" in host:
        consulta = urllib.parse.parse_qs(partes.query)
        consulta["download"] = ["1"]
        return urllib.parse.urlunsplit(partes._replace(
            query=urllib.parse.urlencode(consulta, doseq=True)))
    return url


def baixar_link(url: str, pasta: Path) -> Path:
    """Baixa a relação de um link público e devolve o arquivo salvo."""
    destino_url = url_de_download(url)
    pedido = urllib.request.Request(destino_url, headers={
        "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Helestron"})
    try:
        with urllib.request.urlopen(pedido, timeout=60) as resp:
            dados = resp.read(50 * 1024 * 1024)
            final = resp.geturl()
            nome = ""
            disp = resp.headers.get("Content-Disposition", "")
            m = re.search(r"filename\*?=(?:UTF-8'')?\"?([^\";]+)", disp)
            if m:
                nome = urllib.parse.unquote(m.group(1))
    except Exception as erro:
        raise ListaInvalida(
            f"não consegui baixar o link: {str(erro)[:160]}" + LNLN
            + "Confira se ele está compartilhado como 'qualquer pessoa com o "
              "link'. Se for da rede do tribunal, baixe o arquivo e abra-o "
              "pelo botão 'Abrir arquivo'.") from erro
    amostra = dados[:2000].lower()
    if (b"<html" in amostra and any(s in final for s in (
            "accounts.google.com", "login.microsoftonline.com", "login.live.com"))):
        raise ListaInvalida(
            "o link pede login: ele não é público. Compartilhe como 'qualquer "
            "pessoa com o link' ou baixe o arquivo e abra-o pelo botão 'Abrir "
            "arquivo'.")
    pasta.mkdir(parents=True, exist_ok=True)
    nome = re.sub(r'[<>:"/\\|?*]+', "_", nome).strip() or \
        f"relacao {datetime.now():%Y-%m-%d %Hh%M}"
    destino = pasta / nome
    destino.write_bytes(dados)
    return destino
