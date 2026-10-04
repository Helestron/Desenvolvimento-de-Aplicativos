"""Importar o relatório da pauta exportado do SAJ, do eProc ou de qualquer lugar (8.6).

Planilha (xlsx, xls, ods, csv), HTML (inclusive o ".xls" que é HTML - o
que os sistemas do Judiciário mais exportam), DOCX e PDF. O formato é
reconhecido pelos primeiros bytes (nucleo/listas.formato, a lição do
Assessor SAJ: a extensão mente), e cada formato vira Tabela - daí em
diante, as mesmas regras de reconhecimento da pauta lida no portal
(seção 8.5), com duas diferenças: o sistema é "arquivo" e o tribunal vem
do próprio número do processo (8.02 -> TJAL).

O PDF não tem tabela de verdade, só texto posicionado: as palavras são
agrupadas em linhas pela altura e em colunas pelas posições do cabeçalho
("tabela por texto"). Sem cabeçalho reconhecível, cada linha que tenha
data e número de processo vira uma audiência. Linhas que continuam a de
cima (as partes que não couberam) são juntadas a ela.
"""

from __future__ import annotations

import csv
import io
import logging
import re
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from xml.etree import ElementTree as ET

from ..nucleo import cnj, listas
from . import modelos
from .tabelas import Celula, Reconhecedor, Reconhecimento, Tabela, mapear_cabecalho, \
    tabelas_do_html, tem_valores

log = logging.getLogger("pauta.importacao")

EXTENSOES = (".xlsx", ".xlsm", ".xls", ".ods", ".csv", ".txt", ".htm", ".html", ".pdf", ".docx")
TIPOS_ARQUIVO = [("Relatório da pauta", " ".join("*" + e for e in EXTENSOES))]
MAX_LINHAS = 20000
# Abas que acompanham a pauta (a planilha do próprio Helestron tem Resumo e
# Alterações): com outra aba na planilha, elas não são lidas como pauta - o
# "Quando" das alterações e as contagens do resumo não são audiências.
ABAS_DE_APOIO = {"resumo", "alteracoes", "historico", "estatisticas", "graficos", "totais",
                 "legenda", "instrucoes"}


def _sem_abas_de_apoio(tabelas: list[Tabela]) -> list[Tabela]:
    if len(tabelas) <= 1:
        return tabelas
    principais = [t for t in tabelas if modelos.normalizar_texto(t.identificador)
                  not in ABAS_DE_APOIO]
    return principais or tabelas


class RelatorioInvalido(ValueError):
    """O arquivo não pôde ser lido como relatório de pauta."""


# ================================================================ planilhas
def _descer_mescladas(linhas: list[list], mescladas) -> None:
    """Célula mesclada na VERTICAL (a data escrita uma vez para as audiências do
    dia): o valor vale em todas as linhas que ela cobre, na primeira coluna da
    mescla. Na horizontal (o título do dia na linha inteira), as outras colunas
    ficam vazias - como o colspan do HTML. 'mescladas': (linha0, linha1,
    coluna0), com a linha1 exclusiva, a partir de 0."""
    for l0, l1, c0 in mescladas:
        if l1 - l0 < 2 or l0 >= len(linhas) or c0 >= len(linhas[l0]):
            continue
        valor = linhas[l0][c0]
        for i in range(l0 + 1, min(l1, len(linhas))):
            while len(linhas[i]) <= c0:
                linhas[i].append(None)
            if linhas[i][c0] in (None, ""):
                linhas[i][c0] = valor


def _de_xlsx(caminho: Path) -> list[Tabela]:
    from openpyxl import load_workbook

    try:
        # sem read_only: só assim o openpyxl conta as células mescladas (o relatório
        # da pauta é pequeno)
        livro = load_workbook(io.BytesIO(caminho.read_bytes()), data_only=True)
    except Exception as erro:
        raise RelatorioInvalido(f"não consegui abrir a planilha ({str(erro)[:120]}).") from erro
    saida = []
    try:
        for nome in livro.sheetnames:
            aba = livro[nome]
            if getattr(aba, "sheet_state", "visible") != "visible":
                continue
            linhas = []
            for i, linha in enumerate(aba.iter_rows(values_only=True)):
                if i >= MAX_LINHAS:
                    break
                linhas.append(list(linha))
            try:
                faixas = list(aba.merged_cells.ranges)
            except AttributeError:
                faixas = []
            # iter_rows começa em A1
            _descer_mescladas(linhas, [(f.min_row - 1, f.max_row, f.min_col - 1) for f in faixas])
            saida.append(Tabela.de_textos(linhas, origem=f"{caminho.name} › {nome}",
                                          identificador=nome))
    finally:
        try:
            livro.close()
        except Exception:
            pass
    return _sem_abas_de_apoio(saida)


def _de_xls(caminho: Path) -> list[Tabela]:
    import xlrd

    try:
        try:
            livro = xlrd.open_workbook(str(caminho), formatting_info=True)   # as mescladas
        except NotImplementedError:
            livro = xlrd.open_workbook(str(caminho))
    except Exception as erro:
        raise RelatorioInvalido(
            f"não consegui abrir a planilha antiga ({str(erro)[:120]}).") from erro
    saida = []
    for aba in livro.sheets():
        if getattr(aba, "visibility", 0):
            continue
        linhas = []
        for i in range(min(aba.nrows, MAX_LINHAS)):
            valores = []
            for c in aba.row(i):
                if c.ctype == xlrd.XL_CELL_DATE:
                    try:
                        valores.append(xlrd.xldate_as_datetime(c.value, livro.datemode))
                        continue
                    except Exception:
                        pass
                valores.append(c.value if c.value != "" else None)
            linhas.append(valores)
        _descer_mescladas(linhas, [(l0, l1, c0) for l0, l1, c0, _c1 in
                                   (getattr(aba, "merged_cells", None) or [])])
        saida.append(Tabela.de_textos(linhas, origem=f"{caminho.name} › {aba.name}",
                                      identificador=aba.name))
    return _sem_abas_de_apoio(saida)


_ODS = {"table": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
        "office": "urn:oasis:names:tc:opendocument:xmlns:office:1.0"}
_RE_DURACAO = re.compile(r"PT(\d+)H(\d+)M")


def _valor_ods(celula):
    o = _ODS["office"]
    tipo = celula.get(f"{{{o}}}value-type")
    if tipo == "date":
        texto = celula.get(f"{{{o}}}date-value") or ""
        try:
            return datetime.fromisoformat(texto) if "T" in texto else date.fromisoformat(texto)
        except ValueError:
            pass
    elif tipo == "time":
        m = _RE_DURACAO.match(celula.get(f"{{{o}}}time-value") or "")
        if m:
            return time(int(m.group(1)) % 24, int(m.group(2)))
    elif tipo in ("float", "percentage", "currency"):
        try:
            return float(celula.get(f"{{{o}}}value"))
        except (TypeError, ValueError):
            pass
    texto = " ".join("".join(p.itertext()) for p in celula)
    return texto or None


def _de_ods(caminho: Path) -> list[Tabela]:
    try:
        with zipfile.ZipFile(caminho) as z:
            raiz = ET.fromstring(z.read("content.xml"))
    except Exception as erro:
        raise RelatorioInvalido(
            f"não consegui abrir o arquivo .ods ({str(erro)[:120]}).") from erro
    t = _ODS["table"]
    saida = []
    for tabela in raiz.iter(f"{{{t}}}table"):
        linhas = []
        # coluna -> [valor, linhas que ainda cobre]: a célula mesclada na vertical
        descendo: dict[int, list] = {}
        for linha in tabela.iter(f"{{{t}}}table-row"):
            valores = []
            for celula in linha:
                if not celula.tag.endswith("table-cell") and not celula.tag.endswith(
                        "covered-table-cell"):
                    continue
                coluna = len(valores)
                valor = _valor_ods(celula)
                if celula.tag.endswith("covered-table-cell") and valor in (None, "") and \
                        coluna in descendo:
                    valor = descendo[coluna][0]
                try:
                    repetir = int(celula.get(f"{{{t}}}number-columns-repeated") or 1)
                except ValueError:
                    repetir = 1
                try:
                    altura = int(celula.get(f"{{{t}}}number-rows-spanned") or 1)
                except ValueError:
                    altura = 1
                if altura > 1 and valor not in (None, ""):
                    descendo[coluna] = [valor, altura]
                # o LibreOffice "comprime" colunas repetidas (as vazias até o fim: 1024)
                valores.extend([valor] * min(repetir, 50))
            for coluna in list(descendo):
                descendo[coluna][1] -= 1
                if descendo[coluna][1] <= 0:
                    del descendo[coluna]
            while valores and valores[-1] in (None, ""):
                valores.pop()
            if valores:
                linhas.append(valores)
            if len(linhas) >= MAX_LINHAS:
                break
        nome = tabela.get(f"{{{t}}}name", "")
        saida.append(Tabela.de_textos(linhas, origem=f"{caminho.name} › {nome}",
                                      identificador=nome))
    return _sem_abas_de_apoio(saida)


def _de_csv(caminho: Path) -> list[Tabela]:
    bruto = listas._decodificar(caminho.read_bytes())
    amostra = bruto[:5000]
    try:
        dialeto = csv.Sniffer().sniff(amostra, delimiters=";,\t|")
        delimitador = dialeto.delimiter
    except csv.Error:
        delimitador = max(";,\t|", key=amostra.count)
    linhas = [list(x) for x in csv.reader(io.StringIO(bruto), delimiter=delimitador)][:MAX_LINHAS]
    return [Tabela.de_textos(linhas, origem=caminho.name)]


def _de_html(caminho: Path) -> list[Tabela]:
    return tabelas_do_html(listas._decodificar(caminho.read_bytes()), caminho.name)


# =============================================================== documentos
def _de_docx(caminho: Path) -> tuple[list[Tabela], list[str]]:
    try:
        with zipfile.ZipFile(caminho) as z:
            corpo = ET.fromstring(z.read("word/document.xml"))
    except Exception as erro:
        raise RelatorioInvalido(f"não consegui abrir o documento ({str(erro)[:120]}).") from erro
    tabelas, paragrafos = [], []
    for tipo, conteudo in listas._blocos_word(corpo):
        if tipo == "p":
            if conteudo.strip():
                paragrafos.append(conteudo.strip())
        else:
            tabelas.append(Tabela.de_textos([list(x) for x in conteudo], origem=caminho.name,
                                            legenda=" ".join(paragrafos[-2:])))
    return tabelas, paragrafos


_RODAPE = re.compile(r"^\s*(p[áa]g(ina)?\.?\s*\d+|gerad[oa] (em|pelo|por)|emitid[oa] (em|por)|"
                     r"impress[oa] (em|por))", re.I)
_FOLGA_LINHA = 3.0        # pontos: palavras com o meio a esta distância estão na mesma linha
_FOLGA_COLUNA = 7.0       # pontos: espaço maior que isto separa células


def _linhas_visuais(pagina) -> list[list[tuple[float, float, str]]]:
    """As palavras da página agrupadas em linhas: [(x0, x1, palavra)] por linha."""
    palavras = pagina.get_text("words")
    palavras.sort(key=lambda w: ((w[1] + w[3]) / 2, w[0]))
    linhas: list[list] = []
    centros: list[float] = []
    for w in palavras:
        centro = (w[1] + w[3]) / 2
        if linhas and abs(centro - centros[-1]) <= _FOLGA_LINHA:
            linhas[-1].append((w[0], w[2], w[4]))
        else:
            linhas.append([(w[0], w[2], w[4])])
            centros.append(centro)
    for linha in linhas:
        linha.sort(key=lambda p: p[0])
    return linhas


def _celulas(linha) -> list[tuple[float, float, str]]:
    """Palavras vizinhas viram uma célula; um espaço grande separa as células."""
    celulas: list[list] = []
    for x0, x1, texto in linha:
        if celulas and x0 - celulas[-1][1] <= _FOLGA_COLUNA:
            celulas[-1][1] = x1
            celulas[-1][2] += " " + texto
        else:
            celulas.append([x0, x1, texto])
    return [tuple(c) for c in celulas]


def _de_pdf(caminho: Path, regras) -> tuple[list[Tabela], list[str]]:
    """O PDF como tabela por texto (colunas pelo cabeçalho) e como linhas soltas."""
    try:
        try:
            import pymupdf
        except ImportError:  # pragma: no cover - nome antigo
            import fitz as pymupdf
        documento = pymupdf.open(str(caminho))
    except Exception as erro:
        raise RelatorioInvalido(f"não consegui abrir o PDF ({str(erro)[:120]}).") from erro
    tabelas: list[Tabela] = []
    soltas: list[str] = []
    colunas: list[float] | None = None
    linhas_tabela: list[list[Celula]] = []
    with documento:
        for pagina in documento:
            for linha in _linhas_visuais(pagina):
                celulas = _celulas(linha)
                textos = [c[2] for c in celulas]
                if _RODAPE.search(" ".join(textos)):
                    continue                     # "Página 1 de 3", "Gerado em ..."
                soltas.append(" | ".join(textos))
                rotulos = [Celula(texto=t) for t in textos]
                mapa = mapear_cabecalho(rotulos, regras)
                # "Data: 03/10/2026  Hora: 10:15" da emissão é preâmbulo, não o cabeçalho
                if "data" in mapa and len(mapa) >= 2 and not tem_valores(rotulos):
                    if colunas is None:
                        colunas = [c[0] for c in celulas]
                        linhas_tabela.append(rotulos)
                    continue                     # cabeçalho (repetido a cada página)
                if colunas is None:
                    continue
                alinhadas = [""] * len(colunas)
                for x0, _x1, texto in celulas:
                    i = max((j for j, cx in enumerate(colunas) if cx <= x0 + _FOLGA_COLUNA),
                            default=0)
                    alinhadas[i] = (alinhadas[i] + " " + texto).strip()
                juntas = " ".join(alinhadas)
                continua = (linhas_tabela and len(linhas_tabela) > 1
                            and not cnj.extrair_todos(juntas) and not modelos.ler_data(juntas)
                            and not modelos.ler_hora(juntas))
                if continua:
                    anterior = linhas_tabela[-1]
                    for i, texto in enumerate(alinhadas):
                        if texto and i < len(anterior):
                            anterior[i] = Celula(texto=f"{anterior[i].texto} {texto}".strip())
                    continue
                linhas_tabela.append([Celula(texto=t) for t in alinhadas])
    if colunas is not None:
        tabelas.append(Tabela(linhas=linhas_tabela, origem=caminho.name))
    return tabelas, soltas


def _linhas_soltas(textos: list[str], reconhecedor: Reconhecedor, r: Reconhecimento) -> None:
    """Sem tabela: cada linha com data e processo; linha só com data é a data das de baixo."""
    data_corrente: date | None = None
    for texto in textos:
        celulas = [Celula(texto=x.strip()) for x in texto.split(" | ") if x.strip()]
        if not celulas:
            continue
        juntas = " ".join(c.texto for c in celulas)
        numeros = cnj.extrair_todos(juntas)
        if not numeros:
            d = modelos.ler_data(juntas)
            if d is not None and len(juntas) <= 60:
                data_corrente = d
            continue
        a = reconhecedor.linha_livre(celulas, data_padrao=data_corrente)
        if a is None:
            r.ignoradas += 1
            r.avisar(f"Linha sem data reconhecível (processo {numeros[0].formatado}); ignorada.")
            continue
        r.audiencias.append(a)
    if r.audiencias:
        r.tabelas = max(r.tabelas, 1)


# ===================================================================== entrada
def ler_relatorio(caminho: Path | str, regras, agora: datetime | None = None) -> Reconhecimento:
    """As audiências do relatório. Levanta RelatorioInvalido com a frase para o usuário."""
    caminho = Path(caminho)
    if not caminho.is_file():
        raise RelatorioInvalido(f"não encontrei o arquivo {caminho.name}.")
    if caminho.stat().st_size == 0:
        raise RelatorioInvalido(f"o arquivo {caminho.name} está vazio.")
    try:
        tipo = listas.formato(caminho)
    except listas.ListaInvalida as erro:
        raise RelatorioInvalido(str(erro)) from erro
    if tipo in ("zip-quebrado", "zip-desconhecido"):
        raise RelatorioInvalido(f"o arquivo {caminho.name} parece danificado. Abra-o no Excel ou "
                                "no Word e salve de novo (.xlsx ou .docx).")
    if tipo == "doc":
        raise RelatorioInvalido("documento antigo do Word (.doc): abra-o no Word e salve como "
                                ".docx, ou exporte o relatório em Excel ou PDF.")
    reconhecedor = Reconhecedor(regras, "arquivo", "", origem=caminho.name, fonte="arquivo",
                                estrito=False, agora=agora)
    soltas: list[str] = []
    try:
        if tipo == "xlsx":
            tabelas = _de_xlsx(caminho)
        elif tipo == "xls":
            tabelas = _de_xls(caminho)
        elif tipo == "ods":
            tabelas = _de_ods(caminho)
        elif tipo == "html":
            tabelas = _de_html(caminho)
        elif tipo == "docx":
            tabelas, soltas = _de_docx(caminho)
        elif tipo == "pdf":
            tabelas, soltas = _de_pdf(caminho, regras)
        else:
            tabelas = _de_csv(caminho)
    except ImportError as erro:  # pragma: no cover - instalação incompleta
        raise RelatorioInvalido(f"falta um componente para ler este arquivo ({erro.name}). "
                                "Instale o Helestron de novo com o Helestron-Setup.") from erro
    r = reconhecedor.reconhecer(tabelas)
    if not r.reconhecida and soltas:
        _linhas_soltas(soltas, reconhecedor, r)
    if not r.reconhecida:
        raise RelatorioInvalido(
            f"li o arquivo {caminho.name}, mas não encontrei uma tabela de audiências. O "
            "relatório precisa ter, ao menos, a data da audiência e o número do processo (ou o "
            "tipo da audiência), com o cabeçalho das colunas.")
    log.info("Relatório %s: %d audiência(s), %d linha(s) ignorada(s).", caminho.name,
             len(r.audiencias), r.ignoradas)
    return r
