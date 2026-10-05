"""A pauta em Excel (seção 8.9 da especificação), com openpyxl.

    Documentos\\Helestron\\Pauta\\Pauta de audiências 2026-10-01 a 2026-10-31.xlsx

Três abas:

* Pauta - título e período nas linhas 1 e 2, cabeçalho na 4 (fundo navy,
  letra branca), uma audiência por linha. Data e hora são data e hora DO
  EXCEL (ordenam e filtram como tal), não texto. Filtro automático,
  cabeçalho e as quatro primeiras colunas congelados, larguras ajustadas,
  zebra leve, a audiência de hoje em azul-claro, a cancelada em cinza e
  tachada, a situação colorida. O link da sala virtual é clicável;
* Resumo - quantas por dia, por tipo e por situação;
* Alterações - o histórico das audiências do período.

SIGILO: a planilha fica fora do acervo (não vai para a IA), mas circula -
é impressa, vai por e-mail. Por isso as partes e as observações (onde o
portal costuma pôr o nome do réu preso, da vítima, do advogado) dos
processos em segredo de justiça saem como "(segredo de justiça)", a menos
que o usuário marque "incluir as partes dos sigilosos". Vale também para o
histórico e para o texto da busca no topo (o nome da parte procurada não
aparece quando o resultado tem processo sigiloso).

CARACTERES DE CONTROLE: o texto que vem do PDF, do HTML ou da planilha do
portal pode trazer um caractere de controle (o de código 2, por exemplo)
que o Excel não aceita; ele é tirado da célula (e não derruba a planilha
inteira).

FÓRMULAS: os textos vêm de fora do gabinete (o nome da parte é digitado
pelo advogado no peticionamento). Texto que começa com "=" é gravado como
TEXTO (e com o apóstrofo do Excel), nunca como fórmula: a planilha não
calcula =WEBSERVICE(...) nem =HYPERLINK(...) de ninguém ao ser aberta.
"""

from __future__ import annotations

import logging
import os
import re
import tempfile
from datetime import date, datetime, time
from pathlib import Path

from . import modelos
from .modelos import CAMPOS_MASCARADOS, MASCARA_SIGILO, NOMES_SISTEMA

log = logging.getLogger("pauta.exportacao")

NAVY = "12284A"
NAVY_ESCURO = "0B1A33"
CINZA_TEXTO = "6B7280"
CINZA_ZEBRA = "F4F6F9"
CINZA_BORDA = "E5E7EB"
AZUL_HOJE = "E8F1FF"
AMBAR = "C27C0E"
VERDE = "1F9D55"
VERMELHO = "D93A3A"

COLUNAS = [
    # (título, chave, largura mínima, largura máxima)
    ("Data", "data", 11, 12),
    ("Dia da semana", "dia_semana", 13, 15),
    ("Hora", "hora", 7, 8),
    ("Processo", "processo", 26, 30),
    ("Classe", "classe", 14, 34),
    ("Partes", "partes", 18, 50),
    ("Tipo de audiência", "tipo", 16, 24),
    ("Situação", "situacao", 12, 16),
    ("Local", "local", 12, 34),
    ("Magistrado/Conciliador", "magistrado", 14, 30),
    ("Sistema", "sistema", 9, 20),
    ("Tribunal", "tribunal", 9, 10),
    ("Link", "link", 8, 36),
    ("Observações", "observacoes", 14, 44),
]
LINHA_CABECALHO = 4
# O que o formato do Excel não aceita numa célula (openpyxl ILLEGAL_CHARACTERS_RE)
_ILEGAIS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
COR_SITUACAO = {"Redesignada": AMBAR, "Suspensa": AMBAR, "Realizada": VERDE,
                "Não realizada": VERMELHO, "Cancelada": CINZA_TEXTO}
NOMES_ALTERACAO = {"nova": "Nova", "alterada": "Alterada", "cancelada": "Cancelada",
                   "removida": "Saiu da pauta"}
NOMES_CAMPO = {"data": "Data", "hora": "Hora", "tipo": "Tipo", "situacao": "Situação",
               "local": "Local", "link": "Link", "partes": "Partes", "magistrado": "Magistrado",
               "classe": "Classe", "observacoes": "Observações", "processo": "Processo"}


def nome_do_arquivo(de: date, ate: date) -> str:
    return f"Pauta de audiências {de.isoformat()} a {ate.isoformat()}.xlsx"


def caminho_livre(pasta: Path, nome: str) -> Path:
    """O nome pedido; se já existir (ou estiver aberto no Excel), "(2)", "(3)"..."""
    alvo = pasta / nome
    if not alvo.exists():
        return alvo
    base, extensao = os.path.splitext(nome)
    for i in range(2, 1000):
        alvo = pasta / f"{base} ({i}){extensao}"
        if not alvo.exists():
            return alvo
    return pasta / f"{base} {datetime.now():%Y%m%d-%H%M%S}{extensao}"


def _data(valor) -> date | None:
    if isinstance(valor, date):
        return valor if not isinstance(valor, datetime) else valor.date()
    return modelos.ler_data(valor)


def _hora(texto) -> time | None:
    h = modelos.ler_hora(texto)
    if not h:
        return None
    return time(int(h[:2]), int(h[3:5]))


def _seguro(valor):
    """O texto sem os caracteres de controle que o Excel não aceita."""
    return _ILEGAIS.sub("", valor) if isinstance(valor, str) else valor


def _br(valor: str) -> str:
    """Data ISO de um campo alterado em dd/mm/aaaa (o resto como veio)."""
    d = modelos.ler_data(valor) if isinstance(valor, str) and len(valor) == 10 else None
    return d.strftime("%d/%m/%Y") if d and valor[4:5] == "-" else (valor or "—")


def descrever_campos(alteracao: dict, incluir_sigilosos: bool) -> str:
    audiencia = alteracao.get("audiencia") or {}
    partes = []
    for c in alteracao.get("campos") or []:
        campo = c.get("campo", "")
        antes, depois = c.get("antes") or "", c.get("depois") or ""
        if campo in CAMPOS_MASCARADOS and audiencia.get("sigiloso") and not incluir_sigilosos:
            antes = depois = MASCARA_SIGILO
        nome = NOMES_CAMPO.get(campo, campo.capitalize())
        partes.append(f"{nome}: {_br(antes)} → {_br(depois)}")
    return "; ".join(partes)


def descrever_filtros(filtros: dict | None, ocultar_busca: bool = False) -> str:
    """'ocultar_busca': o texto procurado não é escrito (pode ser o nome da
    parte de um processo em segredo de justiça)."""
    filtros = filtros or {}
    partes = []
    if filtros.get("sistema"):
        partes.append(f"sistema {NOMES_SISTEMA.get(filtros['sistema'], filtros['sistema'])}")
    if filtros.get("situacao"):
        partes.append(f"situação {filtros['situacao']}")
    if filtros.get("busca"):
        partes.append("busca por texto (omitido por causa do segredo de justiça)" if ocultar_busca
                      else f"busca “{filtros['busca']}”")
    return ", ".join(partes)


def _escrever(aba, linha: int, coluna: int, valor):
    """Grava a célula; texto que começa com "=" fica TEXTO, nunca fórmula, e
    sem os caracteres de controle que o Excel não aceita."""
    valor = _seguro(valor)
    c = aba.cell(row=linha, column=coluna, value=valor)
    if isinstance(valor, str) and valor.startswith("="):
        c.data_type = "s"
        try:
            c.quotePrefix = True       # o Excel não o transforma em fórmula nem ao editar
        except Exception:
            pass
    return c


# ================================================================ planilha
def exportar(audiencias: list[dict], alteracoes: list[dict], de: date, ate: date, pasta: Path,
             incluir_partes_sigilosos: bool = False, filtros: dict | None = None,
             agora: datetime | None = None) -> Path:
    """Grava a planilha e devolve o caminho. 'audiencias': dicts da API
    (com 'sigiloso' já apurado - portal ou pasta de sigilosos)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    agora = (agora or datetime.now()).replace(microsecond=0)
    hoje = agora.date()
    pasta = Path(pasta)
    pasta.mkdir(parents=True, exist_ok=True)
    gerado = f"Gerado pelo Helestron em {agora:%d/%m/%Y %H:%M}"
    periodo = (f"De {de:%d/%m/%Y} a {ate:%d/%m/%Y}" if de != ate else f"{de:%d/%m/%Y}")
    total = len(audiencias)
    resumo_txt = f"{periodo} · {total} audiência{'s' if total != 1 else ''}"
    # a busca pelo nome de uma parte, com processo sigiloso no resultado e as partes
    # mascaradas, não pode reaparecer escrita no topo da planilha
    filtros_txt = descrever_filtros(filtros, ocultar_busca=not incluir_partes_sigilosos and any(
        a.get("sigiloso") for a in audiencias))
    if filtros_txt:
        resumo_txt += f" · filtros: {filtros_txt}"
    resumo_txt = _seguro(resumo_txt)          # o texto da busca vem de quem digitou

    fina = Side(style="thin", color=CINZA_BORDA)
    borda = Border(bottom=fina)
    cabecalho_fonte = Font(bold=True, color="FFFFFF", size=11)
    cabecalho_fundo = PatternFill("solid", fgColor=NAVY)
    titulo_fonte = Font(bold=True, size=16, color=NAVY_ESCURO)
    sub_fonte = Font(size=11, color=CINZA_TEXTO)
    rodape_fonte = Font(italic=True, size=9, color=CINZA_TEXTO)
    zebra = PatternFill("solid", fgColor=CINZA_ZEBRA)
    fundo_hoje = PatternFill("solid", fgColor=AZUL_HOJE)
    topo = Alignment(vertical="top")
    quebra = Alignment(vertical="top", wrap_text=True)

    livro = Workbook()
    livro.properties.creator = "Helestron"
    livro.properties.title = f"Pauta de audiências — {periodo}"

    # ------------------------------------------------------------ Pauta
    aba = livro.active
    aba.title = "Pauta"
    aba["A1"] = "Pauta de audiências"
    aba["A1"].font = titulo_fonte
    aba["A2"] = resumo_txt
    aba["A2"].font = sub_fonte
    aba.row_dimensions[1].height = 24
    for j, (titulo, _chave, _mn, _mx) in enumerate(COLUNAS, start=1):
        c = aba.cell(row=LINHA_CABECALHO, column=j, value=titulo)
        c.font = cabecalho_fonte
        c.fill = cabecalho_fundo
        c.alignment = Alignment(vertical="center", wrap_text=True)
    aba.row_dimensions[LINHA_CABECALHO].height = 22
    larguras = [len(t) + 2 for t, *_ in COLUNAS]

    linha = LINHA_CABECALHO
    for i, a in enumerate(audiencias):
        linha += 1
        d = _data(a.get("data"))
        if not incluir_partes_sigilosos:
            a = modelos.mascarar_sigiloso(a)          # partes e observações
        valores = {
            "data": d, "dia_semana": modelos.dia_da_semana(d) if d else "",
            "hora": _hora(a.get("hora")), "processo": a.get("processo") or "",
            "classe": a.get("classe") or "", "partes": a.get("partes") or "",
            "tipo": a.get("tipo") or "", "situacao": a.get("situacao") or "",
            "local": a.get("local") or "", "magistrado": a.get("magistrado") or "",
            "sistema": NOMES_SISTEMA.get(a.get("sistema") or "", a.get("sistema") or ""),
            "tribunal": a.get("tribunal") or "", "link": a.get("link") or "",
            "observacoes": a.get("observacoes") or "",
        }
        valores = {k: _seguro(v) for k, v in valores.items()}
        situacao = valores["situacao"]
        cancelada = situacao == "Cancelada"
        fundo = fundo_hoje if d == hoje else (zebra if i % 2 else None)
        for j, (_t, chave, _mn, _mx) in enumerate(COLUNAS, start=1):
            valor = valores[chave]
            c = _escrever(aba, linha, j, valor if valor not in ("",) else None)
            c.alignment = quebra if chave in ("partes", "observacoes", "classe", "local") else topo
            c.border = borda
            if fundo is not None:
                c.fill = fundo
            if chave == "data" and d is not None:
                c.number_format = "dd/mm/yyyy"
            elif chave == "hora" and valor is not None:
                c.number_format = "hh:mm"
            elif chave == "processo":
                c.font = Font(name="Consolas", size=10, strike=cancelada,
                              color=CINZA_TEXTO if cancelada else None)
            elif chave == "link" and valor:
                # o endereço inteiro, e não "Abrir a sala": a planilha reimportada (ou
                # copiada para outro lugar) não perde o link
                c.hyperlink = valor
                c.font = Font(color="0A66E8", underline="single", strike=cancelada)
            if chave == "situacao" and situacao in COR_SITUACAO:
                c.font = Font(bold=True, color=COR_SITUACAO[situacao], strike=cancelada)
            elif cancelada and chave not in ("processo", "link"):
                c.font = Font(color=CINZA_TEXTO, strike=True)
            texto = (c.value.strftime("%d/%m/%Y") if isinstance(c.value, (date, datetime))
                     else c.value.strftime("%H:%M") if isinstance(c.value, time)
                     else str(c.value or ""))
            larguras[j - 1] = max(larguras[j - 1], min(len(texto) + 2, 80))
    if not audiencias:
        linha += 1
        aba.cell(row=linha, column=1, value="Nenhuma audiência no período.").font = sub_fonte
    ultima = max(linha, LINHA_CABECALHO + 1)
    for j, (_t, _c, minimo, maximo) in enumerate(COLUNAS, start=1):
        aba.column_dimensions[get_column_letter(j)].width = max(minimo, min(maximo,
                                                                          larguras[j - 1]))
    aba.auto_filter.ref = f"A{LINHA_CABECALHO}:{get_column_letter(len(COLUNAS))}{ultima}"
    aba.freeze_panes = f"E{LINHA_CABECALHO + 1}"
    aba.cell(row=ultima + 2, column=1, value=gerado).font = rodape_fonte
    _impressao(aba, gerado, f"{LINHA_CABECALHO}:{LINHA_CABECALHO}")

    # ----------------------------------------------------------- Resumo
    resumo = livro.create_sheet("Resumo")
    resumo["A1"] = "Resumo da pauta"
    resumo["A1"].font = titulo_fonte
    resumo["A2"] = resumo_txt
    resumo["A2"].font = sub_fonte
    por_dia: dict[date, int] = {}
    por_tipo: dict[str, int] = {}
    por_situacao: dict[str, int] = {}
    for a in audiencias:
        d = _data(a.get("data"))
        if d is not None:
            por_dia[d] = por_dia.get(d, 0) + 1
        por_tipo[a.get("tipo") or "Outra"] = por_tipo.get(a.get("tipo") or "Outra", 0) + 1
        s = a.get("situacao") or "Designada"
        por_situacao[s] = por_situacao.get(s, 0) + 1

    def bloco(coluna: int, titulos: list[str], linhas: list[list], formatos: list[str | None]):
        for k, titulo in enumerate(titulos):
            c = resumo.cell(row=LINHA_CABECALHO, column=coluna + k, value=titulo)
            c.font = cabecalho_fonte
            c.fill = cabecalho_fundo
        r = LINHA_CABECALHO
        for n, valores in enumerate(linhas):
            r += 1
            for k, valor in enumerate(valores):
                c = _escrever(resumo, r, coluna + k, valor)
                c.border = borda
                if n % 2:
                    c.fill = zebra
                if formatos[k]:
                    c.number_format = formatos[k]
        r += 1
        resumo.cell(row=r, column=coluna, value="Total").font = Font(bold=True)
        c = resumo.cell(row=r, column=coluna + len(titulos) - 1,
                        value=sum(v[-1] for v in linhas))
        c.font = Font(bold=True)
        c.border = Border(top=Side(style="thin", color=NAVY))

    bloco(1, ["Data", "Dia da semana", "Quantidade"],
          [[d, modelos.dia_da_semana(d), n] for d, n in sorted(por_dia.items())],
          ["dd/mm/yyyy", None, "0"])
    ordem_tipos = {t: i for i, t in enumerate(modelos.TIPOS)}
    bloco(5, ["Tipo de audiência", "Quantidade"],
          [[t, n] for t, n in sorted(por_tipo.items(), key=lambda x: (ordem_tipos.get(x[0], 99),
                                                                       x[0]))],
          [None, "0"])
    ordem_sit = {s: i for i, s in enumerate(modelos.SITUACOES)}
    bloco(8, ["Situação", "Quantidade"],
          [[s, n] for s, n in sorted(por_situacao.items(), key=lambda x: (ordem_sit.get(x[0], 99),
                                                                           x[0]))],
          [None, "0"])
    for letra, largura in (("A", 12), ("B", 16), ("C", 12), ("D", 3), ("E", 24), ("F", 12),
                           ("G", 3), ("H", 16), ("I", 12)):
        resumo.column_dimensions[letra].width = largura
    fim_resumo = LINHA_CABECALHO + max(len(por_dia), len(por_tipo), len(por_situacao)) + 3
    resumo.cell(row=fim_resumo, column=1, value=gerado).font = rodape_fonte
    _impressao(resumo, gerado, None)

    # ------------------------------------------------------- Alterações
    hist = livro.create_sheet("Alterações")
    hist["A1"] = "Alterações da pauta"
    hist["A1"].font = titulo_fonte
    hist_txt = (f"Audiências marcadas de {de:%d/%m/%Y} a {ate:%d/%m/%Y}" if de != ate
                else f"Audiências marcadas para {de:%d/%m/%Y}")
    if filtros_txt:                 # o histórico segue os mesmos filtros da aba Pauta
        hist_txt += f" · filtros: {filtros_txt}"
    hist["A2"] = _seguro(hist_txt)
    hist["A2"].font = sub_fonte
    titulos = ["Quando", "Alteração", "Processo", "Data da audiência", "Hora", "O que mudou",
               "Sistema", "Tribunal"]
    for k, titulo in enumerate(titulos, start=1):
        c = hist.cell(row=LINHA_CABECALHO, column=k, value=titulo)
        c.font = cabecalho_fonte
        c.fill = cabecalho_fundo
    r = LINHA_CABECALHO
    for n, alt in enumerate(alteracoes):
        r += 1
        aud = alt.get("audiencia") or {}
        try:
            quando = datetime.fromisoformat(str(alt.get("quando") or ""))
        except ValueError:
            quando = None
        valores = [quando, NOMES_ALTERACAO.get(alt.get("tipo"), alt.get("tipo") or ""),
                   aud.get("processo") or "", _data(aud.get("data")), _hora(aud.get("hora")),
                   descrever_campos(alt, incluir_partes_sigilosos),
                   NOMES_SISTEMA.get(aud.get("sistema") or "", aud.get("sistema") or ""),
                   aud.get("tribunal") or ""]
        formatos = ["dd/mm/yyyy hh:mm", None, None, "dd/mm/yyyy", "hh:mm", None, None, None]
        for k, valor in enumerate(valores, start=1):
            c = _escrever(hist, r, k, valor if valor != "" else None)
            c.border = borda
            c.alignment = quebra if k == 6 else topo
            if n % 2:
                c.fill = zebra
            if formatos[k - 1] and valor is not None:
                c.number_format = formatos[k - 1]
            if k == 2 and alt.get("tipo") in ("cancelada", "removida"):
                c.font = Font(bold=True, color=VERMELHO if alt.get("tipo") == "cancelada"
                              else CINZA_TEXTO)
    if not alteracoes:
        r += 1
        aviso = "Nenhuma alteração registrada para as audiências do período."
        hist.cell(row=r, column=1, value=aviso).font = sub_fonte
    for letra, largura in (("A", 17), ("B", 14), ("C", 27), ("D", 16), ("E", 7), ("F", 60),
                           ("G", 11), ("H", 9)):
        hist.column_dimensions[letra].width = largura
    if alteracoes:
        hist.auto_filter.ref = f"A{LINHA_CABECALHO}:H{r}"
    hist.freeze_panes = f"A{LINHA_CABECALHO + 1}"
    hist.cell(row=r + 2, column=1, value=gerado).font = rodape_fonte
    _impressao(hist, gerado, f"{LINHA_CABECALHO}:{LINHA_CABECALHO}")

    # ---------------------------------------------------------- gravação
    destino = caminho_livre(pasta, nome_do_arquivo(de, ate))
    descritor, temporario = tempfile.mkstemp(prefix=".helestron-", suffix=".xlsx", dir=str(pasta))
    os.close(descritor)
    try:
        livro.save(temporario)
        os.replace(temporario, destino)
    except BaseException:
        try:
            os.unlink(temporario)
        except OSError:
            pass
        raise
    log.info("Pauta exportada: %s (%d audiências).", destino, total)
    return destino


def _impressao(aba, rodape: str, titulos: str | None) -> None:
    """Paisagem, uma página de largura, cabeçalho repetido e o rodapé."""
    try:
        aba.page_setup.orientation = "landscape"
        aba.page_setup.paperSize = aba.PAPERSIZE_A4
        aba.page_setup.fitToWidth = 1
        aba.page_setup.fitToHeight = 0
        aba.sheet_properties.pageSetUpPr.fitToPage = True
        if titulos:
            aba.print_title_rows = titulos
        aba.oddFooter.left.text = rodape
        aba.oddFooter.left.size = 8
        aba.oddFooter.right.text = "Página &P de &N"
        aba.oddFooter.right.size = 8
    except Exception as erro:          # impressão é enfeite: nunca derruba a exportação
        log.debug("configuração de impressão: %s", erro)
